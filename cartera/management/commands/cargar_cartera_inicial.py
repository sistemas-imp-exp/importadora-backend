"""
Carga inicial de Cartera: el saldo de cada cliente en IMPORTADORA y MARISCOS
SELECTOS a la fecha de corte, desde el Excel de saldos finales del área
(columnas CLIENTE, IMPORTADORA, SELECTOS; la primera hoja que las tenga).

Solo saldos: celular, pagaré y último pago quedan en blanco. Los nombres se
buscan en el catálogo de clientes de Inventario (sin acentos ni espacios de
más) y, si no, en EQUIVALENCIAS (nombres cortos del Excel ya confirmados con
el área). Un nombre sin cliente y con saldo se da de alta; con saldo 0 se omite.

Modo prueba por defecto (valida y reporta sin guardar); --aplicar guarda.
Se niega si ya hay saldos iniciales cargados.

    manage.py cargar_cartera_inicial "docs/CARTERA/SALDOS FINALES MES DE SEPTIEMBRE.xlsx" --fecha-corte 2026-09-30 [--aplicar]
"""
import unicodedata
from decimal import ROUND_HALF_UP, Decimal

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.dateparse import parse_date
from openpyxl import load_workbook

from inventario.models import Cliente, Empresa

from ...models import SaldoInicialCliente

EMPRESAS = {'IMPORTADORA': 'IMPORTADORA', 'SELECTOS': 'MARISCOS SELECTOS'}
CENTAVO = Decimal('0.01')

# Nombre en el Excel de saldos → nombre en el catálogo (confirmado con el área el 2026-10-29).
EQUIVALENCIAS = {
    'HERAY': 'HERAY ACERO MORENO',
    'JOSE LUIS CHILIO': 'JOSE LUIS UYOA VAZQUEZ (CHILIO)',
    'JUAN CARLOS RUIZ': 'JUAN CARLOS RUIZ (XHUMIJAMA)',
    'LILIA CHACON': 'LILIA CRUZ CHACON',
    'MAJOS/PORTELA': 'PRODUCTOS MAJOS',
    'VASA': 'JOSE VAZQUEZ (VASA)',
    'JESUS HERRERA': 'JESUS GABRIEL HERRERA',
    'DIEGO FIDEL VICTORIO DAMIAN': 'DIEGO FIDEL VICTORIO DAMIAN (PAISA)',
    'SR. DANIEL GUATEMALA - GONZALO DE LA CRUZ': 'SR. DANIEL GUATEMALA',
    'LUIS ALBERTO HERNANDEZ OCHOA - ELOHIM': 'PESCADERIA ELOHIM',
    'ARMANDO TOLEDO-BAOS': 'GRUPO VAROCH',
}


class Rollback(Exception):
    pass


def _nombre(valor):
    return ' '.join(str(valor).split()).upper()


def _clave(nombre):
    """Para comparar nombres sin acentos, espacios ni mayúsculas de por medio."""
    sin_acentos = unicodedata.normalize('NFD', nombre).encode('ascii', 'ignore').decode()
    return ' '.join(sin_acentos.upper().split())


def _dinero(valor):
    if valor in (None, ''):
        return Decimal('0')
    if not isinstance(valor, (int, float, Decimal)):
        raise ValueError(f'no es un número: {valor!r}')
    monto = Decimal(str(valor)).quantize(CENTAVO, rounding=ROUND_HALF_UP)
    return Decimal('0') if monto == 0 else monto


class Command(BaseCommand):
    help = 'Carga los saldos iniciales de Cartera por cliente y empresa desde el Excel de saldos finales del área.'

    def add_arguments(self, parser):
        parser.add_argument('archivo')
        parser.add_argument('--fecha-corte', required=True, help='AAAA-MM-DD: último día incluido en el Excel (saldo al cierre); el sistema captura desde el día siguiente.')
        parser.add_argument('--aplicar', action='store_true', help='Guardar (sin esto solo valida y reporta).')
        parser.add_argument('--usuario', help='Usuario que queda registrado (por omisión, el primer superusuario).')

    def handle(self, *args, **opciones):
        corte = parse_date(opciones['fecha_corte'])
        if not corte:
            raise CommandError('La fecha de corte no es válida (AAAA-MM-DD).')
        usuario = self._usuario(opciones.get('usuario'))
        empresas = self._empresas()
        if SaldoInicialCliente.objects.exists():
            raise CommandError('Ya hay saldos iniciales de Cartera cargados: no se vuelve a cargar.')

        filas = self._leer(opciones['archivo'])
        self._verificar_totales(filas)
        self._asignar_clientes(filas)

        try:
            with transaction.atomic():
                resumen = self._cargar(filas, empresas, corte, usuario)
                if not opciones['aplicar']:
                    raise Rollback
        except Rollback:
            pass

        self._reportar(resumen, filas, corte, opciones['aplicar'])

    # --- lectura --------------------------------------------------------

    def _usuario(self, username):
        Usuario = get_user_model()
        if username:
            try:
                return Usuario.objects.get(username=username)
            except Usuario.DoesNotExist:
                raise CommandError(f'No existe el usuario {username}.')
        usuario = Usuario.objects.filter(is_superuser=True).order_by('id').first()
        if usuario is None:
            raise CommandError('No hay superusuario: indica --usuario.')
        return usuario

    def _empresas(self):
        empresas = {}
        for columna, nombre in EMPRESAS.items():
            try:
                empresas[columna] = Empresa.objects.get(nombre=nombre)
            except Empresa.DoesNotExist:
                raise CommandError(f'No existe la empresa {nombre}.')
        return empresas

    def _leer(self, archivo):
        try:
            libro = load_workbook(archivo, data_only=True)
        except FileNotFoundError:
            raise CommandError(f'No se encontró el archivo {archivo}.')

        hoja, encabezado, columnas = None, None, None
        for candidata in libro.worksheets:
            for numero in range(1, 15):
                valores = [_nombre(c.value) if c.value is not None else '' for c in candidata[numero]]
                if {'CLIENTE', 'IMPORTADORA', 'SELECTOS'} <= set(valores):
                    hoja, encabezado = candidata, numero
                    columnas = {v: i for i, v in enumerate(valores) if v}
                    break
            if hoja is not None:
                break
        if hoja is None:
            raise CommandError('No se encontró una hoja con las columnas CLIENTE, IMPORTADORA y SELECTOS.')
        self.hoja = hoja.title.strip()

        c_cliente, c_imp, c_sel = columnas['CLIENTE'], columnas['IMPORTADORA'], columnas['SELECTOS']
        filas, errores = [], []
        self.totales_hoja = None
        for numero, fila in enumerate(hoja.iter_rows(min_row=encabezado + 1, values_only=True), start=encabezado + 1):
            nombre = fila[c_cliente]
            if not isinstance(nombre, str) or not nombre.strip():
                continue
            if _clave(nombre).startswith('TOTAL'):
                # Fin de la tabla: el renglón de totales sirve para verificar la carga.
                self.totales_hoja = {'IMPORTADORA': _dinero(fila[c_imp]), 'SELECTOS': _dinero(fila[c_sel])}
                break
            try:
                saldos = {'IMPORTADORA': _dinero(fila[c_imp]), 'SELECTOS': _dinero(fila[c_sel])}
            except ValueError as error:
                errores.append(f'Fila {numero} ({nombre.strip()}): {error}')
                continue
            filas.append({'fila': numero, 'nombre': _nombre(nombre), 'saldos': saldos})
        if errores:
            raise CommandError('Valores no numéricos:\n  ' + '\n  '.join(errores))
        return filas

    def _verificar_totales(self, filas):
        """La suma de los clientes debe ser igual al renglón TOTALES de la hoja (si lo trae)."""
        if self.totales_hoja is None:
            return
        for columna, esperado in self.totales_hoja.items():
            suma = sum((f['saldos'][columna] for f in filas), Decimal('0'))
            # Cada saldo se redondea a centavos: se tolera un centavo por cliente.
            if abs(suma - esperado) > Decimal('0.01') * len(filas):
                raise CommandError(
                    f'La suma de {columna} ({suma:,.2f}) no cuadra con el renglón TOTALES ({esperado:,.2f}).'
                )

    def _asignar_clientes(self, filas):
        """Cliente del catálogo de cada fila (None = alta nueva); dos filas al mismo cliente es error."""
        catalogo = {_clave(c.nombre): c for c in Cliente.objects.all()}
        equivalencias = {_clave(k): _clave(v) for k, v in EQUIVALENCIAS.items()}
        vistos, repetidos = {}, []
        for fila in filas:
            clave = _clave(fila['nombre'])
            fila['cliente'] = catalogo.get(clave) or catalogo.get(equivalencias.get(clave, ''))
            destino = _clave(fila['cliente'].nombre) if fila['cliente'] else clave
            if destino in vistos:
                repetidos.append(f"{fila['nombre']} (filas {vistos[destino]} y {fila['fila']})")
            vistos[destino] = fila['fila']
        if repetidos:
            raise CommandError('Clientes repetidos en la hoja; corrígelos antes de cargar:\n  ' + '\n  '.join(repetidos))

    # --- carga ----------------------------------------------------------

    def _cargar(self, filas, empresas, corte, usuario):
        resumen = {'existentes': 0, 'nuevos': [], 'omitidos': [], 'saldos': 0, 'con_saldo': 0,
                   'totales': {k: Decimal('0') for k in EMPRESAS}}
        for fila in filas:
            con_saldo = {columna: monto for columna, monto in fila['saldos'].items() if monto != 0}
            cliente = fila['cliente']
            if cliente is None:
                if not con_saldo:
                    resumen['omitidos'].append(fila['nombre'])
                    continue
                cliente = Cliente.objects.create(nombre=fila['nombre'], creado_por=usuario)
                resumen['nuevos'].append(fila['nombre'])
            else:
                resumen['existentes'] += 1
            for columna, monto in con_saldo.items():
                SaldoInicialCliente.objects.create(
                    cliente=cliente, empresa=empresas[columna], monto=monto, fecha=corte, usuario=usuario,
                )
                resumen['saldos'] += 1
                resumen['totales'][columna] += monto
            resumen['con_saldo'] += bool(con_saldo)
        return resumen

    def _reportar(self, resumen, filas, corte, aplicado):
        escribir = self.stdout.write
        escribir(f'Hoja {self.hoja}: {len(filas)} clientes · corte al {corte:%d/%m/%Y}')
        escribir(f"Del catálogo: {resumen['existentes']} · nuevos: {len(resumen['nuevos'])} · "
                 f"sin cliente y en 0 (omitidos): {len(resumen['omitidos'])}")
        for nombre in resumen['nuevos']:
            escribir(f'  + {nombre}')
        for nombre in resumen['omitidos']:
            escribir(f'  - {nombre}')
        escribir(f"Clientes con saldo: {resumen['con_saldo']} · saldos iniciales: {resumen['saldos']}")
        for columna, total in resumen['totales'].items():
            escribir(f'  {EMPRESAS[columna]}: {total:,.2f}')
        escribir(f"  Conjunto: {sum(resumen['totales'].values()):,.2f}")
        if self.totales_hoja is not None:
            escribir('Cuadra con el renglón TOTALES de la hoja.')
        if aplicado:
            escribir(self.style.SUCCESS('Cartera inicial cargada.'))
        else:
            escribir(self.style.WARNING('Modo prueba: no se guardó nada. Usa --aplicar para cargar.'))
