import datetime
import tempfile
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from inventario.models import Cliente, Empresa
from security.permissions import AREA_CARTERA
from security.testing import crear_usuario_con_area
from treasury.models import Divisa, MovimientoDivisa, MovimientoTesoreria

from .models import AjusteCliente, ClienteCartera, PagoCliente, SaldoInicialCliente
from .saldos import saldos_por_cliente

HOY = timezone.localdate()
CORTE = HOY - timedelta(days=10)


class BaseCartera(APITestCase):
    def setUp(self):
        self.usuario = crear_usuario_con_area("cobranza", AREA_CARTERA)
        self.client.force_authenticate(user=self.usuario)
        self.imp, _ = Empresa.objects.get_or_create(nombre="IMPORTADORA")
        self.sel, _ = Empresa.objects.get_or_create(nombre="MARISCOS SELECTOS")
        self.cliente = Cliente.objects.create(nombre="CABEZA DE TORO")
        SaldoInicialCliente.objects.create(cliente=self.cliente, empresa=self.imp, monto=Decimal("1000"),
                                           fecha=CORTE, usuario=self.usuario)
        SaldoInicialCliente.objects.create(cliente=self.cliente, empresa=self.sel, monto=Decimal("-200"),
                                           fecha=CORTE, usuario=self.usuario)

    def _pago(self, monto="300", empresa=None, fecha=None, **extra):
        return self.client.post("/api/cartera/pagos/", {
            "cliente_id": self.cliente.id, "empresa_id": (empresa or self.imp).id, "tipo": "T",
            "fecha": (fecha or HOY).isoformat(), "monto": monto, "referencia": "BBVA 123", **extra,
        }, format="json")

    def _saldo(self, empresa):
        return saldos_por_cliente([self.cliente.id])[self.cliente.id][empresa.id]


class SaldosYPagosTests(BaseCartera):
    def test_pago_baja_el_saldo_y_cancelarlo_lo_devuelve(self):
        r = self._pago("300")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["usuario"]["username"], r.data["tipo_display"]), ("cobranza", "Transferencia"))
        self.assertEqual(self._saldo(self.imp), Decimal("700"))
        r = self.client.post(f"/api/cartera/pagos/{r.data['id']}/cancelar/", {"motivo": "Duplicado"}, format="json")
        self.assertEqual((r.status_code, r.data["cancelado"]), (200, True))
        self.assertEqual(self._saldo(self.imp), Decimal("1000"))

    def test_no_se_capturan_movimientos_anteriores_o_del_dia_de_corte(self):
        r = self._pago(fecha=CORTE)
        self.assertEqual(r.status_code, 400)
        self.assertIn("saldos iniciales", str(r.data["fecha"]))
        self.assertEqual(self._pago(fecha=HOY + timedelta(days=1)).status_code, 400)

    def test_lista_de_clientes_con_saldo_por_empresa_y_ultimo_pago(self):
        ClienteCartera.objects.create(cliente=self.cliente, celular="966 101 2335", pagare=True,
                                      pago_previo_fecha=CORTE, pago_previo_monto=Decimal("50"))
        Cliente.objects.create(nombre="SIN SALDO")
        r = self.client.get("/api/cartera/clientes/", {"con_saldo": "1"})
        self.assertEqual([c["nombre"] for c in r.data], ["CABEZA DE TORO"])
        fila = r.data[0]
        self.assertEqual((fila["saldos"][str(self.imp.id)], fila["saldos"][str(self.sel.id)], fila["saldo_conjunto"]),
                         ("1000.00", "-200.00", "800.00"))
        self.assertEqual((fila["celular"], fila["pagare"]), ("966 101 2335", True))
        # Sin pagos capturados, el último pago es el del Excel; con pagos, el del sistema.
        self.assertEqual(fila["ultimo_pago"], {"fecha": CORTE.isoformat(), "monto": "50.00"})
        self._pago("300")
        self.assertEqual(self.client.get(f"/api/cartera/clientes/{self.cliente.id}/").data["ultimo_pago"],
                         {"fecha": HOY.isoformat(), "monto": "300.00"})

    def test_alta_de_cliente_desde_cartera_va_al_catalogo_de_inventario(self):
        r = self.client.post("/api/cartera/clientes/", {"nombre": " heray  ", "celular": "961", "pagare": False,
                                                        "notas": "", "activo": True}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(Cliente.objects.get(pk=r.data["id"]).nombre, "HERAY")
        r = self.client.post("/api/cartera/clientes/", {"nombre": "Heray", "activo": True}, format="json")
        self.assertEqual(r.status_code, 400)


class IngresoTesoreriaTests(BaseCartera):
    def setUp(self):
        super().setUp()
        self.mxn = Divisa.objects.create(codigo="MXN", nombre="Peso", simbolo="$")
        cajero = get_user_model().objects.create_user("cajero")
        self.ingreso = MovimientoTesoreria.objects.create(
            fecha=HOY - timedelta(days=1), folio="9001", tipo="I", autorizo="CP", beneficiario="CABEZA DE TORO",
            concepto="Abono cliente", usuario=cajero,
        )
        MovimientoDivisa.objects.create(movimiento=self.ingreso, divisa=self.mxn, cantidad=Decimal("450"))

    def test_el_efectivo_se_toma_del_ingreso_de_tesoreria_una_sola_vez(self):
        disponibles = self.client.get("/api/cartera/ingresos-tesoreria/").data
        self.assertEqual([(i["folio"], i["monto"]) for i in disponibles], [("9001", "450.00")])

        r = self.client.post("/api/cartera/pagos/", {
            "cliente_id": self.cliente.id, "empresa_id": self.imp.id, "tipo": "A", "ingreso_tesoreria_id": self.ingreso.id,
        }, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["fecha"], r.data["monto"], r.data["ingreso_tesoreria"]["folio"]),
                         (self.ingreso.fecha.isoformat(), "450.00", "9001"))
        self.assertEqual(self.client.get("/api/cartera/ingresos-tesoreria/").data, [])

        r = self.client.post("/api/cartera/pagos/", {
            "cliente_id": self.cliente.id, "empresa_id": self.sel.id, "tipo": "A", "ingreso_tesoreria_id": self.ingreso.id,
        }, format="json")
        self.assertIn("ya está registrado", str(r.data))

    def test_efectivo_a_tesoreria_no_se_captura_a_mano(self):
        r = self._pago("450", tipo="A")
        self.assertEqual(r.status_code, 400)
        self.assertIn("ingreso de Tesorería", str(r.data))

    def test_cancelar_el_pago_libera_el_ingreso(self):
        pago = PagoCliente.objects.create(cliente=self.cliente, empresa=self.imp, tipo="A", fecha=self.ingreso.fecha,
                                          monto=Decimal("450"), ingreso_tesoreria=self.ingreso, usuario=self.usuario)
        self.client.post(f"/api/cartera/pagos/{pago.id}/cancelar/", {"motivo": "Era de otro cliente"}, format="json")
        self.assertEqual(len(self.client.get("/api/cartera/ingresos-tesoreria/").data), 1)


class AjustesYTraspasosTests(BaseCartera):
    def _ajuste(self, naturaleza, monto, empresa=None):
        return self.client.post("/api/cartera/ajustes/", {
            "cliente_id": self.cliente.id, "empresa_id": (empresa or self.imp).id, "naturaleza": naturaleza,
            "concepto": "  Cargo por   horas extras ", "fecha": HOY.isoformat(), "monto": monto, "notas": "",
        }, format="json")

    def test_cargo_sube_y_abono_baja(self):
        r = self._ajuste("C", "500")
        self.assertEqual((r.status_code, r.data["concepto"]), (201, "Cargo por horas extras"))
        AjusteCliente.objects.filter(pk=r.data["id"]).update(notas="Autorizó CP Lorenzo")
        renglon = self.client.get(f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/").data[-1]
        self.assertEqual(renglon["concepto"], "Cargo por horas extras · Autorizó CP Lorenzo")
        self._ajuste("A", "100")
        self.assertEqual(self._saldo(self.imp), Decimal("1400"))

    def test_traspaso_entre_empresas_y_se_cancela_completo(self):
        r = self.client.post("/api/cartera/ajustes/traspaso/", {
            "cliente_id": self.cliente.id, "origen_id": self.imp.id, "destino_id": self.sel.id,
            "concepto": "Ajuste Selectos", "fecha": HOY.isoformat(), "monto": "600",
        }, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual({a["empresa"]["nombre"]: a["naturaleza"] for a in r.data},
                         {"IMPORTADORA": "A", "MARISCOS SELECTOS": "C"})
        self.assertEqual((self._saldo(self.imp), self._saldo(self.sel)), (Decimal("400"), Decimal("400")))

        abono = AjusteCliente.objects.get(empresa=self.imp)
        r = self.client.put(f"/api/cartera/ajustes/{abono.id}/", {
            "cliente_id": self.cliente.id, "empresa_id": self.imp.id, "naturaleza": "A", "concepto": "X",
            "fecha": HOY.isoformat(), "monto": "1",
        }, format="json")
        self.assertEqual(r.status_code, 400)
        self.client.post(f"/api/cartera/ajustes/{abono.id}/cancelar/", {"motivo": "Error"}, format="json")
        self.assertEqual(AjusteCliente.objects.filter(cancelado=True).count(), 2)
        self.assertEqual((self._saldo(self.imp), self._saldo(self.sel)), (Decimal("1000"), Decimal("-200")))

    def test_traspaso_a_la_misma_empresa_no(self):
        r = self.client.post("/api/cartera/ajustes/traspaso/", {
            "cliente_id": self.cliente.id, "origen_id": self.imp.id, "destino_id": self.imp.id,
            "concepto": "X", "fecha": HOY.isoformat(), "monto": "1",
        }, format="json")
        self.assertEqual(r.status_code, 400)


class EstadoDeCuentaYReportesTests(BaseCartera):
    def test_estado_de_cuenta_con_saldo_corrido(self):
        self._pago("300")
        cancelado = self._pago("50").data["id"]
        self.client.post(f"/api/cartera/pagos/{cancelado}/cancelar/", {"motivo": "x"}, format="json")
        renglones = self.client.get(f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/").data
        self.assertEqual([(r["tipo"], r["saldo"]) for r in renglones],
                         [("inicial", "1000.00"), ("inicial", "800.00"), ("pago", "500.00"), ("pago", None)])
        solo_imp = self.client.get(f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/", {"empresa": self.sel.id}).data
        self.assertEqual([r["saldo"] for r in solo_imp], ["-200.00"])

    def test_descargas(self):
        self._pago("300")
        for ruta in (f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/pdf/",
                     f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/excel/",
                     "/api/cartera/reportes/saldos/pdf/", "/api/cartera/reportes/saldos/excel/"):
            with self.subTest(ruta=ruta):
                self.assertEqual(self.client.get(ruta).status_code, 200)
        hoja = load_workbook(BytesIO(self.client.get("/api/cartera/reportes/saldos/excel/").content)).active
        self.assertEqual([c.value for c in hoja[1]][:4], ["Cliente", "IMPORTADORA", "MARISCOS SELECTOS", "Saldo conjunto"])
        self.assertEqual([c.value for c in hoja[2]][:4], ["CABEZA DE TORO", 700, -200, 500])

    def test_pdf_de_saldos_completo_y_resumido_en_vertical(self):
        import re
        for parametros, horizontal, nombre in (({}, True, "saldos_clientes.pdf"),
                                               ({"resumido": "1"}, False, "saldos_clientes_resumido.pdf")):
            r = self.client.get("/api/cartera/reportes/saldos/pdf/", parametros)
            self.assertEqual(r.status_code, 200)
            self.assertIn(nombre, r["Content-Disposition"])
            ancho, alto = map(float, re.search(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", r.content).groups())
            self.assertEqual(ancho > alto, horizontal)

    def test_resumen_de_inicio(self):
        self._pago("300")
        datos = self.client.get("/api/cartera/resumen/").data
        empresas = {e["nombre"]: e for e in datos["empresas"]}
        self.assertEqual(empresas["IMPORTADORA"]["por_cobrar"], "700.00")
        self.assertEqual(empresas["MARISCOS SELECTOS"]["a_favor"], "200.00")
        self.assertEqual(datos["clientes_con_saldo"], 1)


def _excel_saldos(filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "SEPTIEMBRE "
    ws.append([None, "CLIENTE", "IMPORTADORA", "SELECTOS", "SALDO CONJUNTO"])
    for fila in filas:
        ws.append(fila)
    archivo = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(archivo.name)
    return archivo.name


class CargarCarteraInicialTests(TestCase):
    def setUp(self):
        get_user_model().objects.create_superuser("admin", password="x")
        Empresa.objects.get_or_create(nombre="IMPORTADORA")
        Empresa.objects.get_or_create(nombre="MARISCOS SELECTOS")
        Cliente.objects.create(nombre="HERAY ACERO MORENO")
        Cliente.objects.create(nombre="ALEXANDER VILLATORO")
        Cliente.objects.create(nombre="ANABEL PACHECO")
        self.ruta = _excel_saldos([
            [None, "HERAY", 18888781.04999999, 564240],
            [None, "ALEXANDER  VILLATORO", 375218.8799999999, 0],
            [None, "ANABEL PACHECO ", 4.097806538538862e-10, 0],
            [None, "MAXIMILIANO RAMIREZ", 0, -567145.57],
            [None, "NADIE EN CERO", 0, 0],
            [None, "TOTALES", 19263999.93, -2905.57],
            [None, "DESPUES DEL TOTAL", 5, 0],
        ])

    def _correr(self, *extra):
        salida = StringIO()
        call_command("cargar_cartera_inicial", self.ruta, "--fecha-corte", "2026-09-30", *extra, stdout=salida)
        return salida.getvalue()

    def test_modo_prueba_no_guarda(self):
        salida = self._correr()
        self.assertIn("Del catálogo: 3 · nuevos: 1 · sin cliente y en 0 (omitidos): 1", salida)
        self.assertIn("IMPORTADORA: 19,263,999.93", salida)
        self.assertIn("MARISCOS SELECTOS: -2,905.57", salida)
        self.assertIn("Cuadra con el renglón TOTALES", salida)
        self.assertFalse(SaldoInicialCliente.objects.exists())
        self.assertEqual(Cliente.objects.count(), 3)

    def test_aplica_saldos_con_equivalencias(self):
        self._correr("--aplicar")
        # "HERAY" del Excel es HERAY ACERO MORENO del catálogo (EQUIVALENCIAS).
        heray = SaldoInicialCliente.objects.filter(cliente__nombre="HERAY ACERO MORENO")
        self.assertEqual(sorted(heray.values_list("monto", flat=True)), [Decimal("564240.00"), Decimal("18888781.05")])
        self.assertEqual(heray.first().fecha, datetime.date(2026, 9, 30))
        # Los residuos de redondeo del Excel (4e-10) no generan saldo.
        self.assertFalse(SaldoInicialCliente.objects.filter(cliente__nombre="ANABEL PACHECO").exists())
        # Sin cliente: con saldo se da de alta, en 0 se omite.
        self.assertEqual(SaldoInicialCliente.objects.get(cliente__nombre="MAXIMILIANO RAMIREZ").monto, Decimal("-567145.57"))
        self.assertFalse(Cliente.objects.filter(nombre="NADIE EN CERO").exists())
        self.assertEqual(SaldoInicialCliente.objects.count(), 4)
        with self.assertRaisesMessage(CommandError, "Ya hay saldos iniciales"):
            self._correr("--aplicar")

    def test_si_no_cuadra_con_totales_no_carga(self):
        self.ruta = _excel_saldos([[None, "UNO", 100, 0], [None, "TOTALES", 150, 0]])
        with self.assertRaisesMessage(CommandError, "no cuadra"):
            self._correr()

    def test_clientes_repetidos_se_rechazan(self):
        self.ruta = _excel_saldos([[None, "JOSÉ PEREZ", 1, 0], [None, "jose  perez", 2, 0]])
        with self.assertRaisesMessage(CommandError, "repetidos"):
            self._correr()
        # También si un nombre corto apunta al mismo cliente que otro renglón.
        self.ruta = _excel_saldos([[None, "HERAY", 1, 0], [None, "HERAY ACERO MORENO", 2, 0]])
        with self.assertRaisesMessage(CommandError, "repetidos"):
            self._correr()


def _preparar_ventas(self):
    """Lotes de IMPORTADORA y SELECTOS a $100/kg y salidas de prueba (ver VentasTests)."""
    from inventario.models import Camara, Entrada, EntradaDetalle, LoteGeneral, Producto, Proveedor, Salida, SalidaDetalle
    self.SalidaDetalle = SalidaDetalle
    proveedor = Proveedor.objects.create(nombre="SAN LORENZO")
    self.camara = Camara.objects.create(nombre="REMAINS", tipo=Camara.TIPO_PROPIA)
    self.producto = Producto.objects.create(talla="41-50", tipo="FREEZADO")

    def lote(empresa, codigo):
        entrada = Entrada.objects.create(fecha=CORTE, empresa=empresa)
        general = LoteGeneral.objects.create(codigo=codigo, camara=self.camara, fecha_recibo=CORTE)
        return EntradaDetalle.objects.create(
            entrada=entrada, producto=self.producto, lote_proveedor="L", camara=self.camara, cajas=50,
            peso_por_caja=Decimal("20"), total_kilos=Decimal("1000"), costo_por_kilo=Decimal("100"),
            proveedor_origen=proveedor, lote_general=general,
        )

    def salida(folio, lineas, cliente=None, fecha=HOY):
        s = Salida.objects.create(folio_de_salida=folio, cliente=cliente or self.cliente, fecha=fecha, notas=f"N-{folio}")
        return [SalidaDetalle.objects.create(salida=s, producto=self.producto, entrada_detalle=l, camara=self.camara,
                                             cajas=1, total_kilos=Decimal(k), precio_x_kilo=p and Decimal(p),
                                             total_venta=t and Decimal(t))
                for l, k, p, t in lineas]

    self.lote_imp, self.lote_sel = lote(self.imp, "IMP1"), lote(self.sel, "S-1")
    # SAL-1: una línea con total capturado y otra solo con precio (kilos × precio).
    self.l1, self.l2 = salida("SAL-1", [(self.lote_imp, "100", "130", "13000"), (self.lote_imp, "10", "125.5", None)])
    self.l3, = salida("SAL-2", [(self.lote_sel, "20", "100", "2000")])
    self.sin_precio, = salida("SAL-3", [(self.lote_imp, "5", None, None)])
    salida("SAL-0", [(self.lote_imp, "999", "1", "999")], fecha=CORTE)  # ya en el saldo inicial
    salida("SAL-9", [(self.lote_imp, "1", "1", "1")], cliente=Cliente.objects.create(nombre="OTRO"))


class VentasTests(BaseCartera):
    """Las salidas de Inventario a cliente son las ventas; Cartera las verifica línea por línea."""

    def setUp(self):
        super().setUp()
        _preparar_ventas(self)

    def _ventas(self, **params):
        return self.client.get(f"/api/cartera/clientes/{self.cliente.id}/ventas/", params).data

    def _verificar(self, lineas, estado, comentario=""):
        return self.client.post("/api/cartera/ventas/verificar/", {
            "lineas": [l.id for l in lineas], "estado": estado, "comentario": comentario,
        }, format="json")

    def test_las_salidas_posteriores_al_corte_suman_al_saldo_por_empresa(self):
        self.assertEqual(self._saldo(self.imp), Decimal("1000") + Decimal("13000") + Decimal("1255"))
        self.assertEqual(self._saldo(self.sel), Decimal("-200") + Decimal("2000"))
        ventas = self._ventas()
        self.assertEqual([(v["folio_salida"], v["nota_salida"], v["empresa"]["nombre"], v["total"]) for v in ventas],
                         [("SAL-1", "N-SAL-1", "IMPORTADORA", "14255.00"), ("SAL-2", "N-SAL-2", "MARISCOS SELECTOS", "2000.00"),
                          ("SAL-3", "N-SAL-3", "IMPORTADORA", "0")])
        self.assertEqual([l["estado"] for l in ventas[0]["lineas"]], ["pendiente", "pendiente"])
        self.assertEqual(ventas[2]["lineas"][0]["estado"], "sin_precio")

    def test_verificar_y_reportar_diferencia(self):
        self.assertEqual(self._verificar([self.l1, self.l2], "verificada").status_code, 200)
        r = self._verificar([self.l3], "diferencia")
        self.assertEqual(r.status_code, 400)  # el comentario es obligatorio
        self._verificar([self.l3], "diferencia", "En la nota son 18 kg")
        self.assertEqual(self._verificar([self.sin_precio], "verificada").status_code, 400)

        estados = {l["id"]: (l["estado"], l["verificado_por"]) for v in self._ventas() for l in v["lineas"]}
        self.assertEqual(estados[self.l1.id], ("verificada", "cobranza"))
        self.assertEqual(estados[self.l3.id][0], "diferencia")
        self.assertEqual([v["folio_salida"] for v in self._ventas(estado="diferencia")], ["SAL-2"])
        resumen = self.client.get("/api/cartera/resumen/").data["ventas"]
        self.assertEqual(resumen, {"pendientes": 1, "diferencia": 1, "sin_precio": 1, "verificadas": 2})  # la de OTRO sigue pendiente

    def test_si_inventario_corrige_la_linea_vuelve_a_pendiente_y_se_va_el_aviso(self):
        self._verificar([self.l3], "diferencia", "En la nota son 18 kg")
        inventario = crear_usuario_con_area("almacen", "INV")
        self.client.force_authenticate(user=inventario)
        salida = self.client.get(f"/api/inventario/salidas/{self.l3.salida_id}/").data
        self.assertEqual(salida["detalles"][0]["aviso_cartera"]["comentario"], "En la nota son 18 kg")

        self.SalidaDetalle.objects.filter(pk=self.l3.pk).update(total_kilos=Decimal("18"), total_venta=Decimal("1800"))
        salida = self.client.get(f"/api/inventario/salidas/{self.l3.salida_id}/").data
        self.assertIsNone(salida["detalles"][0]["aviso_cartera"])

        self.client.force_authenticate(user=self.usuario)
        linea = [l for v in self._ventas() for l in v["lineas"] if l["id"] == self.l3.id][0]
        self.assertEqual((linea["estado"], linea["corregida_por_inventario"]), ("pendiente", True))
        self.assertEqual(self._saldo(self.sel), Decimal("-200") + Decimal("1800"))

    def test_diferencia_avisa_a_inventario_y_marca_su_menu(self):
        from notificaciones.models import Notificacion
        inventario = crear_usuario_con_area("almacen", "INV")
        self._verificar([self.l3], "diferencia", "En la nota son 18 kg")
        aviso = Notificacion.objects.get(evento="venta_diferencia")
        self.assertEqual((aviso.area_origen.codigo, aviso.area_destino.codigo), ("CAR", "INV"))
        self.assertEqual(aviso.enlace, "/inventario/salidas?busqueda=SAL-2")
        self.assertIn("En la nota son 18 kg", aviso.mensaje)

        self.client.force_authenticate(user=inventario)
        resumen = self.client.get("/api/notificaciones/resumen/").data
        self.assertEqual(resumen["no_leidas"], 1)
        self.assertEqual(resumen["puntos"], {"/inventario/salidas": 1})

        # Cartera tiene su punto por las ventas pendientes de verificar (l1, l2 y la de OTRO).
        self.client.force_authenticate(user=self.usuario)
        self.assertEqual(self.client.get("/api/notificaciones/resumen/").data["puntos"], {"/cartera/clientes": 3})

    def test_inventario_corrige_una_salida_revisada_y_avisa_a_cartera(self):
        from notificaciones.models import Notificacion
        self._verificar([self.l3], "diferencia", "En la nota son 18 kg")
        inventario = crear_usuario_con_area("almacen", "INV")
        self.client.force_authenticate(user=inventario)
        salida = self.client.get(f"/api/inventario/salidas/{self.l3.salida_id}/").data
        linea = salida["detalles"][0]
        r = self.client.put(f"/api/inventario/salidas/{self.l3.salida_id}/", {
            "folio_de_salida": salida["folio_de_salida"], "cliente_id": self.cliente.id, "fecha": salida["fecha"],
            "notas": salida["notas"],
            "detalles": [{"id": linea["id"], "producto_id": self.producto.id, "entrada_detalle": self.lote_sel.id,
                          "cajas": 1, "total_kilos": "18", "precio_x_kilo": "100", "total_venta": "1800"}],
        }, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        aviso = Notificacion.objects.get(evento="salida_corregida")
        self.assertEqual((aviso.area_origen.codigo, aviso.area_destino.codigo), ("INV", "CAR"))
        self.assertIn("1 con diferencia reportada", aviso.mensaje)
        self.assertEqual(aviso.enlace, f"/cartera/clientes/{self.cliente.id}")
        self.assertEqual(self.client.get("/api/notificaciones/resumen/").data["puntos"], {})

    def test_editar_sin_tocar_lo_revisado_no_avisa(self):
        from notificaciones.models import Notificacion
        self._verificar([self.l3], "verificada")
        inventario = crear_usuario_con_area("almacen", "INV")
        self.client.force_authenticate(user=inventario)
        salida = self.client.get(f"/api/inventario/salidas/{self.l3.salida_id}/").data
        linea = salida["detalles"][0]
        self.client.put(f"/api/inventario/salidas/{self.l3.salida_id}/", {
            "folio_de_salida": salida["folio_de_salida"], "cliente_id": self.cliente.id, "fecha": salida["fecha"],
            "notas": "solo cambia la nota",
            "detalles": [{"id": linea["id"], "producto_id": self.producto.id, "entrada_detalle": self.lote_sel.id,
                          "cajas": 1, "total_kilos": "20", "precio_x_kilo": "100", "total_venta": "2000"}],
        }, format="json")
        self.assertFalse(Notificacion.objects.filter(evento="salida_corregida").exists())

    def test_no_se_verifican_lineas_ajenas(self):
        otra = self.SalidaDetalle.objects.get(salida__folio_de_salida="SAL-0")
        self.assertEqual(self._verificar([otra], "verificada").status_code, 400)

    def test_estado_de_cuenta_trae_cada_salida_con_sus_lineas(self):
        self._verificar([self.l1], "verificada")
        renglones = self.client.get(f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/", {"empresa": self.imp.id}).data
        ventas = [r for r in renglones if r["tipo"] == "venta"]
        self.assertEqual([(v["referencia"], v["cargo"]) for v in ventas],
                         [("Salida SAL-1 · Nota N-SAL-1", "14255.00"), ("Salida SAL-3 · Nota N-SAL-3", "0.00")])
        self.assertIn("sin precio", ventas[1]["concepto"])
        self.assertEqual([(d["importe"], d["estado"]) for d in ventas[0]["detalle"]],
                         [("13000.00", "verificada"), ("1255.00", "pendiente")])
        for formato in ("pdf", "excel"):
            self.assertEqual(self.client.get(f"/api/cartera/clientes/{self.cliente.id}/estado-cuenta/{formato}/").status_code, 200)


class PagosPorIdentificarTests(BaseCartera):
    """Pagos que llegan sin saber de quién son: no mueven saldos hasta identificarlos."""

    def _por_identificar(self, monto="5000", empresa=None, fecha=None):
        return self._pago(monto, empresa=empresa, fecha=fecha, cliente_id=None)

    def test_sin_cliente_no_mueve_saldos_y_se_lista_aparte(self):
        r = self._por_identificar()
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data["cliente"], r.data["por_identificar"]), (None, True))
        self.assertEqual(self._saldo(self.imp), Decimal("1000"))
        self._pago("300")
        listar = lambda v: [p["monto"] for p in self.client.get("/api/cartera/pagos/", {"por_identificar": v}).data]
        self.assertEqual((listar("1"), listar("0")), (["5000.00"], ["300.00"]))
        resumen = self.client.get("/api/cartera/resumen/").data["por_identificar"]
        self.assertEqual(resumen, [{"empresa_id": self.imp.id, "cantidad": 1, "monto": "5000.00"}])

    def test_identificar_aplica_el_pago_con_su_fecha_original(self):
        ayer = HOY - timedelta(days=1)
        pago = self._por_identificar(fecha=ayer).data
        url = f"/api/cartera/pagos/{pago['id']}/identificar/"
        r = self.client.post(url, {"cliente_id": self.cliente.id}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data["cliente"]["nombre"], r.data["fecha"], r.data["identificado_por"]["username"]),
                         ("CABEZA DE TORO", ayer.isoformat(), "cobranza"))
        self.assertEqual(self._saldo(self.imp), Decimal("-4000"))
        self.assertEqual(self.client.post(url, {"cliente_id": self.cliente.id}, format="json").status_code, 400)

    def test_reporte_de_saldos_trae_el_renglon_por_identificar(self):
        self._por_identificar("5000")
        datos = self.client.get("/api/cartera/reportes/saldos/").data
        self.assertEqual(datos["por_identificar"][str(self.imp.id)], "-5000.00")
        hoja = load_workbook(BytesIO(self.client.get("/api/cartera/reportes/saldos/excel/").content)).active
        filas = [[c.value for c in fila][:4] for fila in hoja.iter_rows(min_row=2)]
        self.assertIn(["POR IDENTIFICAR", -5000, 0, -5000], filas)
        self.assertEqual(self.client.get("/api/cartera/reportes/saldos/pdf/").status_code, 200)


class UtilidadRealTests(BaseCartera):
    """Utilidad real: costo por venta capturado por Cartera y gasto operativo por kilo."""

    URL = "/api/cartera/utilidad/"

    def setUp(self):
        super().setUp()
        _preparar_ventas(self)

    def _utilidad(self, **params):
        return self.client.get(self.URL, {"desde": (CORTE + timedelta(days=1)).isoformat(), "hasta": HOY.isoformat(), **params}).data

    def test_filas_y_resumen_con_el_costo_del_lote_y_4_por_kilo(self):
        datos = self._utilidad()
        fila = next(f for f in datos["filas"] if f["linea_id"] == self.l1.id)
        self.assertEqual((fila["lote"], fila["nota"], fila["proveedor"], fila["costo_unitario"], fila["costo_capturado"]),
                         ("IMP1", "N-SAL-1", "REMAINS/SAN LORENZO", "100.0000", False))
        r = datos["resumen"]
        # Completas: l1 13000−10000, l2 1255−1000, l3 2000−2000 y la de OTRO 1−100; sin precio aparte.
        self.assertEqual((r["venta"], r["utilidad"], r["sin_precio"], r["costo_propuesto"]), ("16256.00", "3156.00", 1, 4))
        self.assertEqual((r["kilos"], r["gasto_operativo"], r["utilidad_real"]), ("131.000", "524.00", "2632.00"))
        self.assertEqual(datos["gasto_por_kilo"], "4.0000")

    def test_cartera_captura_el_costo_por_venta_con_historial(self):
        respuesta = self.client.post(f"{self.URL}costos/", {"costos": [{"linea_id": self.l1.id, "costo_unitario": "110"}]}, format="json")
        self.assertEqual(respuesta.data["cambios"], 1)
        r = self._utilidad()["resumen"]
        self.assertEqual((r["utilidad"], r["utilidad_real"], r["costo_propuesto"]), ("2156.00", "1632.00", 3))
        # Sin cambio no ensucia la bitácora; null vuelve al propuesto del lote.
        self.client.post(f"{self.URL}costos/", {"costos": [{"linea_id": self.l1.id, "costo_unitario": "110"}]}, format="json")
        self.client.post(f"{self.URL}costos/", {"costos": [{"linea_id": self.l1.id, "costo_unitario": None}]}, format="json")
        historial = self.client.get(f"{self.URL}costos/{self.l1.id}/historial/").data
        self.assertEqual([(h["anterior"], h["nuevo"]) for h in historial], [("110.0000", None), (None, "110.0000")])
        self.assertEqual(self._utilidad()["resumen"]["utilidad"], "3156.00")

    def test_un_gasto_nuevo_no_cambia_los_dias_anteriores(self):
        self.client.post(f"{self.URL}gastos-operativos/", {"por_kilo": "5", "vigente_desde": (HOY + timedelta(days=1)).isoformat()}, format="json")
        self.assertEqual(self._utilidad()["resumen"]["gasto_operativo"], "524.00")
        self.client.post(f"{self.URL}gastos-operativos/", {"por_kilo": "5", "vigente_desde": HOY.isoformat()}, format="json")
        self.assertEqual(self._utilidad()["resumen"]["gasto_operativo"], "655.00")

    def test_excel_como_la_hoja_registro_y_permisos(self):
        respuesta = self.client.get(f"{self.URL}excel/", {"desde": (CORTE + timedelta(days=1)).isoformat()})
        hoja = load_workbook(BytesIO(respuesta.content))["Registro"]
        self.assertEqual([c.value for c in hoja[1]][:4], ["Fecha", "LOTE", "Nota", "PROVEEDOR"])
        self.assertTrue(any(f[5] == "Utilidad real" for f in hoja.iter_rows(values_only=True)))
        self.assertEqual(self.client.post(f"{self.URL}costos/", {"costos": [{"linea_id": 999999, "costo_unitario": "1"}]},
                                          format="json").status_code, 400)
        from security.permissions import AREA_INVENTARIO
        self.client.force_authenticate(user=crear_usuario_con_area("almacen", AREA_INVENTARIO))
        self.assertEqual(self.client.get(self.URL).status_code, 403)


class ReportePagosTests(BaseCartera):
    def test_pdf_y_excel_con_los_filtros_de_la_pantalla(self):
        self._pago("300")
        cancelado = self._pago("50").data
        self.client.post(f"/api/cartera/pagos/{cancelado['id']}/cancelar/", {"motivo": "Duplicado"}, format="json")
        params = {"desde": (HOY - timedelta(days=1)).isoformat(), "hasta": HOY.isoformat(), "por_identificar": "false"}

        pdf = self.client.get("/api/cartera/pagos/pdf/", params)
        self.assertEqual((pdf.status_code, pdf["Content-Type"]), (200, "application/pdf"))
        hoja = load_workbook(BytesIO(self.client.get("/api/cartera/pagos/excel/", params).content))["Pagos de clientes"]
        filas = list(hoja.iter_rows(values_only=True))
        self.assertEqual(filas[0][:5], ("Fecha", "Cliente", "Empresa", "Tipo", "Monto"))
        self.assertEqual(sum(1 for f in filas if f[1] == "CABEZA DE TORO"), 2)
        self.assertEqual(next(f for f in filas if f[1] == "Total (activos)")[4], 300)  # el cancelado no suma
        self.assertIn("Pagos identificados", [f[0] for f in filas])
        # Solo activos: el cancelado ya no sale.
        hoja = load_workbook(BytesIO(self.client.get("/api/cartera/pagos/excel/", {**params, "estado": "activo"}).content)).active
        self.assertEqual(sum(1 for f in hoja.iter_rows(values_only=True) if f[1] == "CABEZA DE TORO"), 1)
