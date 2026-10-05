"""
Reporte de utilidad: lo que se ganó con las ventas reales de un periodo.

Cada línea de salida a un cliente (los traslados entre cámaras también son
salidas, pero sin cliente, y no son ventas) aporta:

    venta    = total_venta de la línea (o kilos × precio/kg si solo trae precio)
    costo    = kilos vendidos × costo/kg del lote del que salió
    utilidad = venta − costo           margen = utilidad ÷ venta

Una línea sin precio o cuyo lote no tiene costo no entra en los totales (se
vería como utilidad o pérdida falsa): se cuenta aparte para avisarlo.
"""
from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

from .models import SalidaDetalle
from .reportes import (
    COLOR_ENCABEZADO_RL,
    ESTILOS,
    _celda,
    _documento_base,
    _encabezado_documento,
    _estilizar_encabezado,
    _pie_pagina_pdf,
)

CERO = Decimal("0")
AGRUPACIONES = {
    "producto": "Producto",
    "cliente": "Cliente",
    "proveedor": "Proveedor",
    "salida": "Salida",
}
ESTADO_COMPLETA = "completa"
ESTADO_SIN_PRECIO = "sin_precio"
ESTADO_SIN_COSTO = "sin_costo"


def _margen(utilidad, venta):
    return (utilidad / venta * 100).quantize(Decimal("0.1")) if venta else None


def lineas_vendidas(filtros):
    """Líneas de salida a clientes con fecha entre desde y hasta (inclusive) y los filtros opcionales."""
    lineas = (
        SalidaDetalle.objects
        .filter(salida__cliente__isnull=False, salida__fecha__gte=filtros["desde"], salida__fecha__lte=filtros["hasta"])
        .select_related(
            "salida__cliente", "producto", "camara",
            "entrada_detalle__proveedor_origen", "entrada_detalle__entrada__empresa",
        )
        .order_by("salida__fecha", "salida__folio_de_salida", "id")
    )
    if filtros.get("empresa"):
        lineas = lineas.filter(entrada_detalle__entrada__empresa_id=filtros["empresa"])
    if filtros.get("cliente"):
        lineas = lineas.filter(salida__cliente_id=filtros["cliente"])
    if filtros.get("proveedor"):
        lineas = lineas.filter(entrada_detalle__proveedor_origen_id=filtros["proveedor"])
    if filtros.get("producto"):
        lineas = lineas.filter(producto_id=filtros["producto"])
    if filtros.get("camara"):
        lineas = lineas.filter(camara_id=filtros["camara"])
    return lineas


def _linea(d):
    lote = d.entrada_detalle
    kilos = d.total_kilos
    if d.total_venta is not None:
        venta = d.total_venta
    elif d.precio_x_kilo is not None:
        venta = (kilos * d.precio_x_kilo).quantize(Decimal("0.01"))
    else:
        venta = None
    costo_kg = lote.costo_por_kilo if lote else None
    costo = (kilos * costo_kg).quantize(Decimal("0.01")) if costo_kg is not None else None
    if venta is None:
        estado = ESTADO_SIN_PRECIO
    elif costo is None:
        estado = ESTADO_SIN_COSTO
    else:
        estado = ESTADO_COMPLETA
    utilidad = venta - costo if estado == ESTADO_COMPLETA else None
    return {
        "salida_id": d.salida_id,
        "fecha": d.salida.fecha,
        "folio": d.salida.folio_de_salida,
        "cliente_id": d.salida.cliente_id,
        "cliente": d.salida.cliente.nombre,
        "producto_id": d.producto_id,
        "producto": f"{d.producto.talla} {d.producto.tipo}",
        "lote": lote.lote_proveedor if lote else "—",
        "proveedor_id": lote.proveedor_origen_id if lote else None,
        "proveedor": lote.proveedor_origen.nombre if lote and lote.proveedor_origen_id else "—",
        "camara": d.camara.nombre if d.camara_id else "—",
        "kilos": kilos,
        "precio_kg": (venta / kilos).quantize(Decimal("0.01")) if venta is not None and kilos else None,
        "venta": venta,
        "costo_kg": costo_kg,
        "costo": costo,
        "utilidad": utilidad,
        "margen": _margen(utilidad, venta) if utilidad is not None else None,
        "estado": estado,
    }


def _clave_grupo(linea, agrupar):
    if agrupar == "cliente":
        return linea["cliente_id"], linea["cliente"]
    if agrupar == "proveedor":
        return linea["proveedor_id"], linea["proveedor"]
    if agrupar == "salida":
        return linea["salida_id"], f"{linea['folio']} · {linea['cliente']}"
    return linea["producto_id"], linea["producto"]


def _acumular(destino, linea):
    destino["lineas"] += 1
    if linea["estado"] != ESTADO_COMPLETA:
        destino["incompletas"] += 1
        return
    destino["kilos"] += linea["kilos"]
    destino["venta"] += linea["venta"]
    destino["costo"] += linea["costo"]
    destino["utilidad"] += linea["utilidad"]


def _vacio():
    return {"lineas": 0, "incompletas": 0, "kilos": CERO, "venta": CERO, "costo": CERO, "utilidad": CERO}


def construir_reporte(filtros):
    agrupar = filtros.get("agrupar") if filtros.get("agrupar") in AGRUPACIONES else "producto"
    lineas = [_linea(d) for d in lineas_vendidas(filtros)]

    resumen = _vacio()
    avisos = {
        ESTADO_SIN_PRECIO: {"lineas": 0, "kilos": CERO},
        ESTADO_SIN_COSTO: {"lineas": 0, "kilos": CERO},
    }
    grupos = defaultdict(lambda: {**_vacio(), "detalle": []})
    etiquetas = {}
    for linea in lineas:
        _acumular(resumen, linea)
        if linea["estado"] != ESTADO_COMPLETA:
            avisos[linea["estado"]]["lineas"] += 1
            avisos[linea["estado"]]["kilos"] += linea["kilos"]
        clave, etiqueta = _clave_grupo(linea, agrupar)
        etiquetas[clave] = etiqueta
        _acumular(grupos[clave], linea)
        grupos[clave]["detalle"].append(linea)

    resultado = []
    for clave, grupo in grupos.items():
        grupo["etiqueta"] = etiquetas[clave]
        grupo["margen"] = _margen(grupo["utilidad"], grupo["venta"])
        resultado.append(grupo)
    # Lo que más deja arriba; los grupos sin ventas completas al final.
    resultado.sort(key=lambda g: (g["lineas"] == g["incompletas"], -g["utilidad"], g["etiqueta"]))
    resumen["margen"] = _margen(resumen["utilidad"], resumen["venta"])
    return {"agrupar": agrupar, "resumen": resumen, "avisos": avisos, "grupos": resultado}


# ---------------------------------------------------------------- JSON
def _texto(valor):
    if valor is None:
        return None
    if isinstance(valor, Decimal):
        return str(valor)
    return valor


def serializar_reporte(reporte):
    def numeros(d):
        return {k: _texto(v) for k, v in d.items() if k != "detalle"}

    return {
        "agrupar": reporte["agrupar"],
        "resumen": numeros(reporte["resumen"]),
        "avisos": {k: numeros(v) for k, v in reporte["avisos"].items()},
        "grupos": [
            {
                **numeros(g),
                "detalle": [
                    {k: (v.isoformat() if k == "fecha" else _texto(v)) for k, v in linea.items()}
                    for linea in g["detalle"]
                ],
            }
            for g in reporte["grupos"]
        ],
    }


# ---------------------------------------------------------------- Excel
FORMATO_PESOS = '"$"#,##0.00'
FORMATO_KG = "#,##0.00"
ESTADOS_TEXTO = {ESTADO_COMPLETA: "", ESTADO_SIN_PRECIO: "Sin precio", ESTADO_SIN_COSTO: "Sin costo"}


def _numero(valor):
    return float(valor) if valor is not None else None


def construir_excel(reporte, subtitulo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    agrupado = AGRUPACIONES[reporte["agrupar"]]
    encabezados = [agrupado, "Líneas", "Kg vendidos", "Venta", "Costo", "Utilidad", "Margen %", "Líneas incompletas"]
    ws.append(encabezados)
    _estilizar_encabezado(ws, len(encabezados))
    for g in reporte["grupos"]:
        ws.append([g["etiqueta"], g["lineas"], _numero(g["kilos"]), _numero(g["venta"]), _numero(g["costo"]),
                   _numero(g["utilidad"]), _numero(g["margen"]), g["incompletas"] or None])
    r = reporte["resumen"]
    if reporte["grupos"]:
        ws.append(["Total", r["lineas"], _numero(r["kilos"]), _numero(r["venta"]), _numero(r["costo"]),
                   _numero(r["utilidad"]), _numero(r["margen"]), r["incompletas"] or None])
        for celda in ws[ws.max_row]:
            celda.font = Font(bold=True)
    for (celda,) in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        celda.number_format = FORMATO_KG
    for col in (4, 5, 6):
        for (celda,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format = FORMATO_PESOS
    ws.append([])
    ws.append([f"Filtros: {subtitulo}"])
    ws.append(["Utilidad = venta − kg vendidos × costo/kg del lote. Las líneas sin precio o sin costo no suman."])
    for col, ancho in zip("ABCDEFGH", (30, 9, 14, 15, 15, 15, 11, 18)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "B2"

    det = wb.create_sheet("Detalle")
    encabezados = ["Fecha", "Folio", "Cliente", "Producto", "Lote", "Proveedor", "Cámara", "Kg", "Precio/kg",
                   "Venta", "Costo/kg", "Costo", "Utilidad", "Margen %", "Observación"]
    det.append(encabezados)
    _estilizar_encabezado(det, len(encabezados))
    for g in reporte["grupos"]:
        for l in g["detalle"]:
            det.append([l["fecha"], l["folio"], l["cliente"], l["producto"], l["lote"], l["proveedor"], l["camara"],
                        _numero(l["kilos"]), _numero(l["precio_kg"]), _numero(l["venta"]), _numero(l["costo_kg"]),
                        _numero(l["costo"]), _numero(l["utilidad"]), _numero(l["margen"]), ESTADOS_TEXTO[l["estado"]]])
    for (celda,) in det.iter_rows(min_row=2, min_col=1, max_col=1):
        celda.number_format = "DD/MM/YYYY"
    for (celda,) in det.iter_rows(min_row=2, min_col=8, max_col=8):
        celda.number_format = FORMATO_KG
    for col in (9, 10, 11, 12, 13):
        for (celda,) in det.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format = FORMATO_PESOS
    for i, ancho in enumerate((12, 14, 24, 18, 14, 22, 18, 11, 11, 14, 11, 14, 14, 10, 13), start=1):
        det.column_dimensions[det.cell(row=1, column=i).column_letter].width = ancho
    det.freeze_panes = "C2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------- PDF
def construir_pdf(reporte, subtitulo):
    buffer, documento = _documento_base("Reporte de utilidad")
    elementos = _encabezado_documento("Utilidad", [subtitulo] if subtitulo else [])

    r = reporte["resumen"]

    def pesos(v):
        return f"${v:,.2f}" if v is not None else "—"

    def margen(v):
        return f"{v}%" if v is not None else "—"

    estilo_resumen = ParagraphStyle("resumen-utilidad", parent=ESTILOS["Normal"], fontSize=10, leading=13)
    elementos.append(Paragraph(
        f"<b>Venta:</b> {pesos(r['venta'])} &nbsp;&nbsp; <b>Costo:</b> {pesos(r['costo'])} &nbsp;&nbsp; "
        f"<b>Utilidad:</b> {pesos(r['utilidad'])} &nbsp;&nbsp; <b>Margen:</b> {margen(r['margen'])} &nbsp;&nbsp; "
        f"<b>Kg vendidos:</b> {r['kilos']:,.2f}",
        estilo_resumen,
    ))
    incompletas = r["incompletas"]
    if incompletas:
        a = reporte["avisos"]
        elementos.append(Paragraph(
            f"No incluidas: {a[ESTADO_SIN_PRECIO]['lineas']} línea(s) sin precio y "
            f"{a[ESTADO_SIN_COSTO]['lineas']} sin costo.",
            ParagraphStyle("aviso-utilidad", parent=estilo_resumen, fontSize=8.5, textColor=colors.HexColor("#B45309")),
        ))
    elementos.append(Spacer(1, 0.4 * cm))

    if not reporte["grupos"]:
        elementos.append(Paragraph("No hay ventas en el periodo con los filtros seleccionados.", ESTILOS["Normal"]))
    else:
        estilo_enc = ParagraphStyle("enc-utilidad", parent=ESTILOS["Normal"], fontSize=8, textColor=colors.white,
                                    fontName="Helvetica-Bold", alignment=TA_CENTER)
        encabezados = [AGRUPACIONES[reporte["agrupar"]], "Líneas", "Kg vendidos", "Venta", "Costo", "Utilidad", "Margen"]
        filas = [[Paragraph(h, estilo_enc) for h in encabezados]]
        for g in reporte["grupos"]:
            filas.append([
                _celda(g["etiqueta"]), _celda(f"{g['lineas']:,}", TA_RIGHT), _celda(f"{g['kilos']:,.2f}", TA_RIGHT),
                _celda(pesos(g["venta"]), TA_RIGHT), _celda(pesos(g["costo"]), TA_RIGHT),
                _celda(pesos(g["utilidad"]), TA_RIGHT), _celda(margen(g["margen"]), TA_RIGHT),
            ])
        negrita = ParagraphStyle("total-utilidad", parent=ESTILOS["Normal"], fontSize=7.5, fontName="Helvetica-Bold",
                                 alignment=TA_RIGHT)
        filas.append([
            Paragraph("Total", ParagraphStyle("total-izq", parent=negrita, alignment=0)),
            Paragraph(f"{r['lineas']:,}", negrita), Paragraph(f"{r['kilos']:,.2f}", negrita),
            Paragraph(pesos(r["venta"]), negrita), Paragraph(pesos(r["costo"]), negrita),
            Paragraph(pesos(r["utilidad"]), negrita), Paragraph(margen(r["margen"]), negrita),
        ])
        tabla = Table(filas, colWidths=[c * cm for c in (8.5, 1.8, 3, 3.2, 3.2, 3.2, 2.2)], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F7F7F7")]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E9ECEF")),
        ]))
        elementos.append(tabla)

    documento.build(elementos, onFirstPage=_pie_pagina_pdf, onLaterPages=_pie_pagina_pdf)
    buffer.seek(0)
    return buffer
