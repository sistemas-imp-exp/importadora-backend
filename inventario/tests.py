import importlib
import io
from datetime import timedelta
from decimal import Decimal

from django.db import IntegrityError
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APITestCase

from security.permissions import AREA_INVENTARIO
from security.testing import crear_usuario_con_area

from .alertas import clasificar_nivel, obtener_lotes_por_vencer
from .models import Camara, Cliente, EdicionEntrada, Empresa, Proveedor, Producto, Entrada, EntradaDetalle, LoteGeneral, Salida, SalidaDetalle, MovimientoCamara
from .testing import empresa_importadora


class CatalogosModelTests(TestCase):
    def test_camara_propia_tiene_empresa_y_camara_de_tercero_no(self):
        empresa = Empresa.objects.create(nombre="Importadora y Exportadora de Mariscos")
        propia = Camara.objects.create(
            nombre="IMPORTADORA", ubicacion="Arriaga", tipo="propia", empresa=empresa
        )
        tercero = Camara.objects.create(
            nombre="REMAINS3", ubicacion="Tapachula", tipo="tercero"
        )

        self.assertEqual(propia.empresa, empresa)
        self.assertIsNone(tercero.empresa)
        self.assertEqual(str(propia), "IMPORTADORA")

    def test_proveedor_y_cliente_solo_requieren_nombre(self):
        proveedor = Proveedor.objects.create(nombre="CACESA")
        cliente = Cliente.objects.create(nombre="HERAY ACERO MORENO")

        self.assertEqual(str(proveedor), "CACESA")
        self.assertEqual(str(cliente), "HERAY ACERO MORENO")


class ProductoModelTests(TestCase):
    def test_talla_y_tipo_no_se_repiten(self):
        Producto.objects.create(talla="41-50", tipo="FREEZADO", categoria="Camarón")

        with self.assertRaises(IntegrityError):
            Producto.objects.create(talla="41-50", tipo="FREEZADO")

    def test_presentacion_es_opcional_y_no_se_deriva_de_la_talla(self):
        producto = Producto.objects.create(
            talla="16-20", tipo="MARQUETA", presentacion=Producto.PRESENTACION_COLAS
        )

        self.assertEqual(producto.presentacion, Producto.PRESENTACION_COLAS)
        self.assertEqual(str(producto), "16-20 MARQUETA")


class EntradaModelTests(TestCase):
    def setUp(self):
        self.proveedor = Proveedor.objects.create(nombre="CACESA")
        self.camara = Camara.objects.create(nombre="IMPORTADORA", tipo=Camara.TIPO_PROPIA)
        self.producto = Producto.objects.create(talla="30-40", tipo="FREEZADO")

    def test_entrada_detalle_a_resguardo_requiere_lote_proveedor_y_puede_agrupar_en_lote_general(self):
        entrada = Entrada.objects.create(
            fecha="2026-01-05", proveedor=self.proveedor, factura="FACT 1070"
        )
        lote_general = LoteGeneral.objects.create(
            codigo="IMP1765", entrada=entrada, camara=self.camara, fecha_recibo="2026-01-05"
        )
        detalle = EntradaDetalle.objects.create(
            entrada=entrada,
            producto=self.producto,
            lote_general=lote_general,
            lote_proveedor="LOTE-001",
            camara=self.camara,
            cajas=1,
            peso_por_caja=Decimal("18.00"),
            total_kilos=Decimal("18.00"),
            costo_por_kilo=Decimal("100.00"),
            precio_venta_planeado=Decimal("150.00"),
        )

        self.assertEqual(detalle.lote_general.codigo, "IMP1765")
        self.assertEqual(detalle.entrada.proveedor, self.proveedor)

    def test_entrada_detalle_de_venta_directa_no_requiere_lote_general_ni_camara(self):
        entrada = Entrada.objects.create(fecha="2026-01-02", proveedor=self.proveedor)
        detalle = EntradaDetalle.objects.create(
            entrada=entrada,
            producto=self.producto,
            lote_proveedor="LOTE-DIRECTO",
            cajas=25,
            peso_por_caja=Decimal("20.00"),
            total_kilos=Decimal("500.00"),
        )

        self.assertIsNone(detalle.lote_general)
        self.assertIsNone(detalle.camara)
        self.assertIsNone(detalle.costo_por_kilo)

    def test_lote_proveedor_vacio_es_rechazado_por_la_base_de_datos(self):
        entrada = Entrada.objects.create(fecha="2026-01-02", proveedor=self.proveedor)

        with self.assertRaises(IntegrityError):
            EntradaDetalle.objects.create(
                entrada=entrada,
                producto=self.producto,
                lote_proveedor="",
                cajas=10,
                peso_por_caja=Decimal("10.00"),
                total_kilos=Decimal("100.00"),
            )


class CorregirPesoPorCajaMigracionTests(TestCase):
    """La 0018 rellena los peso_por_caja que la 0017 dejó en 0."""

    def test_deriva_peso_de_total_kilos_entre_cajas(self):
        from django.apps import apps
        migracion = importlib.import_module("inventario.migrations.0018_corregir_peso_por_caja")

        proveedor = Proveedor.objects.create(nombre="CACESA")
        producto = Producto.objects.create(talla="30-40", tipo="FREEZADO")
        entrada = Entrada.objects.create(fecha="2026-01-05", proveedor=proveedor)
        con_cajas = EntradaDetalle.objects.create(
            entrada=entrada, producto=producto, lote_proveedor="LOTE-1",
            cajas=3, peso_por_caja=Decimal("1.00"), total_kilos=Decimal("100.00"),
        )
        sin_cajas = EntradaDetalle.objects.create(
            entrada=entrada, producto=producto, lote_proveedor="LOTE-2",
            cajas=0, peso_por_caja=Decimal("1.00"), total_kilos=Decimal("7.50"),
        )
        # update() salta el validador, igual que el default=0 de la 0017.
        EntradaDetalle.objects.update(peso_por_caja=Decimal("0"))

        migracion.corregir_peso_por_caja(apps, None)

        con_cajas.refresh_from_db()
        sin_cajas.refresh_from_db()
        self.assertEqual(con_cajas.peso_por_caja, Decimal("33.33"))
        self.assertEqual(sin_cajas.peso_por_caja, Decimal("7.50"))


class SalidaModelTests(TestCase):
    def setUp(self):
        self.proveedor = Proveedor.objects.create(nombre="CACESA")
        self.cliente = Cliente.objects.create(nombre="HERAY ACERO MORENO")
        self.camara = Camara.objects.create(nombre="IMPORTADORA", tipo=Camara.TIPO_PROPIA)
        self.producto = Producto.objects.create(talla="41-50", tipo="FREEZADO")

        self.entrada = Entrada.objects.create(fecha="2026-01-05", proveedor=self.proveedor)
        self.lote_a = EntradaDetalle.objects.create(
            entrada=self.entrada,
            producto=self.producto,
            lote_proveedor="LOTE-A",
            camara=self.camara,
            cajas=60,
            peso_por_caja=Decimal("18.00"),
            total_kilos=Decimal("1080.00"),
        )
        self.lote_b = EntradaDetalle.objects.create(
            entrada=self.entrada,
            producto=self.producto,
            lote_proveedor="LOTE-B",
            camara=self.camara,
            cajas=40,
            peso_por_caja=Decimal("18.00"),
            total_kilos=Decimal("720.00"),
        )

    def test_venta_que_se_completa_con_dos_lotes_se_parte_en_dos_lineas(self):
        salida = Salida.objects.create(
            folio_de_salida="SI9000", cliente=self.cliente, fecha="2026-02-01"
        )
        SalidaDetalle.objects.create(
            salida=salida,
            producto=self.producto,
            entrada_detalle=self.lote_a,
            camara=self.camara,
            cajas=60,
            total_kilos=Decimal("1080.00"),
            factura_proveedor="FACT 1070",
            precio_x_kilo=Decimal("140.00"),
        )
        SalidaDetalle.objects.create(
            salida=salida,
            producto=self.producto,
            entrada_detalle=self.lote_b,
            camara=self.camara,
            cajas=40,
            total_kilos=Decimal("720.00"),
            factura_proveedor="FACT 1090",
            precio_x_kilo=Decimal("145.00"),
        )

        self.assertEqual(salida.detalles.count(), 2)
        facturas = set(salida.detalles.values_list('factura_proveedor', flat=True))
        self.assertEqual(facturas, {"FACT 1070", "FACT 1090"})

    def test_venta_parcial_descuenta_cajas_y_kilos_disponibles_del_lote(self):
        # lote_a entró con 60 cajas / 1080.00 kilos (18kg/caja)
        salida = Salida.objects.create(
            folio_de_salida="SI9100", cliente=self.cliente, fecha="2026-02-02"
        )
        SalidaDetalle.objects.create(
            salida=salida,
            producto=self.producto,
            entrada_detalle=self.lote_a,
            camara=self.camara,
            cajas=20,
            total_kilos=Decimal("360.00"),
            precio_x_kilo=Decimal("140.00"),
        )

        self.lote_a.refresh_from_db()
        self.assertEqual(self.lote_a.cajas_disponibles, 40)
        self.assertEqual(self.lote_a.kilos_disponibles, Decimal("720.00"))
        # el lote que no tuvo venta no se ve afectado
        self.assertEqual(self.lote_b.cajas_disponibles, 40)
        self.assertEqual(self.lote_b.kilos_disponibles, Decimal("720.00"))


class MovimientoCamaraModelTests(TestCase):
    def test_movimiento_enlaza_salida_origen_con_entrada_destino_sin_datos_comerciales(self):
        proveedor = Proveedor.objects.create(nombre="CACESA")
        cliente = Cliente.objects.create(nombre="IMPORTADORA-MEXIDELI")
        camara_origen = Camara.objects.create(nombre="IMPORTADORA", tipo=Camara.TIPO_PROPIA)
        camara_destino = Camara.objects.create(nombre="MEXIDELI", tipo=Camara.TIPO_PROPIA)
        producto = Producto.objects.create(talla="31-35", tipo="FREEZADO")

        entrada_origen = Entrada.objects.create(fecha="2026-01-20", proveedor=proveedor)
        lote_origen = EntradaDetalle.objects.create(
            entrada=entrada_origen,
            producto=producto,
            lote_proveedor="LOTE-ORIGEN",
            camara=camara_origen,
            cajas=500,
            peso_por_caja=Decimal("20.00"),
            total_kilos=Decimal("10000.00"),
        )

        salida = Salida.objects.create(
            folio_de_salida="SI8085", cliente=cliente, fecha="2026-01-26"
        )
        salida_detalle = SalidaDetalle.objects.create(
            salida=salida,
            producto=producto,
            entrada_detalle=lote_origen,
            camara=camara_origen,
            cajas=500,
            total_kilos=Decimal("10000.00"),
        )

        entrada_destino = Entrada.objects.create(fecha="2026-01-26", proveedor=proveedor)
        lote_destino = EntradaDetalle.objects.create(
            entrada=entrada_destino,
            producto=producto,
            lote_proveedor="LOTE-ORIGEN",
            camara=camara_destino,
            cajas=500,
            peso_por_caja=Decimal("20.00"),
            total_kilos=Decimal("10000.00"),
        )

        movimiento = MovimientoCamara.objects.create(
            entrada_detalle_origen=lote_origen,
            salida_detalle=salida_detalle,
            entrada_detalle_destino=lote_destino,
            camara_origen=camara_origen,
            camara_destino=camara_destino,
            fecha="2026-01-26",
            cajas=500,
        )

        self.assertEqual(salida_detalle.movimiento_camara, movimiento)
        self.assertEqual(lote_destino.movimiento_camara_como_destino, movimiento)

    def test_una_salida_puede_mover_dos_lotes_distintos_a_la_vez(self):
        proveedor = Proveedor.objects.create(nombre="CACESA")
        cliente = Cliente.objects.create(nombre="IMPORTADORA-MEXIDELI")
        camara_origen = Camara.objects.create(nombre="IMPORTADORA2", tipo=Camara.TIPO_PROPIA)
        camara_destino = Camara.objects.create(nombre="MEXIDELI2", tipo=Camara.TIPO_PROPIA)
        producto = Producto.objects.create(talla="61-70", tipo="FREEZADO")

        entrada = Entrada.objects.create(fecha="2026-01-20", proveedor=proveedor)
        lote_1 = EntradaDetalle.objects.create(
            entrada=entrada, producto=producto, lote_proveedor="LOTE-1",
            camara=camara_origen, cajas=100, peso_por_caja=Decimal("10.00"),
            total_kilos=Decimal("1000.00"),
        )
        lote_2 = EntradaDetalle.objects.create(
            entrada=entrada, producto=producto, lote_proveedor="LOTE-2",
            camara=camara_origen, cajas=50, peso_por_caja=Decimal("10.00"),
            total_kilos=Decimal("500.00"),
        )

        salida = Salida.objects.create(folio_de_salida="SI9999", cliente=cliente, fecha="2026-01-27")
        salida_detalle_1 = SalidaDetalle.objects.create(
            salida=salida, producto=producto, entrada_detalle=lote_1,
            camara=camara_origen, cajas=100, total_kilos=Decimal("1000.00"),
        )
        salida_detalle_2 = SalidaDetalle.objects.create(
            salida=salida, producto=producto, entrada_detalle=lote_2,
            camara=camara_origen, cajas=50, total_kilos=Decimal("500.00"),
        )

        entrada_destino = Entrada.objects.create(fecha="2026-01-27", proveedor=proveedor)
        lote_1_destino = EntradaDetalle.objects.create(
            entrada=entrada_destino, producto=producto, lote_proveedor="LOTE-1",
            camara=camara_destino, cajas=100, peso_por_caja=Decimal("10.00"),
            total_kilos=Decimal("1000.00"),
        )
        lote_2_destino = EntradaDetalle.objects.create(
            entrada=entrada_destino, producto=producto, lote_proveedor="LOTE-2",
            camara=camara_destino, cajas=50, peso_por_caja=Decimal("10.00"),
            total_kilos=Decimal("500.00"),
        )

        movimiento_1 = MovimientoCamara.objects.create(
            entrada_detalle_origen=lote_1, salida_detalle=salida_detalle_1,
            entrada_detalle_destino=lote_1_destino, camara_origen=camara_origen,
            camara_destino=camara_destino, fecha="2026-01-27", cajas=100,
        )
        movimiento_2 = MovimientoCamara.objects.create(
            entrada_detalle_origen=lote_2, salida_detalle=salida_detalle_2,
            entrada_detalle_destino=lote_2_destino, camara_origen=camara_origen,
            camara_destino=camara_destino, fecha="2026-01-27", cajas=50,
        )

        self.assertEqual(salida.detalles.count(), 2)
        self.assertNotEqual(movimiento_1, movimiento_2)
        self.assertEqual(salida_detalle_1.movimiento_camara, movimiento_1)
        self.assertEqual(salida_detalle_2.movimiento_camara, movimiento_2)


class CamaraApiTests(APITestCase):
    def setUp(self):
        user = crear_usuario_con_area("tester", AREA_INVENTARIO)
        self.client.force_authenticate(user=user)

    def test_crear_y_listar_camaras(self):
        response = self.client.post(
            "/api/inventario/camaras/",
            {"nombre": "IMPORTADORA", "tipo": "propia"},
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)

        response = self.client.get("/api/inventario/camaras/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)


class EntradaApiTests(APITestCase):
    def setUp(self):
        user = crear_usuario_con_area("tester2", AREA_INVENTARIO)
        self.client.force_authenticate(user=user)
        self.proveedor = Proveedor.objects.create(nombre="CACESA")
        self.camara = Camara.objects.create(nombre="IMPORTADORA2", tipo=Camara.TIPO_PROPIA)
        self.producto = Producto.objects.create(talla="70-80", tipo="FREEZADO")

    def test_crear_entrada(self):
        response = self.client.post(
            "/api/inventario/entradas/",
            {
                "fecha": "2026-01-05",
                "empresa_id": empresa_importadora().id,
                "proveedor_id": self.proveedor.id,
                "factura": "FACT 1070",
                "detalles": [
                    {
                        "producto_id": self.producto.id,
                        "lote_proveedor": "LOTE-TEST",
                        "camara": self.camara.id,
                        "cajas": 10,
                        "peso_por_caja": "18.00",
                        "total_kilos": "180.00",
                    }
                ],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)


class EntradaCreacionAnidadaApiTests(APITestCase):
    def setUp(self):
        self.user = crear_usuario_con_area("tester3", AREA_INVENTARIO)
        self.client.force_authenticate(user=self.user)
        self.proveedor = Proveedor.objects.create(nombre="CACESA")
        self.camara = Camara.objects.create(nombre="IMPORTADORA3", tipo=Camara.TIPO_PROPIA)
        self.producto = Producto.objects.create(talla="80-100", tipo="FREEZADO")

    def test_crear_entrada_con_dos_lineas_en_un_solo_post(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "FACT 2001",
            "pedimento": "",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-A",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                },
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-B",
                    "camara": self.camara.id,
                    "cajas": 5,
                    "peso_por_caja": "18.00",
                    "total_kilos": "90.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        entrada = Entrada.objects.get(id=response.data["id"])
        self.assertEqual(entrada.detalles.count(), 2)

    def test_crear_entrada_sin_lineas_falla(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "",
            "pedimento": "",
            "detalles": [],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 400)

    def test_entrada_internacional_sin_pedimento_falla(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "es_internacional": True,
            "factura": "FACT 3001",
            "pedimento": "",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-INTL",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("pedimento", response.data)

    def test_entrada_nacional_sin_pedimento_se_crea(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "es_internacional": False,
            "factura": "FACT 3002",
            "pedimento": "",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-NAL",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)

    def test_entrada_sin_factura_falla(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-SF",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("factura", response.data)

    def test_textos_libres_se_normalizan_a_mayusculas(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "fact 3003",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "lote-minusculas",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                    "observaciones": "revisar con calma",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["factura"], "FACT 3003")
        self.assertEqual(response.data["detalles"][0]["lote_proveedor"], "LOTE-MINUSCULAS")
        self.assertEqual(response.data["detalles"][0]["observaciones"], "REVISAR CON CALMA")

    def test_las_lineas_heredan_el_proveedor_de_la_entrada_como_proveedor_origen(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "FACT 3005",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-D",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["detalles"][0]["proveedor_origen"]["id"], self.proveedor.id)

    def test_recibo_ingreso_registra_al_usuario_logueado_en_el_lote_general(self):
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "FACT 3004",
            "recibo_ingreso": "imp-9999",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-C",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "18.00",
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        lote_general = LoteGeneral.objects.get(codigo="IMP-9999")
        self.assertEqual(lote_general.creado_por_id, self.user.id)

    def test_peso_por_caja_en_cero_se_rechaza(self):
        # cajas_disponibles divide entre peso_por_caja: un 0 tumbaría existencias.
        payload = {
            "fecha": "2026-02-01",
            "empresa_id": empresa_importadora().id,
            "proveedor_id": self.proveedor.id,
            "factura": "FACT 3006",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "lote_proveedor": "LOTE-CERO",
                    "camara": self.camara.id,
                    "cajas": 10,
                    "peso_por_caja": "0.00",
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/entradas/", payload, format="json")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("detalles", response.data)
        self.assertFalse(Entrada.objects.filter(factura="FACT 3006").exists())


class SalidaCreacionAnidadaApiTests(APITestCase):
    def setUp(self):
        user = crear_usuario_con_area("tester4", AREA_INVENTARIO)
        self.client.force_authenticate(user=user)
        self.proveedor = Proveedor.objects.create(nombre="CACESA")
        self.cliente = Cliente.objects.create(nombre="HERAY ACERO MORENO")
        self.camara = Camara.objects.create(nombre="IMPORTADORA4", tipo=Camara.TIPO_PROPIA)
        self.producto = Producto.objects.create(talla="90-110", tipo="FREEZADO")

        entrada = Entrada.objects.create(fecha="2026-02-01", proveedor=self.proveedor)
        self.lote_a = EntradaDetalle.objects.create(
            entrada=entrada, producto=self.producto, lote_proveedor="LOTE-A",
            camara=self.camara, cajas=60, peso_por_caja="18.00", total_kilos="1080.00",
        )
        self.lote_b = EntradaDetalle.objects.create(
            entrada=entrada, producto=self.producto, lote_proveedor="LOTE-B",
            camara=self.camara, cajas=40, peso_por_caja="18.00", total_kilos="720.00",
        )

    def test_crear_salida_con_lineas_de_lotes_distintos_en_un_solo_post(self):
        payload = {
            "folio_de_salida": "SI9500",
            "cliente_id": self.cliente.id,
            "fecha": "2026-02-05",
            "notas": "2999/25A",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "entrada_detalle": self.lote_a.id,
                    "camara": self.camara.id,
                    "cajas": 60,
                    "total_kilos": "1080.00",
                    "factura_proveedor": "FACT 1070",
                    "precio_x_kilo": "140.00",
                },
                {
                    "producto_id": self.producto.id,
                    "entrada_detalle": self.lote_b.id,
                    "camara": self.camara.id,
                    "cajas": 40,
                    "total_kilos": "720.00",
                    "factura_proveedor": "FACT 1090",
                    "precio_x_kilo": "145.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/salidas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        salida = Salida.objects.get(id=response.data["id"])
        self.assertEqual(salida.detalles.count(), 2)

    def test_camara_de_la_linea_se_hereda_del_lote_e_ignora_lo_que_mande_el_cliente(self):
        # Una venta no puede "mover" mercancía a otra cámara — eso es exclusivo de
        # MovimientoCamara — así que aunque el payload mande una cámara distinta a
        # la real del lote, el backend debe ignorarla y usar la del lote de origen.
        otra_camara = Camara.objects.create(nombre="OTRA_CAMARA", tipo=Camara.TIPO_PROPIA)
        payload = {
            "folio_de_salida": "SI9600",
            "cliente_id": self.cliente.id,
            "fecha": "2026-02-06",
            "notas": "NS-9600",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "entrada_detalle": self.lote_a.id,
                    "camara": otra_camara.id,
                    "cajas": 10,
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/salidas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["detalles"][0]["camara"], self.camara.id)
        salida_detalle = SalidaDetalle.objects.get(id=response.data["detalles"][0]["id"])
        self.assertEqual(salida_detalle.camara_id, self.camara.id)

    def test_salida_sin_nota_de_salida_es_rechazada(self):
        # La nota de salida es el folio del documento físico con el que sale la
        # mercancía: sin ella la salida del sistema no se amarra con nada.
        payload = {
            "folio_de_salida": "SI9601",
            "cliente_id": self.cliente.id,
            "fecha": "2026-02-06",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "entrada_detalle": self.lote_a.id,
                    "cajas": 10,
                    "total_kilos": "180.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/salidas/", payload, format="json")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("notas", response.data)
        self.assertFalse(Salida.objects.filter(folio_de_salida="SI9601").exists())

    def test_salida_con_cajas_en_cero_registra_kilos_sueltos(self):
        # lote_a entró con 60 cajas / 1080.00 kilos (18kg/caja). Vender kilos
        # sueltos de una caja ya abierta no debe mover ninguna caja completa.
        payload = {
            "folio_de_salida": "SI9700",
            "cliente_id": self.cliente.id,
            "fecha": "2026-02-07",
            "notas": "NS-9700",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "entrada_detalle": self.lote_a.id,
                    "cajas": 0,
                    "total_kilos": "5.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/salidas/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.lote_a.refresh_from_db()
        self.assertEqual(self.lote_a.kilos_disponibles, Decimal("1075.00"))
        # 1075.00 / 18.00 sigue dando piso 59: no se perdió ninguna caja completa.
        self.assertEqual(self.lote_a.cajas_disponibles, 59)

    def test_salida_que_excede_los_kilos_disponibles_del_lote_se_rechaza(self):
        # lote_b entró con 40 cajas / 720.00 kilos: pedir más de eso debe
        # rechazarse aunque las cajas pedidas sí alcancen (aquí ni se piden).
        payload = {
            "folio_de_salida": "SI9701",
            "cliente_id": self.cliente.id,
            "fecha": "2026-02-07",
            "notas": "NS-9701",
            "detalles": [
                {
                    "producto_id": self.producto.id,
                    "entrada_detalle": self.lote_b.id,
                    "cajas": 0,
                    "total_kilos": "800.00",
                },
            ],
        }
        response = self.client.post("/api/inventario/salidas/", payload, format="json")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("detalles", response.data)
        self.assertFalse(Salida.objects.filter(folio_de_salida="SI9701").exists())


class MovimientoCamaraApiTests(APITestCase):
    def setUp(self):
        user = crear_usuario_con_area("tester5", AREA_INVENTARIO)
        self.client.force_authenticate(user=user)
        self.proveedor = Proveedor.objects.create(nombre="ACUAMAYA")
        self.camara_origen = Camara.objects.create(nombre="FRIGARSA", tipo=Camara.TIPO_TERCERO)
        self.camara_destino = Camara.objects.create(nombre="MEXIDELI", tipo=Camara.TIPO_PROPIA)
        self.producto = Producto.objects.create(talla="41-50", tipo="FREEZADO")

        entrada = Entrada.objects.create(fecha="2026-02-01", proveedor=self.proveedor, factura="FACT 4001")
        self.lote_origen = EntradaDetalle.objects.create(
            entrada=entrada, producto=self.producto, lote_proveedor="LOTE-MOV",
            camara=self.camara_origen, cajas=100, peso_por_caja=Decimal("20.00"),
            total_kilos=Decimal("2000.00"),
            proveedor_origen=self.proveedor, fecha_caducidad="2026-06-01",
        )

    def test_el_lote_destino_hereda_el_proveedor_de_origen_no_del_entrada_de_llegada(self):
        # La "entrada de llegada" que arma el movimiento no tiene proveedor propio
        # (no es una compra real), pero el lote sigue siendo del mismo proveedor —
        # antes se perdía esta trazabilidad y las existencias mostraban "sin proveedor".
        payload = {
            "entrada_detalle_origen": self.lote_origen.id,
            "camara_destino": self.camara_destino.id,
            "fecha": "2026-02-10",
            "cajas": 40,
            "total_kilos": "800.00",
        }
        response = self.client.post("/api/inventario/movimientos-camara/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        lote_destino = EntradaDetalle.objects.get(id=response.data["entrada_detalle_destino"])
        self.assertEqual(lote_destino.proveedor_origen_id, self.proveedor.id)
        self.assertIsNone(lote_destino.entrada.proveedor_id)

    def test_el_lote_destino_hereda_la_fecha_de_caducidad_del_lote_de_origen(self):
        payload = {
            "entrada_detalle_origen": self.lote_origen.id,
            "camara_destino": self.camara_destino.id,
            "fecha": "2026-02-10",
            "cajas": 40,
            "total_kilos": "800.00",
        }
        response = self.client.post("/api/inventario/movimientos-camara/", payload, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        lote_destino = EntradaDetalle.objects.get(id=response.data["entrada_detalle_destino"])
        self.assertEqual(str(lote_destino.fecha_caducidad), "2026-06-01")

    def _mover(self, origen=None, destino=None, cajas=40):
        response = self.client.post("/api/inventario/movimientos-camara/", {
            "entrada_detalle_origen": (origen or self.lote_origen).id,
            "camara_destino": (destino or self.camara_destino).id,
            "fecha": "2026-02-10", "cajas": cajas, "total_kilos": f"{cajas * 20}.00",
        }, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def _recibo(self, movimiento_id, recibo):
        return self.client.patch(
            f"/api/inventario/movimientos-camara/{movimiento_id}/recibo/", {"recibo": recibo}, format="json",
        )

    def test_el_recibo_en_destino_se_guarda_sin_tocar_el_origen_y_deja_bitacora(self):
        LoteGeneral.objects.create(
            codigo="IMP-ORIG", entrada=self.lote_origen.entrada, camara=self.camara_origen, fecha_recibo="2026-02-01",
        )
        EntradaDetalle.objects.filter(id=self.lote_origen.id).update(lote_general=LoteGeneral.objects.get(codigo="IMP-ORIG"))
        movimiento = self._mover()
        self.assertEqual((movimiento["recibo_origen"], movimiento["recibo_destino"]), ("IMP-ORIG", "IMP-ORIG"))
        self.assertFalse(movimiento["recibo_destino_propio"])

        respuesta = self._recibo(movimiento["id"], "imp-dest")

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        consultado = self.client.get(f"/api/inventario/movimientos-camara/{movimiento['id']}/").data
        self.assertEqual((consultado["recibo_origen"], consultado["recibo_destino"]), ("IMP-ORIG", "IMP-DEST"))
        self.assertTrue(consultado["recibo_destino_propio"])
        self.lote_origen.refresh_from_db()
        self.assertEqual(self.lote_origen.lote_general.codigo, "IMP-ORIG")
        self.assertEqual(self.lote_origen.entrada.factura, "FACT 4001")
        existencias = {e["detalle_id"]: e for e in self.client.get("/api/inventario/existencias/").data}
        destino = existencias[movimiento["entrada_detalle_destino"]]
        self.assertEqual((destino["recibo_ingreso"], destino["factura"]), ("IMP-DEST", "FACT 4001"))
        self.assertEqual(existencias[self.lote_origen.id]["recibo_ingreso"], "IMP-ORIG")
        bitacora = EdicionEntrada.objects.get(campo="recibo_ingreso")
        self.assertEqual((bitacora.valor_anterior, bitacora.valor_nuevo), ("", "IMP-DEST"))
        self.assertEqual(bitacora.entrada_detalle_id, movimiento["entrada_detalle_destino"])

    def test_el_recibo_en_destino_se_puede_capturar_al_registrar_el_movimiento(self):
        response = self.client.post("/api/inventario/movimientos-camara/", {
            "entrada_detalle_origen": self.lote_origen.id, "camara_destino": self.camara_destino.id,
            "fecha": "2026-02-10", "cajas": 40, "total_kilos": "800.00", "recibo_destino": " imp-nuevo ",
        }, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["recibo_destino"], response.data["recibo_destino_propio"]), ("IMP-NUEVO", True))
        self.assertIsNone(EntradaDetalle.objects.get(id=self.lote_origen.id).lote_general_id)
        repetido = self.client.post("/api/inventario/movimientos-camara/", {
            "entrada_detalle_origen": self.lote_origen.id, "camara_destino": self.camara_destino.id,
            "fecha": "2026-02-11", "cajas": 1, "total_kilos": "20.00", "recibo_destino": "IMP-NUEVO",
        }, format="json")
        self.assertEqual(repetido.status_code, 400)

    def test_el_listado_trae_kilos_documentos_y_lo_que_queda_en_destino(self):
        movimiento = self._mover()
        salida = Salida.objects.create(folio_de_salida="V-1", fecha="2026-02-12")
        SalidaDetalle.objects.create(
            salida=salida, producto=self.producto, camara=self.camara_destino, cajas=5, total_kilos=Decimal("100.00"),
            entrada_detalle_id=movimiento["entrada_detalle_destino"],
        )

        fila = self.client.get("/api/inventario/movimientos-camara/").data[0]

        self.assertEqual(
            (fila["kilos"], fila["proveedor"], fila["factura"], fila["fecha_caducidad"]),
            ("800.00", "ACUAMAYA", "FACT 4001", "2026-06-01"),
        )
        self.assertEqual((fila["cajas_disponibles_destino"], fila["kilos_disponibles_destino"]), (35, "700.00"))
        self.assertEqual(fila["creado_por"]["username"], "tester5")

    def test_un_traslado_posterior_hereda_el_recibo_del_destino(self):
        primero = self._mover()
        self._recibo(primero["id"], "IMP-MEX")
        tercera = Camara.objects.create(nombre="REMAINS 3", tipo=Camara.TIPO_TERCERO)
        segundo = self._mover(origen=EntradaDetalle.objects.get(id=primero["entrada_detalle_destino"]), destino=tercera, cajas=10)
        self.assertEqual(segundo["recibo_destino"], "IMP-MEX")

    def test_vaciar_el_recibo_vuelve_al_heredado(self):
        movimiento = self._mover()
        self._recibo(movimiento["id"], "IMP-1")
        respuesta = self._recibo(movimiento["id"], "")
        self.assertEqual((respuesta.data["recibo_destino"], respuesta.data["recibo_destino_propio"]), ("", False))
        self.assertFalse(LoteGeneral.objects.filter(codigo="IMP-1").exists())

    def test_recibo_repetido_o_sin_dato_se_rechaza_y_solo_lectura_no_edita(self):
        movimiento = self._mover()
        LoteGeneral.objects.create(
            codigo="IMP-USADO", entrada=self.lote_origen.entrada, camara=self.camara_origen, fecha_recibo="2026-02-01",
        )
        self.assertEqual(self._recibo(movimiento["id"], "IMP-USADO").status_code, 400)
        self.assertEqual(
            self.client.patch(f"/api/inventario/movimientos-camara/{movimiento['id']}/recibo/", {}, format="json").status_code,
            400,
        )
        lector = crear_usuario_con_area("lector-mov", AREA_INVENTARIO, solo_lectura=True)
        self.client.force_authenticate(user=lector)
        self.assertEqual(self._recibo(movimiento["id"], "IMP-X").status_code, 403)


class ClasificarNivelTests(TestCase):
    def test_vencido_si_dias_restantes_es_negativo(self):
        self.assertEqual(clasificar_nivel(-1), "vencido")
        self.assertEqual(clasificar_nivel(-30), "vencido")

    def test_critico_hasta_7_dias(self):
        self.assertEqual(clasificar_nivel(0), "critico")
        self.assertEqual(clasificar_nivel(7), "critico")

    def test_urgente_de_8_a_15_dias(self):
        self.assertEqual(clasificar_nivel(8), "urgente")
        self.assertEqual(clasificar_nivel(15), "urgente")

    def test_por_vencer_de_16_a_30_dias(self):
        self.assertEqual(clasificar_nivel(16), "por_vencer")
        self.assertEqual(clasificar_nivel(30), "por_vencer")

    def test_ninguno_fuera_de_la_ventana_de_30_dias(self):
        self.assertIsNone(clasificar_nivel(31))
        self.assertIsNone(clasificar_nivel(365))


class AlertasCaducidadTests(TestCase):
    def setUp(self):
        self.proveedor = Proveedor.objects.create(nombre="ACUAMAYA")
        self.camara = Camara.objects.create(nombre="FRIGARSA", tipo=Camara.TIPO_TERCERO)
        self.producto = Producto.objects.create(talla="41-50", tipo="FREEZADO")
        self.hoy = timezone.localdate()

    def _crear_lote(self, cajas, dias_para_caducar=None, cajas_vendidas=0):
        entrada = Entrada.objects.create(fecha=self.hoy, proveedor=self.proveedor)
        fecha_caducidad = self.hoy + timedelta(days=dias_para_caducar) if dias_para_caducar is not None else None
        # peso_por_caja=10.00: consistente con total_kilos=100.00 solo porque todas
        # las llamadas de este helper usan cajas=10 — si eso cambia, hay que ajustar
        # total_kilos junto con cajas para no desalinear cajas_disponibles.
        lote = EntradaDetalle.objects.create(
            entrada=entrada, producto=self.producto, lote_proveedor="LOTE-X",
            camara=self.camara, cajas=cajas, peso_por_caja=Decimal("10.00"),
            total_kilos=Decimal("100.00"),
            proveedor_origen=self.proveedor, fecha_caducidad=fecha_caducidad,
        )
        if cajas_vendidas:
            cliente = Cliente.objects.create(nombre=f"CLIENTE-{lote.id}")
            salida = Salida.objects.create(folio_de_salida=f"SI-{lote.id}", cliente=cliente, fecha=self.hoy)
            SalidaDetalle.objects.create(
                salida=salida, producto=self.producto, entrada_detalle=lote,
                camara=self.camara, cajas=cajas_vendidas, total_kilos=Decimal(cajas_vendidas * 10),
            )
        return lote

    def test_excluye_lotes_sin_caducidad(self):
        self._crear_lote(cajas=10, dias_para_caducar=None)
        alertas = obtener_lotes_por_vencer({})
        self.assertEqual(len(alertas), 0)

    def test_excluye_lotes_fuera_de_la_ventana_de_30_dias(self):
        self._crear_lote(cajas=10, dias_para_caducar=45)
        alertas = obtener_lotes_por_vencer({})
        self.assertEqual(len(alertas), 0)

    def test_excluye_lotes_sin_cajas_disponibles(self):
        self._crear_lote(cajas=10, dias_para_caducar=5, cajas_vendidas=10)
        alertas = obtener_lotes_por_vencer({})
        self.assertEqual(len(alertas), 0)

    def test_incluye_y_clasifica_correctamente_por_nivel(self):
        self._crear_lote(cajas=10, dias_para_caducar=-2)
        self._crear_lote(cajas=10, dias_para_caducar=3)
        self._crear_lote(cajas=10, dias_para_caducar=10)
        self._crear_lote(cajas=10, dias_para_caducar=25)

        alertas = obtener_lotes_por_vencer({})

        self.assertEqual(len(alertas), 4)
        self.assertEqual([a['nivel'] for a in alertas], ["vencido", "critico", "urgente", "por_vencer"])

    def test_filtro_por_nivel(self):
        self._crear_lote(cajas=10, dias_para_caducar=3)
        self._crear_lote(cajas=10, dias_para_caducar=25)

        alertas = obtener_lotes_por_vencer({'nivel': 'critico'})

        self.assertEqual(len(alertas), 1)
        self.assertEqual(alertas[0]['nivel'], 'critico')


class AlertasCaducidadApiTests(APITestCase):
    def setUp(self):
        user = crear_usuario_con_area("tester6", AREA_INVENTARIO)
        self.client.force_authenticate(user=user)
        self.proveedor = Proveedor.objects.create(nombre="ACUAMAYA")
        self.camara = Camara.objects.create(nombre="FRIGARSA", tipo=Camara.TIPO_TERCERO)
        self.producto = Producto.objects.create(talla="41-50", tipo="FREEZADO")

    def test_endpoint_devuelve_totales_y_conteo_por_nivel(self):
        hoy = timezone.localdate()
        entrada = Entrada.objects.create(fecha=hoy, proveedor=self.proveedor)
        EntradaDetalle.objects.create(
            entrada=entrada, producto=self.producto, lote_proveedor="LOTE-Y",
            camara=self.camara, cajas=10, peso_por_caja=Decimal("10.00"),
            total_kilos=Decimal("100.00"),
            proveedor_origen=self.proveedor, fecha_caducidad=hoy + timedelta(days=3),
        )

        response = self.client.get("/api/inventario/alertas/caducidad/")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["total"], 1)
        self.assertEqual(response.data["conteo_por_nivel"]["critico"], 1)
        self.assertEqual(response.data["alertas"][0]["nivel"], "critico")


class EmpresaEntradaApiTests(APITestCase):
    """Cada entrada pertenece a una empresa del grupo (IMPORTADORA, MARISCOS SELECTOS)."""

    def setUp(self):
        self.client.force_authenticate(user=crear_usuario_con_area("almacen", AREA_INVENTARIO))
        self.importadora = empresa_importadora()
        self.selectos = Empresa.objects.get(nombre="MARISCOS SELECTOS")
        self.proveedor = Proveedor.objects.create(nombre="ACUAMAYA")
        self.camara = Camara.objects.create(nombre="IMPORTADORA", tipo=Camara.TIPO_PROPIA)
        self.camara_destino = Camara.objects.create(nombre="MEXIDELI", tipo=Camara.TIPO_TERCERO)
        self.producto = Producto.objects.create(talla="41-50", tipo="FREEZADO")

    def _crear(self, empresa_id, factura):
        payload = {
            "fecha": "2026-09-01",
            "proveedor_id": self.proveedor.id,
            "factura": factura,
            "detalles": [{
                "producto_id": self.producto.id, "lote_proveedor": f"L-{factura}",
                "camara": self.camara.id, "cajas": 10, "peso_por_caja": "20.00", "total_kilos": "200.00",
            }],
        }
        if empresa_id is not None:
            payload["empresa_id"] = empresa_id
        return self.client.post("/api/inventario/entradas/", payload, format="json")

    def test_la_empresa_es_obligatoria(self):
        response = self._crear(None, "F-SIN")

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("empresa_id", response.data)

    def test_la_entrada_devuelve_su_empresa_y_se_filtra_por_ella(self):
        creada = self._crear(self.selectos.id, "F-SEL")
        self._crear(self.importadora.id, "F-IMP")

        self.assertEqual(creada.status_code, 201, creada.data)
        self.assertEqual(creada.data["empresa"]["nombre"], "MARISCOS SELECTOS")
        listado = self.client.get(f"/api/inventario/entradas/?empresa={self.selectos.id}")
        self.assertEqual([e["factura"] for e in listado.data["results"]], ["F-SEL"])

    def test_existencias_y_traslados_respetan_la_empresa(self):
        self._crear(self.importadora.id, "F-IMP")
        entrada_sel = self._crear(self.selectos.id, "F-SEL").data
        lote_sel = entrada_sel["detalles"][0]["id"]
        # Mover parte del lote de Selectos a otra cámara no debe cambiarlo de dueño.
        movimiento = self.client.post("/api/inventario/movimientos-camara/", {
            "entrada_detalle_origen": lote_sel, "camara_destino": self.camara_destino.id,
            "fecha": "2026-09-02", "cajas": 4, "total_kilos": "80.00",
        }, format="json")
        self.assertEqual(movimiento.status_code, 201, movimiento.data)

        existencias = self.client.get(f"/api/inventario/existencias/?empresa={self.selectos.id}").data
        self.assertEqual(
            sorted((e["detalle_id"], e["camara_nombre"]) for e in existencias),
            sorted([(lote_sel, "IMPORTADORA"), (movimiento.data["entrada_detalle_destino"], "MEXIDELI")]),
        )
        self.assertEqual({e["empresa_nombre"] for e in existencias}, {"MARISCOS SELECTOS"})
        excel = self.client.get(f"/api/inventario/reportes/existencias/excel/?modo=lote&empresa={self.selectos.id}")
        self.assertEqual(excel.status_code, 200)

        # El Excel replica la tabla de la pantalla: mismas columnas, factura del
        # lote raíz también en el lote trasladado, y fila de totales al final.
        from openpyxl import load_workbook
        filas = list(load_workbook(io.BytesIO(excel.content)).active.iter_rows(values_only=True))
        self.assertEqual(filas[0][:6], ("Cámara", "Producto", "Proveedor", "Fecha entrada", "Recibo ingreso", "Factura"))
        self.assertEqual(len(filas[0]), 16)
        self.assertEqual(filas[0][12:14], ("Total", "Utilidad"))
        self.assertEqual({f[5] for f in filas[1:-1]}, {"F-SEL"})
        self.assertEqual(filas[-1][1], "Total (2 lotes)")
        pdf = self.client.get(f"/api/inventario/reportes/existencias/pdf/?modo=lote&empresa={self.selectos.id}")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))


class EntradaEmpresaMigracionTests(TestCase):
    def test_asigna_importadora_a_las_entradas_existentes(self):
        from django.apps import apps
        migracion = importlib.import_module("inventario.migrations.0019_entrada_empresa")
        entrada = Entrada.objects.create(fecha="2026-01-05", factura="VIEJA")

        migracion.crear_empresas_y_asignar(apps, None)

        entrada.refresh_from_db()
        self.assertEqual(entrada.empresa.nombre, "IMPORTADORA")
        self.assertEqual(Empresa.objects.filter(nombre__in=["IMPORTADORA", "MARISCOS SELECTOS"]).count(), 2)


class QuitarEmpresaDuplicadaMigracionTests(TestCase):
    def test_pasa_entradas_y_camaras_a_importadora_y_borra_la_repetida(self):
        from django.apps import apps
        migracion = importlib.import_module("inventario.migrations.0020_quitar_empresa_duplicada")
        duplicada = Empresa.objects.create(nombre="IMPORTADORA DE MARISCOS")
        entrada = Entrada.objects.create(fecha="2026-01-05", factura="X", empresa=duplicada)
        camara = Camara.objects.create(nombre="CAM-DUP", tipo=Camara.TIPO_PROPIA, empresa=duplicada)

        migracion.quitar_duplicada(apps, None)

        entrada.refresh_from_db()
        camara.refresh_from_db()
        self.assertEqual((entrada.empresa, camara.empresa), (empresa_importadora(), empresa_importadora()))
        self.assertFalse(Empresa.objects.filter(nombre="IMPORTADORA DE MARISCOS").exists())


class UtilidadExistenciasTests(TestCase):
    def test_utilidad_solo_con_precio_y_costo(self):
        from .reportes import filas_existencias_lote, obtener_lotes_filtrados
        proveedor = Proveedor.objects.create(nombre="P")
        camara = Camara.objects.create(nombre="CAM-UT", tipo=Camara.TIPO_PROPIA)
        entrada = Entrada.objects.create(fecha="2026-01-05", proveedor=proveedor, factura="F", empresa=empresa_importadora())
        for talla, costo, precio in (("21-25", "100", "130"), ("26-30", "100", None)):
            EntradaDetalle.objects.create(
                entrada=entrada, producto=Producto.objects.create(talla=talla, tipo="FREEZADO"), lote_proveedor="L",
                camara=camara, cajas=10, peso_por_caja=Decimal("20"), total_kilos=Decimal("200"),
                costo_por_kilo=Decimal(costo), precio_venta_planeado=Decimal(precio) if precio else None,
                proveedor_origen=proveedor,
            )

        filas, totales = filas_existencias_lote(obtener_lotes_filtrados({}))

        self.assertEqual(sorted(f["utilidad"] for f in filas if f["utilidad"] is not None), [Decimal("6000")])
        self.assertEqual(totales["utilidad"], Decimal("6000"))


class ExcelTablasApiTests(APITestCase):
    """Entradas y Salidas se descargan en Excel con los filtros de la pantalla."""

    def setUp(self):
        # Solo lectura también descarga: es un GET.
        self.client.force_authenticate(user=crear_usuario_con_area("lector-xls", AREA_INVENTARIO, solo_lectura=True))
        proveedor = Proveedor.objects.create(nombre="CACESA")
        cliente = Cliente.objects.create(nombre="HERAY")
        camara = Camara.objects.create(nombre="CAM-XLS", tipo=Camara.TIPO_PROPIA)
        producto = Producto.objects.create(talla="21-25", tipo="FREEZADO")
        for fecha, factura in (("2026-03-01", "F-MAR"), ("2026-04-01", "F-ABR")):
            entrada = Entrada.objects.create(fecha=fecha, proveedor=proveedor, factura=factura, empresa=empresa_importadora())
            lote = EntradaDetalle.objects.create(
                entrada=entrada, producto=producto, lote_proveedor=f"L-{factura}", camara=camara, cajas=10,
                peso_por_caja=Decimal("20.00"), total_kilos=Decimal("200.00"), proveedor_origen=proveedor,
                lote_general=LoteGeneral.objects.create(
                    codigo=f"IMP-{factura[2:]}", entrada=entrada, camara=camara, fecha_recibo=fecha,
                ),
            )
        salida = Salida.objects.create(folio_de_salida="V-XLS", cliente=cliente, fecha="2026-04-05")
        SalidaDetalle.objects.create(
            salida=salida, producto=producto, entrada_detalle=lote, camara=camara, cajas=2,
            total_kilos=Decimal("40.00"), precio_x_kilo=Decimal("100.00"), total_venta=Decimal("4000.00"),
        )

    def _hojas(self, url):
        from openpyxl import load_workbook
        respuesta = self.client.get(url)
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(io.BytesIO(respuesta.content))
        return {ws.title: list(ws.iter_rows(values_only=True)) for ws in libro}

    def test_entradas_respeta_filtros_y_trae_detalle(self):
        hojas = self._hojas("/api/inventario/entradas/excel/?desde=2026-04-01")

        self.assertEqual(list(hojas), ["Entradas", "Detalle"])
        self.assertEqual([f[3] for f in hojas["Entradas"][1:-1]], ["F-ABR"])
        self.assertEqual(hojas["Entradas"][-1][1], "Total (1 entradas)")
        self.assertEqual(hojas["Entradas"][1][12], "Sí")  # con salidas
        self.assertEqual(hojas["Detalle"][1][6], "L-F-ABR")

    def test_salidas_con_totales_y_detalle(self):
        hojas = self._hojas("/api/inventario/salidas/excel/?busqueda=heray")

        self.assertEqual(hojas["Salidas"][1][1:8], ("V-XLS", "HERAY", None, 1, 2, 40, 4000))
        self.assertEqual(hojas["Detalle"][0][7], "Recibo")
        self.assertEqual(hojas["Detalle"][1][4:9], ("L-F-ABR", "CACESA", "CAM-XLS", "IMP-ABR", "—"))
