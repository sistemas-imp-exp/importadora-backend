"""Reporte de utilidad: ventas reales del periodo menos el costo de sus lotes."""
import io
from decimal import Decimal

from rest_framework.test import APITestCase

from security.permissions import AREA_INVENTARIO
from security.testing import crear_usuario_con_area

from .models import Camara, Cliente, Entrada, EntradaDetalle, MovimientoCamara, Producto, Proveedor, Salida, SalidaDetalle
from .testing import empresa_importadora

URL = "/api/inventario/reportes/utilidad/"
PERIODO = {"desde": "2026-03-01", "hasta": "2026-03-31"}


class ReporteUtilidadApiTests(APITestCase):
    def setUp(self):
        self.client.force_authenticate(user=crear_usuario_con_area("util", AREA_INVENTARIO, solo_lectura=True))
        self.proveedor = Proveedor.objects.create(nombre="CACESA")
        self.cliente = Cliente.objects.create(nombre="HERAY")
        self.otro_cliente = Cliente.objects.create(nombre="MOSTRADOR")
        self.camara = Camara.objects.create(nombre="CAM-U", tipo=Camara.TIPO_PROPIA)
        self.p1 = Producto.objects.create(talla="21-25", tipo="FREEZADO")
        self.p2 = Producto.objects.create(talla="26-30", tipo="FREEZADO")
        entrada = Entrada.objects.create(fecha="2026-02-01", proveedor=self.proveedor, factura="F", empresa=empresa_importadora())
        self.lote1 = self._lote(entrada, self.p1, Decimal("100"))
        self.lote2 = self._lote(entrada, self.p2, None)  # sin costo

    def _lote(self, entrada, producto, costo):
        return EntradaDetalle.objects.create(
            entrada=entrada, producto=producto, lote_proveedor="L", camara=self.camara, cajas=50,
            peso_por_caja=Decimal("20"), total_kilos=Decimal("1000"), costo_por_kilo=costo,
            proveedor_origen=self.proveedor,
        )

    def _vender(self, folio, fecha, lote, kilos, total, cliente=None):
        salida = Salida.objects.create(folio_de_salida=folio, cliente=cliente or self.cliente, fecha=fecha)
        return SalidaDetalle.objects.create(
            salida=salida, producto=lote.producto, entrada_detalle=lote, camara=self.camara, cajas=1,
            total_kilos=Decimal(kilos), total_venta=Decimal(total) if total else None,
        )

    def test_utilidad_y_margen_con_avisos_de_lineas_incompletas(self):
        self._vender("V1", "2026-03-05", self.lote1, "100", "13000")  # utilidad 3000
        self._vender("V2", "2026-03-06", self.lote1, "50", "4500", self.otro_cliente)  # pérdida 500
        self._vender("V3", "2026-03-07", self.lote1, "10", None)  # sin precio
        self._vender("V4", "2026-03-08", self.lote2, "20", "3000")  # lote sin costo
        self._vender("V5", "2026-04-01", self.lote1, "10", "9999")  # fuera del periodo

        datos = self.client.get(URL, PERIODO).data

        r = datos["resumen"]
        self.assertEqual((r["venta"], r["costo"], r["utilidad"], r["margen"]), ("17500.00", "15000.00", "2500.00", "14.3"))
        self.assertEqual((r["lineas"], r["incompletas"]), (4, 2))
        self.assertEqual((datos["avisos"]["sin_precio"]["lineas"], datos["avisos"]["sin_costo"]["lineas"]), (1, 1))
        self.assertEqual([g["etiqueta"] for g in datos["grupos"]], ["21-25 FREEZADO", "26-30 FREEZADO"])

    def test_agrupa_por_cliente_y_filtra(self):
        self._vender("V1", "2026-03-05", self.lote1, "100", "13000")
        self._vender("V2", "2026-03-06", self.lote1, "50", "4500", self.otro_cliente)

        por_cliente = self.client.get(URL, {**PERIODO, "agrupar": "cliente"}).data
        self.assertEqual([(g["etiqueta"], g["utilidad"]) for g in por_cliente["grupos"]],
                         [("HERAY", "3000.00"), ("MOSTRADOR", "-500.00")])
        filtrado = self.client.get(URL, {**PERIODO, "cliente": self.otro_cliente.id}).data
        self.assertEqual(filtrado["resumen"]["utilidad"], "-500.00")

    def test_los_traslados_entre_camaras_no_son_ventas(self):
        destino = Camara.objects.create(nombre="CAM-D", tipo=Camara.TIPO_PROPIA)
        self.client.force_authenticate(user=crear_usuario_con_area("mov-u", AREA_INVENTARIO))
        respuesta = self.client.post("/api/inventario/movimientos-camara/", {
            "entrada_detalle_origen": self.lote1.id, "camara_destino": destino.id,
            "fecha": "2026-03-10", "cajas": 5, "total_kilos": "100.00",
        }, format="json")
        self.assertEqual(respuesta.status_code, 201, respuesta.data)
        movido = MovimientoCamara.objects.get().entrada_detalle_destino
        self._vender("V1", "2026-03-11", movido, "50", "6000")  # el lote movido conserva el costo

        datos = self.client.get(URL, PERIODO).data
        self.assertEqual((datos["resumen"]["lineas"], datos["resumen"]["utilidad"]), (1, "1000.00"))

    def test_fechas_obligatorias_y_rango_valido(self):
        self.assertEqual(self.client.get(URL).status_code, 400)
        self.assertEqual(self.client.get(URL, {"desde": "2026-03-31", "hasta": "2026-03-01"}).status_code, 400)

    def test_excel_y_pdf(self):
        from openpyxl import load_workbook
        self._vender("V1", "2026-03-05", self.lote1, "100", "13000")

        excel = self.client.get(URL + "excel/", PERIODO)
        libro = load_workbook(io.BytesIO(excel.content))
        self.assertEqual(libro.sheetnames, ["Resumen", "Detalle"])
        self.assertEqual(libro["Resumen"]["F2"].value, 3000)
        pdf = self.client.get(URL + "pdf/", PERIODO)
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
