"""
Carga los resguardos pendientes AL CORTE de Cartera (30/09) desde el control
del área (`docs/CARTERA/CONTROL DE RESGUARDO A CLIENTES 2026.xlsx`) como
ventas en resguardo. Lo posterior al corte (ventas y retiros de octubre en
adelante) lo captura Inventario en el sistema.

- Pendiente al corte = pendiente de la hoja "RESUMEN RESG." + lo entregado
  después del corte según la hoja de cada cliente (filas SI/SM bajo el bloque de
  la nota y producto). Las líneas que quedan en 0 se omiten.
- Las ventas con fecha posterior al corte no se cargan (las captura Inventario).
- Una venta por cliente + nota; cajas = kilos / peso por caja del último lote
  del producto (hacia abajo); el proveedor del control queda en las notas.
- Al ser anteriores al corte, Cartera no las vuelve a cobrar (ya están en el
  saldo inicial).

Por defecto solo muestra lo que haría; con --aplicar lo guarda. No repite una
venta ya cargada (mismo cliente, nota y fecha). El historial de resguardos
terminados se queda en el Excel.
"""
import datetime
import re
from collections import defaultdict
from decimal import ROUND_DOWN, Decimal

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from cartera.saldos import fecha_corte
from inventario.models import Cliente, Empresa, EntradaDetalle, Producto, Salida, SalidaDetalle

HOJA_RESUMEN = 'RESUMEN RESG.'
# Encabezado del bloque en el resumen → (cliente del catálogo, hoja con sus entregas).
CLIENTES = {
    'RAYMUNDO PEÑA': ('RAYMUNDO PEÑA', 'RAYMUNDO PEÑA'),
    'HERAY ACERO MORENO': ('HERAY ACERO MORENO', 'HERAY ACERO 1'),
    'DE JOSE LUIS UYOA VAZQUEZ': ('JOSE LUIS UYOA VAZQUEZ (CHILIO)', 'JOSE LUIS ULLOA'),
}
# Cómo escribe el control los productos que no son "talla tipo".
PRODUCTOS = {
    'FILETE TILAPIA DELIPEZ': ('FILETE DELIPEZ', 'FILETE'),
    'FILETE TILAPIA DE PAZ': ('FILETE DE PAZ', 'FILETE'),
    'FILETE DE PAZ': ('FILETE DE PAZ', 'FILETE'),
    'TENTACULO DE CALAMAR': ('TENTACULO', 'CALAMAR'),
}
KG = Decimal('0.001')
# Salidas de cámara de IMPORTADORA (SI) y de MARISCOS SELECTOS (SM).
SALIDA = re.compile(r'S[IM]\s?\d+')
NOTA = re.compile(r'\d{3,4}/\d{2}[A-Z]')


def _texto(valor):
    return ' '.join(str(valor).split()).upper() if valor is not None else ''


def _nota(valor):
    return _texto(valor).removeprefix('NOTA ')


def _producto(nombre):
    if not nombre or ' ' not in nombre:
        return None
    talla, tipo = PRODUCTOS.get(nombre) or nombre.rsplit(' ', 1)
    return Producto.objects.filter(talla=talla, tipo=tipo).first()


def _kilos(valor):
    return Decimal(str(valor)).quantize(KG)


def leer_resumen(libro):
    """[(nombre del bloque, nota, fecha, producto, proveedor, pendiente, precio)] de todas las líneas."""
    filas, bloque = [], None
    for fila in libro[HOJA_RESUMEN].iter_rows(values_only=True):
        _, fecha, nota, talla, proveedor, kilos, precio = (tuple(fila) + (None,) * 7)[:7]
        if isinstance(fecha, str) and _texto(fecha).startswith('RESGUARDO'):
            bloque = _texto(fecha)[len('RESGUARDO '):]
            if bloque not in CLIENTES:
                raise CommandError(f'Cliente del resumen sin equivalencia: "{bloque}".')
            continue
        if not isinstance(fecha, datetime.datetime) or kilos is None:
            continue
        producto = _producto(_texto(talla))
        if producto is None:
            raise CommandError(f'No existe el producto "{talla}" (nota {nota}) en el catálogo.')
        filas.append((bloque, _nota(nota), fecha.date(), producto, _texto(proveedor),
                      max(_kilos(kilos), Decimal('0')), Decimal(str(precio))))
    return filas


def entregas_despues(hoja, notas, corte):
    """
    {(nota, producto_id): kilos} entregados después del corte: filas SI/SM con
    fecha > corte bajo el encabezado del bloque de esa nota y producto.
    """
    entregado, bloque = defaultdict(Decimal), None
    for fila in hoja.iter_rows(values_only=True):
        celdas = [_texto(v) for v in fila[:5] if isinstance(v, str)]
        if any(c.startswith('RESGUARDO TERMINADO') for c in celdas):
            bloque = None
            continue
        if any(SALIDA.fullmatch(c) for c in celdas):
            fecha = next((v for v in fila[:3] if isinstance(v, datetime.datetime)), None)
            # Los kilos entregados: primera cifra después de las columnas de texto.
            kilos = next((v for v in fila[5:11] if isinstance(v, (int, float))), None)
            if bloque and fecha and fecha.date() > corte and kilos:
                entregado[bloque] += _kilos(kilos)
            continue
        # Encabezado de otra venta: cierra el bloque anterior aunque no sea de las cargadas.
        nota = next((_nota(c) for c in celdas if NOTA.search(c)), None)
        if nota:
            producto = next((p for p in map(_producto, celdas) if p), None)
            bloque = (nota, producto.id) if nota in notas and producto else None
    return entregado


def _cajas(producto, kilos):
    lote = EntradaDetalle.objects.filter(producto=producto, peso_por_caja__gt=0).order_by('-id').first()
    return int((kilos / lote.peso_por_caja).to_integral_value(ROUND_DOWN)) if lote else 0


class Command(BaseCommand):
    help = 'Carga los resguardos pendientes al corte de Cartera desde el control del área (modo prueba salvo --aplicar).'

    def add_arguments(self, parser):
        parser.add_argument('ruta', help='Ruta del .xlsx de control de resguardos.')
        parser.add_argument('--aplicar', action='store_true', help='Guarda los cambios; sin esto solo los muestra.')

    def handle(self, *args, **opciones):
        corte = fecha_corte()
        if corte is None:
            raise CommandError('Cartera no tiene fecha de corte (saldos iniciales).')
        libro = load_workbook(opciones['ruta'], data_only=True)
        resumen = leer_resumen(libro)
        notas = {fila[1] for fila in resumen}
        despues = {}
        for bloque, (_, hoja) in CLIENTES.items():
            despues[bloque] = entregas_despues(libro[hoja], notas, corte)

        ventas = defaultdict(list)
        for bloque, nota, fecha, producto, proveedor, pendiente, precio in resguardo_filas(resumen, corte, self):
            entregado = despues[bloque].get((nota, producto.id), Decimal('0'))
            kilos = pendiente + entregado
            if kilos > 0:
                ventas[(bloque, nota, fecha)].append((producto, proveedor, kilos, precio, pendiente, entregado))

        empresa = Empresa.objects.get(nombre='IMPORTADORA')
        with transaction.atomic():
            totales = self._cargar(ventas, empresa, corte)
            if not opciones['aplicar']:
                transaction.set_rollback(True)
        self.stdout.write(
            f"\n{totales['ventas']} venta(s), {totales['lineas']} línea(s), {totales['kilos']:,.3f} kg pendientes "
            f"al {corte:%d/%m/%Y}; {totales['omitidas']} venta(s) ya cargadas."
        )
        self.stdout.write(self.style.SUCCESS('Guardado.') if opciones['aplicar'] else 'Modo prueba: no se guardó nada.')

    def _cargar(self, ventas, empresa, corte):
        totales = {'ventas': 0, 'lineas': 0, 'kilos': Decimal('0'), 'omitidas': 0}
        for (bloque, nota, fecha), lineas in ventas.items():
            cliente = Cliente.objects.get(nombre=CLIENTES[bloque][0])
            if Salida.objects.filter(es_resguardo=True, cliente=cliente, notas=nota, fecha=fecha).exists():
                totales['omitidas'] += 1
                self.stdout.write(f'· Ya cargada: {cliente.nombre} nota {nota} ({fecha}).')
                continue
            venta = Salida.objects.create(cliente=cliente, fecha=fecha, notas=nota, es_resguardo=True, empresa=empresa)
            self.stdout.write(f'\n{cliente.nombre} · nota {nota} · {fecha:%d/%m/%Y}')
            for producto, proveedor, kilos, precio, pendiente, entregado in lineas:
                cajas = _cajas(producto, kilos)
                SalidaDetalle.objects.create(
                    salida=venta, producto=producto, cajas=cajas, total_kilos=kilos, precio_x_kilo=precio,
                    total_venta=(kilos * precio).quantize(Decimal('0.01')), cajas_resguardo=cajas, kilos_resguardo=kilos,
                    notas=f'Proveedor según control de resguardos: {proveedor}' if proveedor else '',
                )
                totales['lineas'] += 1
                totales['kilos'] += kilos
                self.stdout.write(
                    f'   {producto}: {kilos:,.3f} kg al {corte:%d/%m} ({cajas} cajas) × ${precio}'
                    f'  = hoy {pendiente:,.3f} + entregado después {entregado:,.3f}'
                )
            totales['ventas'] += 1
        return totales


def resguardo_filas(resumen, corte, comando):
    """Las líneas del resumen de ventas hasta el corte; las posteriores las captura Inventario."""
    for fila in resumen:
        if fila[2] > corte:
            if fila[5] > 0:
                comando.stdout.write(f'· Posterior al corte, no se carga: nota {fila[1]} ({fila[2]}) {fila[3]}.')
            continue
        yield fila
