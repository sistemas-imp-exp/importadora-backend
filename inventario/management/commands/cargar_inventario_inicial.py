"""
Carga el inventario inicial (corte 30/09/2026) desde el CONSOLIDADO del área:
pestañas IMPORTADORA y SELECTOS con el formato del kardex (FECHA, TALLA, TIPO,
PESO MASTER (KG), PROVEEDOR, LOTE, ENTREGADO A/ GUARDADO EN, ENTRADA, SALIDA,
SALDO, OBSERVACIONES, CADUCIDAD, LOTE PROVEEDOR, FACTURA, COSTO/KG, NOTAS).

Cada recibo (columna LOTE, p. ej. IMP25) es una Entrada con la fecha del
kardex, de la empresa de su pestaña, con un LoteGeneral (el recibo) y una
línea por renglón con el SALDO como kilos. Las cajas salen de saldo / peso
master. Da de alta los productos (talla + tipo) y proveedores que no existan.

Por defecto es modo prueba: hace toda la carga dentro de una transacción y la
deshace, así que valida contra la base real sin guardar nada. Con --aplicar la
guarda. --borrar-movimientos elimina antes las entradas, salidas y traslados
existentes (solo para quitar capturas de prueba; se niega si hay muchas).
"""
import datetime
import re
from collections import defaultdict
from decimal import ROUND_HALF_UP, Decimal

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from inventario.models import (
    Camara,
    EdicionEntrada,
    Empresa,
    Entrada,
    EntradaDetalle,
    LoteGeneral,
    MovimientoCamara,
    Producto,
    Proveedor,
    Salida,
    SalidaDetalle,
)

# Pestaña del archivo -> nombre de la Empresa en el sistema.
EMPRESAS = {"IMPORTADORA": "IMPORTADORA", "SELECTOS": "MARISCOS SELECTOS"}
ENCABEZADOS = [
    "FECHA", "TALLA", "TIPO", "PESO MASTER (KG)", "PROVEEDOR", "LOTE", "ENTREGADO A/ GUARDADO EN",
    "ENTRADA", "SALIDA", "SALDO", "OBSERVACIONES", "CADUCIDAD", "LOTE PROVEEDOR", "FACTURA", "COSTO/KG", "NOTAS",
]
SIN_LOTE = "SIN LOTE"  # lote_proveedor no admite vacío
MAX_ENTRADAS_A_BORRAR = 20  # --borrar-movimientos es para pruebas, no para datos reales
DOS = Decimal("0.01")


class Rollback(Exception):
    """Deshace la transacción del modo prueba."""


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split()).upper()


def decimal(valor, nombre, fila):
    if not isinstance(valor, (int, float)) or isinstance(valor, bool):
        raise CommandError(f"{fila}: {nombre} no es un número ({valor!r}).")
    return Decimal(str(valor)).quantize(DOS, rounding=ROUND_HALF_UP)


def caducidad(valor):
    """(fecha o None, aviso). Acepta fechas de Excel y textos dd/mm/aaaa o dd-mm-aaaa."""
    if isinstance(valor, datetime.datetime):
        return valor.date(), ""
    if isinstance(valor, datetime.date):
        return valor, ""
    t = texto(valor)
    if t in ("", "N/A", "NA"):
        return None, ""
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", t)
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))), ""
        except ValueError:
            pass
    return None, f"Caducidad en el archivo: {t}"


def leer_archivo(ruta):
    """{(pestaña, recibo): [renglón]} con los valores ya normalizados."""
    libro = load_workbook(ruta, data_only=True)
    recibos = defaultdict(list)
    for pestana in EMPRESAS:
        if pestana not in libro.sheetnames:
            raise CommandError(f"Falta la pestaña {pestana}.")
        hoja = libro[pestana]
        encabezado = [texto(c.value) for c in hoja[2]][: len(ENCABEZADOS)]
        if encabezado != ENCABEZADOS:
            raise CommandError(f"{pestana}: los encabezados de la fila 2 no son los esperados: {encabezado}")
        for n, r in enumerate(hoja.iter_rows(min_row=3, max_col=len(ENCABEZADOS), values_only=True), start=3):
            fila = f"{pestana} fila {n}"
            v = dict(zip(ENCABEZADOS, r))
            recibo = texto(v["LOTE"])
            if not recibo:
                continue  # total, leyenda o renglones vacíos
            if not isinstance(v["FECHA"], (datetime.date, datetime.datetime)):
                raise CommandError(f"{fila}: FECHA no es una fecha ({v['FECHA']!r}).")
            fecha_cad, aviso_cad = caducidad(v["CADUCIDAD"])
            saldo = decimal(v["SALDO"], "SALDO", fila)
            peso = decimal(v["PESO MASTER (KG)"], "PESO MASTER", fila)
            if saldo <= 0:
                raise CommandError(f"{fila}: SALDO debe ser mayor que 0.")
            if peso <= 0:
                raise CommandError(f"{fila}: PESO MASTER debe ser mayor que 0.")
            costo = v["COSTO/KG"]
            recibos[(pestana, recibo)].append({
                "fila": fila,
                "fecha": v["FECHA"].date() if isinstance(v["FECHA"], datetime.datetime) else v["FECHA"],
                "talla": texto(v["TALLA"]),
                "tipo": texto(v["TIPO"]),
                "peso": peso,
                "proveedor": texto(v["PROVEEDOR"]),
                "camara": texto(v["ENTREGADO A/ GUARDADO EN"]),
                "saldo": saldo,
                "caducidad": fecha_cad,
                "lote_proveedor": texto(v["LOTE PROVEEDOR"]) or SIN_LOTE,
                "factura": texto(v["FACTURA"]),
                "costo": decimal(costo, "COSTO/KG", fila) if costo not in (None, "") else None,
                "notas": [x for x in (
                    f"Observaciones: {texto(v['OBSERVACIONES'])}" if texto(v["OBSERVACIONES"]) else "",
                    f"Notas: {texto(v['NOTAS'])}" if texto(v["NOTAS"]) else "",
                    aviso_cad,
                ) if x],
            })
            for campo in ("talla", "tipo", "proveedor", "camara"):
                if not recibos[(pestana, recibo)][-1][campo]:
                    raise CommandError(f"{fila}: falta {campo}.")
    return recibos


class Command(BaseCommand):
    help = "Carga el inventario inicial desde el CONSOLIDADO (modo prueba salvo --aplicar)."

    def add_arguments(self, parser):
        parser.add_argument("ruta", help="Ruta del CONSOLIDADO .xlsx (pestañas IMPORTADORA y SELECTOS).")
        parser.add_argument("--aplicar", action="store_true", help="Guarda la carga; sin esto solo la valida.")
        parser.add_argument("--borrar-movimientos", action="store_true",
                            help="Elimina antes las entradas, salidas y traslados existentes (capturas de prueba).")
        parser.add_argument("--usuario", help="Usuario que queda como creador (por defecto, el primer superusuario).")

    def handle(self, *args, **opciones):
        recibos = leer_archivo(opciones["ruta"])
        usuario = self._usuario(opciones.get("usuario"))
        try:
            with transaction.atomic():
                resumen = self._cargar(recibos, usuario, opciones["borrar_movimientos"])
                if not opciones["aplicar"]:
                    raise Rollback
        except Rollback:
            self._imprimir(resumen)
            self.stdout.write(self.style.NOTICE("Modo prueba: se validó todo y no se guardó nada. Usa --aplicar para cargarlo."))
            return
        self._imprimir(resumen)
        self.stdout.write(self.style.SUCCESS("Inventario inicial cargado."))

    def _usuario(self, nombre):
        User = get_user_model()
        if nombre:
            try:
                return User.objects.get(username=nombre)
            except User.DoesNotExist:
                raise CommandError(f"No existe el usuario {nombre}.")
        usuario = User.objects.filter(is_superuser=True).order_by("id").first()
        if usuario is None:
            raise CommandError("No hay superusuario; indica --usuario.")
        return usuario

    def _borrar_movimientos(self):
        borrado = {
            "entradas": Entrada.objects.count(),
            "lineas de entrada": EntradaDetalle.objects.count(),
            "salidas": Salida.objects.count(),
            "traslados": MovimientoCamara.objects.count(),
        }
        if borrado["entradas"] > MAX_ENTRADAS_A_BORRAR:
            raise CommandError(
                f"Hay {borrado['entradas']} entradas: --borrar-movimientos es solo para capturas de prueba."
            )
        # En orden, por las llaves PROTECT entre ellos.
        MovimientoCamara.objects.all().delete()
        SalidaDetalle.objects.all().delete()
        Salida.objects.all().delete()
        EdicionEntrada.objects.all().delete()
        EntradaDetalle.objects.all().delete()
        LoteGeneral.objects.all().delete()
        Entrada.objects.all().delete()
        return borrado

    def _cargar(self, recibos, usuario, borrar):
        resumen = {"borrado": self._borrar_movimientos() if borrar else None}

        existentes = set(LoteGeneral.objects.filter(codigo__in=[r for _, r in recibos]).values_list("codigo", flat=True))
        if existentes:
            raise CommandError(f"Ya existen {len(existentes)} recibos en el sistema (p. ej. {sorted(existentes)[:5]}): "
                               "¿ya se cargó? Usa --borrar-movimientos si eran pruebas.")

        empresas = {}
        for pestana, nombre in EMPRESAS.items():
            try:
                empresas[pestana] = Empresa.objects.get(nombre=nombre)
            except Empresa.DoesNotExist:
                raise CommandError(f"No existe la empresa {nombre}.")

        lineas = [l for ls in recibos.values() for l in ls]
        camaras = {c.nombre: c for c in Camara.objects.filter(nombre__in={l["camara"] for l in lineas})}
        faltan = {l["camara"] for l in lineas} - set(camaras)
        if faltan:
            raise CommandError(f"No existen las cámaras: {sorted(faltan)}.")

        proveedores = {p.nombre.upper(): p for p in Proveedor.objects.all()}
        nuevos_proveedores = sorted({l["proveedor"] for l in lineas} - set(proveedores))
        for nombre in nuevos_proveedores:
            proveedores[nombre] = Proveedor.objects.create(nombre=nombre, creado_por=usuario)

        productos = {(p.talla.upper(), p.tipo.upper()): p for p in Producto.objects.all()}
        nuevos_productos = sorted({(l["talla"], l["tipo"]) for l in lineas} - set(productos))
        for talla, tipo in nuevos_productos:
            productos[(talla, tipo)] = Producto.objects.create(talla=talla, tipo=tipo, creado_por=usuario)

        kilos = defaultdict(Decimal)
        for (pestana, recibo), renglones in sorted(recibos.items()):
            entrada = Entrada.objects.create(
                fecha=renglones[0]["fecha"],
                empresa=empresas[pestana],
                creado_por=usuario,
            )
            lote_general = LoteGeneral.objects.create(
                codigo=recibo, camara=camaras[renglones[0]["camara"]],
                fecha_recibo=entrada.fecha, creado_por=usuario,
            )
            for l in renglones:
                notas = ["Carga inicial al 30/09/2026"]
                detalle = EntradaDetalle(
                    entrada=entrada,
                    producto=productos[(l["talla"], l["tipo"])],
                    proveedor_origen=proveedores[l["proveedor"]],
                    factura=l["factura"][:50],
                    lote_general=lote_general,
                    lote_proveedor=l["lote_proveedor"][:50],
                    camara=camaras[l["camara"]],
                    cajas=int(l["saldo"] // l["peso"]),
                    peso_por_caja=l["peso"],
                    total_kilos=l["saldo"],
                    costo_por_kilo=l["costo"],
                    fecha_caducidad=l["caducidad"],
                    observaciones=". ".join(notas + l["notas"]),
                )
                detalle.full_clean()
                detalle.save()
                kilos[pestana] += l["saldo"]

        resumen.update({
            "recibos": len(recibos),
            "lineas": len(lineas),
            "kilos": dict(kilos),
            "proveedores_nuevos": nuevos_proveedores,
            "productos_nuevos": nuevos_productos,
            "sin_lote": sum(1 for l in lineas if l["lote_proveedor"] == SIN_LOTE),
            "sin_costo": sum(1 for l in lineas if l["costo"] is None),
            "avisos": [f"{l['fila']}: {n}" for l in lineas for n in l["notas"] if n.startswith("Caducidad en el archivo")],
        })
        return resumen

    def _imprimir(self, r):
        if r.get("borrado"):
            self.stdout.write("Borrado previo (capturas de prueba): " + ", ".join(f"{v} {k}" for k, v in r["borrado"].items()))
        self.stdout.write(f"Recibos (entradas): {r['recibos']} · líneas: {r['lineas']}")
        for pestana, kg in r["kilos"].items():
            self.stdout.write(f"  {pestana}: {kg:,.2f} kg")
        self.stdout.write(f"  Total: {sum(r['kilos'].values()):,.2f} kg")
        self.stdout.write(f"Proveedores nuevos: {len(r['proveedores_nuevos'])} · productos nuevos: {len(r['productos_nuevos'])}")
        for talla, tipo in r["productos_nuevos"]:
            self.stdout.write(f"  + {talla} {tipo}")
        self.stdout.write(f"Líneas sin lote proveedor (quedan '{SIN_LOTE}'): {r['sin_lote']} · sin costo: {r['sin_costo']}")
        for aviso in r["avisos"]:
            self.stdout.write(self.style.WARNING(aviso))
