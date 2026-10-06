"""
Carga el catálogo de productos desde el Excel depurado por el área
(columnas TALLA, TIPO, CATEGORIA, PRESENTACION en la primera hoja).

Por defecto solo muestra lo que haría; con --aplicar lo guarda. Crea los
productos nuevos y actualiza categoría y presentación de los existentes
(misma talla + tipo). Nunca borra ni desactiva: un producto que ya no viene
en el archivo se queda como está, porque puede tener entradas.
"""
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from inventario.models import Producto

PRESENTACIONES = {
    "COLA": Producto.PRESENTACION_COLAS,
    "COLAS": Producto.PRESENTACION_COLAS,
    "ENTERO": Producto.PRESENTACION_ENTERO,
    "": "",
}


def _texto(valor):
    return " ".join(str(valor).split()).upper() if valor is not None else ""


def leer_catalogo(ruta):
    """[(fila, talla, tipo, categoria, presentacion)] sin duplicados (gana la primera) y la lista de avisos."""
    hoja = load_workbook(ruta, data_only=True, read_only=True).worksheets[0]
    filas, vistos, avisos = [], {}, []
    largo = {c: Producto._meta.get_field(c).max_length for c in ("talla", "tipo", "categoria")}
    for numero, fila in enumerate(hoja.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
        talla, tipo, categoria, presentacion = (_texto(v) for v in (tuple(fila) + (None,) * 4)[:4])
        if not (talla or tipo or categoria or presentacion):
            continue
        if not talla or not tipo:
            avisos.append(f"Fila {numero}: falta talla o tipo, se omite.")
            continue
        if presentacion not in PRESENTACIONES:
            avisos.append(f"Fila {numero}: presentación desconocida '{presentacion}', se omite.")
            continue
        for campo, valor in (("talla", talla), ("tipo", tipo), ("categoria", categoria)):
            if len(valor) > largo[campo]:
                raise CommandError(f"Fila {numero}: {campo} excede {largo[campo]} caracteres.")
        clave = (talla, tipo)
        if clave in vistos:
            avisos.append(f"Fila {numero}: {talla} {tipo} repetido (ya en la fila {vistos[clave]}), se omite.")
            continue
        vistos[clave] = numero
        filas.append((numero, talla, tipo, categoria, PRESENTACIONES[presentacion]))
    return filas, avisos


class Command(BaseCommand):
    help = "Carga el catálogo de productos desde Excel (modo prueba salvo --aplicar)."

    def add_arguments(self, parser):
        parser.add_argument("ruta", help="Ruta del .xlsx (TALLA, TIPO, CATEGORIA, PRESENTACION).")
        parser.add_argument("--aplicar", action="store_true", help="Guarda los cambios; sin esto solo los muestra.")

    def handle(self, *args, **opciones):
        filas, avisos = leer_catalogo(opciones["ruta"])
        existentes = {(p.talla, p.tipo): p for p in Producto.objects.all()}
        nuevos, cambios = [], []
        for numero, talla, tipo, categoria, presentacion in filas:
            producto = existentes.get((talla, tipo))
            if producto is None:
                nuevos.append(Producto(talla=talla, tipo=tipo, categoria=categoria, presentacion=presentacion))
                continue
            diferencias = {
                campo: (getattr(producto, campo), valor)
                for campo, valor in (("categoria", categoria), ("presentacion", presentacion))
                if getattr(producto, campo) != valor
            }
            if diferencias:
                cambios.append((producto, diferencias))

        for aviso in avisos:
            self.stdout.write(self.style.WARNING(aviso))
        for producto, diferencias in cambios:
            detalle = ", ".join(f"{c}: '{a}' → '{n}'" for c, (a, n) in diferencias.items())
            self.stdout.write(f"Actualizar {producto.talla} {producto.tipo}: {detalle}")
        sin_cambio = len(filas) - len(nuevos) - len(cambios)
        fuera = len(set(existentes) - {(f[1], f[2]) for f in filas})
        self.stdout.write(
            f"Archivo: {len(filas)} productos · nuevos: {len(nuevos)} · a actualizar: {len(cambios)} · "
            f"sin cambio: {sin_cambio} · en la base y no en el archivo (se conservan): {fuera}"
        )

        if not opciones["aplicar"]:
            self.stdout.write(self.style.NOTICE("Modo prueba: no se guardó nada. Usa --aplicar para cargarlo."))
            return
        with transaction.atomic():
            Producto.objects.bulk_create(nuevos)
            for producto, diferencias in cambios:
                for campo, (_, valor) in diferencias.items():
                    setattr(producto, campo, valor)
                producto.save(update_fields=list(diferencias) + ["modificado"])
        self.stdout.write(self.style.SUCCESS(f"Cargado: {len(nuevos)} nuevos, {len(cambios)} actualizados."))
