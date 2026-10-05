"""
Excel de las tablas de Entradas y Salidas: lo que se ve en pantalla con los
mismos filtros (búsqueda, fechas, empresa), sin paginar.

Cada libro trae dos hojas: una fila por documento, con las columnas de la
tabla, y "Detalle" con una fila por línea (lo que en pantalla se despliega).
Las columnas siguen a EntradasTable.tsx / SalidasTable.tsx: cambiar ambos juntos.
"""
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from .reportes import _estilizar_encabezado

FORMATO_KG = "#,##0.00"
FORMATO_PESOS = '"$"#,##0.00'
FORMATO_FECHA = "DD/MM/YYYY"


def _numero(valor):
    return float(valor) if valor is not None else None


def _hoja(ws, titulo, encabezados, filas, formatos, anchos, total=None):
    """Llena una hoja: encabezado, filas, formatos por columna (1-based) y fila de total opcional."""
    ws.title = titulo
    ws.append(encabezados)
    _estilizar_encabezado(ws, len(encabezados))
    for fila in filas:
        ws.append(fila)
    if total and filas:
        ws.append(total)
        for celda in ws[ws.max_row]:
            celda.font = Font(bold=True)
    for col, formato in formatos.items():
        for (celda,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            celda.number_format = formato
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = ancho
    ws.freeze_panes = "C2"


def libro_entradas(entradas):
    """`entradas`: queryset de EntradaViewSet (detalles con consumo anotado)."""
    resumen, detalle = [], []
    tot_cajas, tot_kilos = 0, Decimal("0")
    for e in entradas:
        lineas = list(e.detalles.all())
        cajas = sum(d.cajas for d in lineas)
        kilos = sum((d.total_kilos for d in lineas), Decimal("0"))
        caducidades = [d.fecha_caducidad for d in lineas if d.fecha_caducidad]
        con_salidas = any(d.total_kilos != d.kilos_disponibles for d in lineas)
        recibos = ", ".join(sorted({d.lote_general.codigo for d in lineas if d.lote_general_id}))
        proveedor = e.proveedor.nombre if e.proveedor_id else "—"
        empresa = e.empresa.nombre if e.empresa_id else "—"
        resumen.append([
            e.fecha, proveedor, empresa, e.factura or "—", e.pedimento or "", recibos or "—",
            len(lineas), cajas, _numero(kilos), min(caducidades) if caducidades else None,
            "Importación" if e.es_internacional else "Nacional",
            "Sí" if getattr(e, "fue_editada", False) else "No",
            "Sí" if con_salidas else "No",
        ])
        tot_cajas += cajas
        tot_kilos += kilos
        for d in lineas:
            detalle.append([
                e.fecha, proveedor, empresa, e.factura or "—",
                d.lote_general.codigo if d.lote_general_id else "—",
                f"{d.producto.talla} {d.producto.tipo}", d.lote_proveedor,
                d.camara.nombre if d.camara_id else "Venta directa",
                d.cajas, _numero(d.peso_por_caja), _numero(d.total_kilos), _numero(d.kilos_disponibles),
                _numero(d.costo_por_kilo), d.fecha_caducidad, d.observaciones or "",
            ])

    wb = Workbook()
    _hoja(
        wb.active, "Entradas",
        ["Fecha", "Proveedor", "Empresa", "Factura", "Pedimento", "Recibo ingreso", "Líneas", "Cajas",
         "Total kg", "Caducidad próxima", "Tipo", "Editada", "Con salidas"],
        resumen,
        {1: FORMATO_FECHA, 8: "#,##0", 9: FORMATO_KG, 10: FORMATO_FECHA},
        [12, 24, 20, 16, 16, 18, 8, 10, 13, 16, 13, 9, 11],
        total=[None, f"Total ({len(resumen)} entradas)", None, None, None, None, None,
               tot_cajas, _numero(tot_kilos), None, None, None, None],
    )
    _hoja(
        wb.create_sheet(), "Detalle",
        ["Fecha", "Proveedor", "Empresa", "Factura", "Recibo", "Producto", "Lote proveedor", "Cámara",
         "Cajas", "Kg/caja", "Total kg", "Disponible kg", "Costo/kg", "Caducidad", "Observaciones"],
        detalle,
        {1: FORMATO_FECHA, 9: "#,##0", 10: FORMATO_KG, 11: FORMATO_KG, 12: FORMATO_KG,
         13: FORMATO_PESOS, 14: FORMATO_FECHA},
        [12, 24, 20, 16, 14, 20, 16, 20, 9, 10, 12, 13, 12, 12, 30],
    )
    return wb


def libro_salidas(salidas):
    """`salidas`: queryset de SalidaViewSet."""
    resumen, detalle = [], []
    tot_cajas, tot_kilos, tot_venta = 0, Decimal("0"), Decimal("0")
    for s in salidas:
        lineas = list(s.detalles.all())
        cajas = sum(d.cajas for d in lineas)
        kilos = sum((d.total_kilos for d in lineas), Decimal("0"))
        venta = sum((d.total_venta or Decimal("0") for d in lineas), Decimal("0"))
        sin_precio = sum(1 for d in lineas if not d.total_venta)
        cliente = s.cliente.nombre if s.cliente_id else "—"
        resumen.append([
            s.fecha, s.folio_de_salida, cliente, s.notas or "", len(lineas), cajas,
            _numero(kilos), _numero(venta) if venta else None, sin_precio or None,
        ])
        tot_cajas += cajas
        tot_kilos += kilos
        tot_venta += venta
        for d in lineas:
            origen = d.entrada_detalle
            detalle.append([
                s.fecha, s.folio_de_salida, cliente, f"{d.producto.talla} {d.producto.tipo}",
                origen.lote_proveedor if origen else "—",
                origen.proveedor_origen.nombre if origen and origen.proveedor_origen_id else "—",
                d.camara.nombre if d.camara_id else "—", d.factura_proveedor or "—",
                d.cajas, _numero(d.total_kilos), _numero(d.precio_x_kilo), _numero(d.total_venta), d.notas or "",
            ])

    wb = Workbook()
    _hoja(
        wb.active, "Salidas",
        ["Fecha", "Folio", "Cliente", "Nota de salida", "Líneas", "Cajas", "Total kg", "Total venta",
         "Líneas sin precio"],
        resumen,
        {1: FORMATO_FECHA, 6: "#,##0", 7: FORMATO_KG, 8: FORMATO_PESOS},
        [12, 16, 26, 30, 8, 10, 13, 15, 15],
        total=[None, f"Total ({len(resumen)} salidas)", None, None, None,
               tot_cajas, _numero(tot_kilos), _numero(tot_venta), None],
    )
    _hoja(
        wb.create_sheet(), "Detalle",
        ["Fecha", "Folio", "Cliente", "Producto", "Lote origen", "Proveedor", "Cámara", "Factura proveedor",
         "Cajas", "Total kg", "Precio/kg", "Total venta", "Nota"],
        detalle,
        {1: FORMATO_FECHA, 9: "#,##0", 10: FORMATO_KG, 11: FORMATO_PESOS, 12: FORMATO_PESOS},
        [12, 16, 26, 20, 16, 24, 20, 18, 9, 12, 12, 14, 30],
    )
    return wb
