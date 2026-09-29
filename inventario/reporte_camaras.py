"""
Reporte "Existencias por cámara": el formato que el área ya usa en Excel
(referencia: docs/MEXIDELI 24.09.26.pdf del workspace).

Por cámara, dos bloques —Entero y Colas— con su subtotal y un total; dentro
de cada bloque, talla → tipo → una fila por proveedor y peso por caja, con
cajas ("master"), kilos y la existencia total de la talla. Puede sacarse a
una fecha pasada: cuenta las entradas hasta ese día y resta las salidas
hasta ese día (los traslados entre cámaras ya son una salida más una
entrada, así que quedan contados en su fecha).
"""
import re
from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from django.db.models import DecimalField, F, Q, Sum, Value
from django.db.models.functions import Coalesce
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from .models import EntradaDetalle, Producto

CERO = Decimal("0")
VERDE = "1F7A1F"  # encabezados, como en el formato del área

# Orden de los bloques. Un producto sin presentación capturada no se pierde:
# cae en su propio bloque, que además avisa que falta completar el catálogo.
BLOQUES = [
    (Producto.PRESENTACION_ENTERO, "Entero"),
    (Producto.PRESENTACION_COLAS, "Colas"),
    ("", "Sin presentación"),
]
COLUMNAS = ["TALLA", "TIPO", "PROVEEDOR", "PESO MASTER", "MASTER", "KILOGRAMOS", "EXISTENCIA"]


def _clave_talla(talla):
    """30-40 antes que 100-120: ordena por los números, no alfabéticamente."""
    numeros = [int(n) for n in re.findall(r"\d+", talla)]
    return (numeros or [10**9], talla)


# ---------------------------------------------------------------- cálculo
def lotes_con_existencia(fecha, camara_id=None, empresa_id=None):
    """
    Lotes (EntradaDetalle) con kilos disponibles > 0 al cierre de `fecha`,
    solo los que están en una cámara. `kilos_a_fecha` = kilos que entraron
    menos los que salieron hasta ese día.
    """
    lotes = (
        EntradaDetalle.objects
        .filter(entrada__fecha__lte=fecha, camara__isnull=False)
        .select_related("camara", "producto", "proveedor_origen")
        .annotate(
            kilos_salidos=Coalesce(
                Sum("salidas_detalle__total_kilos", filter=Q(salidas_detalle__salida__fecha__lte=fecha)),
                Value(CERO),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            ),
        )
        .annotate(kilos_a_fecha=F("total_kilos") - F("kilos_salidos"))
        .filter(kilos_a_fecha__gt=0)
    )
    if camara_id:
        lotes = lotes.filter(camara_id=camara_id)
    if empresa_id:
        lotes = lotes.filter(entrada__empresa_id=empresa_id)
    return lotes


def construir_reporte(fecha, camara_id=None, empresa_id=None):
    """
    Estructura lista para pintar (JSON, PDF o Excel):

    [{camara, bloques: [{presentacion, etiqueta, tallas: [{talla, existencia,
      tipos: [{tipo, filas: [{proveedor, peso_por_caja, cajas, kilos}]}]}],
      subtotal}], total}]

    Proveedores sin existencia no aparecen (decisión del área, "hasta nuevo
    aviso"): solo se agregan lotes con kilos > 0.
    """
    # cámara -> presentación -> talla -> tipo -> (proveedor, peso) -> [cajas, kilos]
    arbol = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: defaultdict(
        lambda: defaultdict(lambda: [0, CERO])))))
    camaras = {}
    for lote in lotes_con_existencia(fecha, camara_id, empresa_id):
        camaras[lote.camara_id] = lote.camara.nombre
        proveedor = lote.proveedor_origen.nombre if lote.proveedor_origen_id else "—"
        celda = arbol[lote.camara_id][lote.producto.presentacion][lote.producto.talla][lote.producto.tipo][
            (proveedor, lote.peso_por_caja)
        ]
        # Cajas completas de cada lote, igual que EntradaDetalle.cajas_disponibles.
        celda[0] += int(lote.kilos_a_fecha // lote.peso_por_caja)
        celda[1] += lote.kilos_a_fecha

    reporte = []
    for cam_id, nombre in sorted(camaras.items(), key=lambda kv: kv[1]):
        bloques = []
        for presentacion, etiqueta in BLOQUES:
            por_talla = arbol[cam_id].get(presentacion)
            if not por_talla:
                continue
            tallas = []
            for talla in sorted(por_talla, key=_clave_talla):
                tipos = []
                for tipo in sorted(por_talla[talla]):
                    filas = [
                        {"proveedor": prov, "peso_por_caja": peso, "cajas": cajas, "kilos": kilos}
                        for (prov, peso), (cajas, kilos) in sorted(por_talla[talla][tipo].items())
                    ]
                    tipos.append({"tipo": tipo, "filas": filas})
                existencia = sum((f["kilos"] for t in tipos for f in t["filas"]), CERO)
                tallas.append({"talla": talla, "existencia": existencia, "tipos": tipos})
            subtotal = sum((t["existencia"] for t in tallas), CERO)
            bloques.append({"presentacion": presentacion, "etiqueta": etiqueta, "tallas": tallas, "subtotal": subtotal})
        total = sum((b["subtotal"] for b in bloques), CERO)
        reporte.append({"camara": {"id": cam_id, "nombre": nombre}, "bloques": bloques, "total": total})
    return reporte


def filas_planas(bloque):
    """
    Las filas de un bloque en el orden en que se dibujan, con los tamaños de
    las celdas combinadas: (talla, span_talla, tipo, span_tipo, fila, es_primera_de_talla, es_primera_de_tipo).
    """
    for talla in bloque["tallas"]:
        span_talla = sum(len(t["filas"]) for t in talla["tipos"])
        primera_talla = True
        for tipo in talla["tipos"]:
            primera_tipo = True
            for fila in tipo["filas"]:
                yield talla, span_talla, tipo, len(tipo["filas"]), fila, primera_talla, primera_tipo
                primera_talla = primera_tipo = False


def serializar_reporte(reporte):
    """Decimales a texto para JSON (el front los formatea)."""
    def num(valor):
        return str(valor.quantize(Decimal("0.01")))

    return [
        {
            "camara": c["camara"],
            "total": num(c["total"]),
            "bloques": [
                {
                    "presentacion": b["presentacion"],
                    "etiqueta": b["etiqueta"],
                    "subtotal": num(b["subtotal"]),
                    "tallas": [
                        {
                            "talla": t["talla"],
                            "existencia": num(t["existencia"]),
                            "tipos": [
                                {
                                    "tipo": tp["tipo"],
                                    "filas": [
                                        {
                                            "proveedor": f["proveedor"],
                                            "peso_por_caja": num(f["peso_por_caja"]),
                                            "cajas": f["cajas"],
                                            "kilos": num(f["kilos"]),
                                        }
                                        for f in tp["filas"]
                                    ],
                                }
                                for tp in t["tipos"]
                            ],
                        }
                        for t in b["tallas"]
                    ],
                }
                for b in c["bloques"]
            ],
        }
        for c in reporte
    ]


def _peso(valor):
    """18.00 -> "18", 18.50 -> "18.5": el formato del área muestra el peso master sin decimales de más."""
    return f"{valor.normalize():f}" if valor == valor.to_integral() else f"{valor:.2f}".rstrip("0")


# ---------------------------------------------------------------- PDF
def construir_pdf(reporte, fecha, subtitulo=""):
    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=letter, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm, title="Existencias por cámara",
    )
    estilo = ParagraphStyle("celda", fontName="Helvetica", fontSize=7.5, leading=9)
    negrita = ParagraphStyle("negrita", parent=estilo, fontName="Helvetica-Bold")
    # El color va en el estilo del Paragraph: el TEXTCOLOR de la tabla no aplica a celdas Paragraph.
    encabezado = ParagraphStyle("encabezado", parent=negrita, textColor=colors.HexColor(f"#{VERDE}"))

    def p(texto, alineacion=TA_LEFT, estilo_base=estilo):
        return Paragraph(str(texto), ParagraphStyle("x", parent=estilo_base, alignment=alineacion))

    anchos = [2.6 * cm, 2.4 * cm, 4.2 * cm, 2.3 * cm, 2.2 * cm, 2.6 * cm, 2.6 * cm]
    elementos = []
    if not reporte:
        elementos.append(Paragraph(f"No hay existencias al {fecha:%d/%m/%Y} con los filtros seleccionados.", estilo))

    for indice, camara in enumerate(reporte):
        if indice:
            elementos.append(PageBreak())
        titulo = f"EXISTENCIAS {camara['camara']['nombre']}"
        filas = [
            [p(titulo, TA_LEFT, negrita), "", p(subtitulo), "", "", p(f"{fecha:%d/%m/%Y}", TA_CENTER, negrita), ""],
            [p(c, TA_CENTER, encabezado) for c in COLUMNAS],
        ]
        estilos = [
            ("SPAN", (0, 0), (1, 0)), ("SPAN", (2, 0), (4, 0)),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ]
        for bloque in camara["bloques"]:
            if len(camara["bloques"]) > 1:
                filas.append([p(bloque["etiqueta"].upper(), TA_LEFT, negrita)] + [""] * 6)
                estilos.append(("SPAN", (0, len(filas) - 1), (-1, len(filas) - 1)))
            for talla, span_t, tipo, span_tp, fila, prim_talla, prim_tipo in filas_planas(bloque):
                r = len(filas)
                filas.append([
                    p(talla["talla"], TA_CENTER) if prim_talla else "",
                    p(tipo["tipo"], TA_CENTER) if prim_tipo else "",
                    p(fila["proveedor"]),
                    p(_peso(fila["peso_por_caja"]), TA_CENTER),
                    p(f"{fila['cajas']:,}", TA_RIGHT),
                    p(f"{fila['kilos']:,.2f}", TA_RIGHT),
                    p(f"{talla['existencia']:,.2f}", TA_RIGHT) if prim_talla else "",
                ])
                if prim_talla and span_t > 1:
                    estilos += [("SPAN", (0, r), (0, r + span_t - 1)), ("SPAN", (6, r), (6, r + span_t - 1))]
                if prim_tipo and span_tp > 1:
                    estilos.append(("SPAN", (1, r), (1, r + span_tp - 1)))
            filas.append(["", "", "", "", p("SUBTOTAL", TA_CENTER, negrita), p(f"{bloque['subtotal']:,.2f}", TA_RIGHT, negrita), ""])
        filas.append(["", "", "", "", p("TOTAL", TA_CENTER, negrita), p(f"{camara['total']:,.2f}", TA_RIGHT, negrita), ""])

        tabla = Table(filas, colWidths=anchos, repeatRows=2)
        tabla.setStyle(TableStyle(estilos))
        elementos.append(tabla)

    documento.build(elementos)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------- Excel
def construir_excel(reporte, fecha, subtitulo=""):
    wb = Workbook()
    wb.remove(wb.active)
    delgado = Side(style="thin", color="000000")
    borde = Border(left=delgado, right=delgado, top=delgado, bottom=delgado)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    negrita = Font(name="Calibri", bold=True)
    verde = Font(name="Calibri", bold=True, color=VERDE)

    if not reporte:
        ws = wb.create_sheet("Existencias")
        ws["A1"] = f"No hay existencias al {fecha:%d/%m/%Y} con los filtros seleccionados."

    usados = set()
    for camara in reporte:
        nombre = re.sub(r"[\\/*?:\[\]]", "", camara["camara"]["nombre"])[:31] or "Cámara"
        while nombre in usados:
            nombre = nombre[:28] + f"_{len(usados)}"
        usados.add(nombre)
        ws = wb.create_sheet(nombre)

        ws.cell(1, 1, f"EXISTENCIAS {camara['camara']['nombre']}").font = negrita
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        if subtitulo:
            ws.cell(1, 3, subtitulo)
            ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=5)
        ws.cell(1, 6, fecha).number_format = "dd/mm/yyyy"
        ws.cell(1, 6).font = negrita
        ws.cell(1, 6).alignment = centro
        for col, titulo in enumerate(COLUMNAS, start=1):
            celda = ws.cell(2, col, titulo)
            celda.font = verde
            celda.alignment = centro

        r = 3
        for bloque in camara["bloques"]:
            if len(camara["bloques"]) > 1:
                ws.cell(r, 1, bloque["etiqueta"].upper()).font = negrita
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
                r += 1
            for talla, span_t, tipo, span_tp, fila, prim_talla, prim_tipo in filas_planas(bloque):
                if prim_talla:
                    ws.cell(r, 1, talla["talla"]).alignment = centro
                    ws.cell(r, 7, float(talla["existencia"])).number_format = "#,##0.00"
                    ws.cell(r, 7).alignment = Alignment(vertical="center")
                    if span_t > 1:
                        ws.merge_cells(start_row=r, start_column=1, end_row=r + span_t - 1, end_column=1)
                        ws.merge_cells(start_row=r, start_column=7, end_row=r + span_t - 1, end_column=7)
                if prim_tipo:
                    ws.cell(r, 2, tipo["tipo"]).alignment = centro
                    if span_tp > 1:
                        ws.merge_cells(start_row=r, start_column=2, end_row=r + span_tp - 1, end_column=2)
                ws.cell(r, 3, fila["proveedor"])
                ws.cell(r, 4, float(fila["peso_por_caja"])).alignment = centro
                ws.cell(r, 5, fila["cajas"]).number_format = "#,##0"
                ws.cell(r, 6, float(fila["kilos"])).number_format = "#,##0.00"
                r += 1
            ws.cell(r, 5, "SUBTOTAL").font = negrita
            ws.cell(r, 6, float(bloque["subtotal"])).number_format = "#,##0.00"
            ws.cell(r, 6).font = negrita
            r += 1
        ws.cell(r, 5, "TOTAL").font = negrita
        ws.cell(r, 6, float(camara["total"])).number_format = "#,##0.00"
        ws.cell(r, 6).font = negrita

        for fila_celdas in ws.iter_rows(min_row=1, max_row=r, max_col=7):
            for celda in fila_celdas:
                celda.border = borde
        for col, ancho in enumerate([12, 12, 26, 13, 11, 15, 15], start=1):
            ws.column_dimensions[get_column_letter(col)].width = ancho
        ws.freeze_panes = "A3"
        ws.print_title_rows = "1:2"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

