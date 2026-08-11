from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db.models import Count, DecimalField, F, Sum
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from openpyxl import Workbook

from .models import Banco, NominaDetalle, Rancho
from .reportes import (
    COLOR_ENCABEZADO_RL,
    ESTILOS,
    _celda_detalle,
    _estilizar_encabezado,
    _parsear_fecha,
    _pie_pagina_pdf,
)

PREVIEW_LIMIT = 20


def obtener_detalles_filtrados(params):
    """
    params: QueryDict (request.query_params). Devuelve un queryset de
    NominaDetalle —una fila por empleado dentro de una nómina semanal—
    con los filtros del reporte ya aplicados. Una nómina "cae" dentro del
    rango si su semana se solapa con el rango pedido (no que esté contenida
    por completo), igual que se esperaría de un reporte por rango de fechas.
    """
    fecha_inicio_raw = params.get("fecha_inicio")
    fecha_fin_raw = params.get("fecha_fin")

    if not fecha_inicio_raw or not fecha_fin_raw:
        raise ValidationError("Debes indicar la fecha de inicio y la fecha de fin.")

    fecha_inicio = _parsear_fecha(fecha_inicio_raw, "fecha de inicio")
    fecha_fin = _parsear_fecha(fecha_fin_raw, "fecha de fin")

    if fecha_inicio > fecha_fin:
        raise ValidationError("La fecha de inicio no puede ser posterior a la fecha de fin.")

    detalles = (
        NominaDetalle.objects
        .select_related("nomina", "empleado", "empleado__rancho", "empleado__puesto", "empleado__banco")
        .filter(nomina__fecha_inicio__lte=fecha_fin, nomina__fecha_fin__gte=fecha_inicio)
    )

    rancho_ids = params.getlist("rancho") if hasattr(params, "getlist") else None
    if rancho_ids:
        detalles = detalles.filter(empleado__rancho_id__in=rancho_ids)

    banco_ids = params.getlist("banco") if hasattr(params, "getlist") else None
    if banco_ids:
        detalles = detalles.filter(empleado__banco_id__in=banco_ids)

    return detalles.order_by("nomina__fecha_inicio", "empleado__rancho__nombre", "empleado__nombre")


def calcular_resumen_por_rancho(detalles):
    campo_decimal = DecimalField(max_digits=18, decimal_places=2)

    return list(
        detalles.values("empleado__rancho_id", "empleado__rancho__nombre")
        .annotate(
            empleados=Count("empleado_id", distinct=True),
            total_bruto=Sum(F("dias_trabajados") * F("salario_diario"), output_field=campo_decimal),
            total_descuento=Sum("descuento"),
        )
        .order_by("empleado__rancho__nombre")
    )


def calcular_resumen_por_banco(detalles):
    campo_decimal = DecimalField(max_digits=18, decimal_places=2)

    return list(
        detalles.values("empleado__banco_id", "empleado__banco__nombre")
        .annotate(
            empleados=Count("empleado_id", distinct=True),
            total_bruto=Sum(F("dias_trabajados") * F("salario_diario"), output_field=campo_decimal),
            total_descuento=Sum("descuento"),
        )
        .order_by("empleado__banco__nombre")
    )


def describir_filtros_nomina(params):
    """
    Arma una lista de frases legibles ("Ranchos: Norte, Sur") con los
    filtros distintos al rango de fechas, para mostrarlas como subtítulo
    en el PDF. El rango de fechas se muestra aparte porque siempre está
    presente.
    """
    partes = []

    rancho_ids = params.getlist("rancho") if hasattr(params, "getlist") else None
    if rancho_ids:
        nombres = list(Rancho.objects.filter(id__in=rancho_ids).order_by("nombre").values_list("nombre", flat=True))
        if nombres:
            partes.append("Ranchos: " + ", ".join(nombres))

    banco_ids = params.getlist("banco") if hasattr(params, "getlist") else None
    if banco_ids:
        nombres = list(Banco.objects.filter(id__in=banco_ids).order_by("nombre").values_list("nombre", flat=True))
        if nombres:
            partes.append("Bancos: " + ", ".join(nombres))

    return partes


def _llenar_hoja_resumen(ws, encabezado_primera_columna, resumen, campo_nombre, sin_valor):
    encabezados_resumen = [encabezado_primera_columna, "Empleados", "Total bruto", "Descuento", "Total neto"]
    ws.append(encabezados_resumen)
    _estilizar_encabezado(ws, len(encabezados_resumen))

    for fila_datos in resumen:
        total_bruto = fila_datos["total_bruto"] or Decimal("0")
        total_descuento = fila_datos["total_descuento"] or Decimal("0")
        ws.append([
            fila_datos[campo_nombre] or sin_valor,
            fila_datos["empleados"],
            float(total_bruto),
            float(total_descuento),
            float(total_bruto - total_descuento),
        ])
        fila = ws.max_row
        for col in (3, 4, 5):
            ws.cell(row=fila, column=col).number_format = "#,##0.00"

    anchos_resumen = {"A": 18, "B": 12, "C": 16, "D": 14, "E": 16}
    for col, ancho in anchos_resumen.items():
        ws.column_dimensions[col].width = ancho


def construir_libro_excel_nomina(detalles, resumen, resumen_banco):
    wb = Workbook()

    ws = wb.active
    ws.title = "Detalle"
    encabezados = [
        "Semana", "Rancho", "Empleado", "Puesto", "Días trabajados",
        "Salario diario", "Descuento", "Total bruto", "Total neto",
        "Banco", "Número de cuenta", "Nombre de la cuenta",
    ]
    ws.append(encabezados)
    _estilizar_encabezado(ws, len(encabezados))

    for detalle in detalles:
        ws.append([
            f"{detalle.nomina.fecha_inicio.strftime('%d/%m/%Y')} - {detalle.nomina.fecha_fin.strftime('%d/%m/%Y')}",
            detalle.empleado.rancho.nombre,
            detalle.empleado.nombre,
            detalle.empleado.puesto.nombre,
            float(detalle.dias_trabajados),
            float(detalle.salario_diario),
            float(detalle.descuento),
            float(detalle.total_bruto),
            float(detalle.total_neto),
            detalle.empleado.banco.nombre if detalle.empleado.banco else "",
            detalle.empleado.numero_cuenta,
            detalle.empleado.nombre_cuenta,
        ])
        fila = ws.max_row
        for col in (5, 6, 7, 8, 9):
            ws.cell(row=fila, column=col).number_format = "#,##0.00"

    anchos = {
        "A": 22, "B": 16, "C": 24, "D": 16, "E": 14, "F": 14, "G": 12, "H": 14, "I": 14,
        "J": 16, "K": 18, "L": 22,
    }
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    ws_resumen = wb.create_sheet("Resumen por rancho")
    _llenar_hoja_resumen(ws_resumen, "Rancho", resumen, "empleado__rancho__nombre", "Sin rancho")

    ws_resumen_banco = wb.create_sheet("Resumen por banco")
    _llenar_hoja_resumen(ws_resumen_banco, "Banco", resumen_banco, "empleado__banco__nombre", "Sin banco")

    return wb


def _tabla_resumen_pdf(encabezados, resumen, campo_nombre, sin_valor):
    filas = [encabezados]
    for fila in resumen:
        total_bruto = fila["total_bruto"] or Decimal("0")
        total_descuento = fila["total_descuento"] or Decimal("0")
        filas.append([
            fila[campo_nombre] or sin_valor,
            str(fila["empleados"]),
            f'{total_bruto:,.2f}',
            f'{total_descuento:,.2f}',
            f'{(total_bruto - total_descuento):,.2f}',
        ])

    tabla = Table(filas, colWidths=[4 * cm, 3 * cm, 3.5 * cm, 3.2 * cm, 3.5 * cm])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return tabla, len(filas)


def construir_pdf_nomina(detalles, resumen, resumen_banco, fecha_inicio, fecha_fin, filtros_descripcion):
    detalles = list(detalles)

    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title="Reporte de nómina semanal",
    )

    estilo_empresa = ParagraphStyle(
        "empresa_nomina", parent=ESTILOS["Heading1"], fontSize=15, textColor=COLOR_ENCABEZADO_RL, spaceAfter=1,
    )
    estilo_titulo = ParagraphStyle(
        "titulo_nomina", parent=ESTILOS["Normal"], fontSize=12, textColor=colors.black, spaceAfter=6,
    )
    estilo_subtitulo = ParagraphStyle(
        "subtitulo_nomina", parent=ESTILOS["Normal"], fontSize=9.5, textColor=colors.HexColor("#444444"), spaceAfter=2,
    )
    estilo_generado = ParagraphStyle(
        "generado_nomina", parent=ESTILOS["Normal"], fontSize=8, textColor=colors.HexColor("#999999"), spaceBefore=4,
    )
    estilo_h2 = ParagraphStyle(
        "h2_nomina", parent=ESTILOS["Heading2"], fontSize=12, textColor=COLOR_ENCABEZADO_RL, spaceAfter=4,
    )

    elementos = [
        Paragraph("Importadora y Exportadora de Mariscos", estilo_empresa),
        Paragraph("Reporte de nómina semanal", estilo_titulo),
        Paragraph(f"Del {fecha_inicio.strftime('%d/%m/%Y')} al {fecha_fin.strftime('%d/%m/%Y')}", estilo_subtitulo),
    ]
    if filtros_descripcion:
        elementos.append(Paragraph("Filtros: " + " · ".join(filtros_descripcion), estilo_subtitulo))
    elementos.append(Paragraph(
        f"Generado el {timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M')}",
        estilo_generado,
    ))
    elementos.append(Spacer(1, 0.6 * cm))

    # --- Resumen por rancho ---
    elementos.append(Paragraph("Resumen por rancho", estilo_h2))
    tabla_resumen, filas_resumen_n = _tabla_resumen_pdf(
        ["Rancho", "Empleados", "Total bruto", "Descuento", "Total neto"],
        resumen, "empleado__rancho__nombre", "Sin rancho",
    )
    if filas_resumen_n == 1:
        elementos.append(Paragraph("No hay nóminas para los filtros seleccionados.", ESTILOS["Normal"]))
    else:
        elementos.append(tabla_resumen)

    elementos.append(Spacer(1, 0.7 * cm))

    # --- Resumen por banco ---
    elementos.append(Paragraph("Resumen por banco", estilo_h2))
    tabla_resumen_banco, filas_resumen_banco_n = _tabla_resumen_pdf(
        ["Banco", "Empleados", "Total bruto", "Descuento", "Total neto"],
        resumen_banco, "empleado__banco__nombre", "Sin banco",
    )
    if filas_resumen_banco_n == 1:
        elementos.append(Paragraph("No hay nóminas para los filtros seleccionados.", ESTILOS["Normal"]))
    else:
        elementos.append(tabla_resumen_banco)

    elementos.append(Spacer(1, 0.7 * cm))

    # --- Detalle ---
    elementos.append(Paragraph("Detalle por empleado", estilo_h2))
    encabezados_detalle = [
        "Semana", "Rancho", "Empleado", "Puesto", "Días", "Salario diario", "Descuento", "Total bruto", "Total neto",
    ]
    filas_detalle = [encabezados_detalle]

    for detalle in detalles:
        semana = f"{detalle.nomina.fecha_inicio.strftime('%d/%m/%Y')} - {detalle.nomina.fecha_fin.strftime('%d/%m/%Y')}"
        filas_detalle.append([
            _celda_detalle(semana, TA_CENTER),
            _celda_detalle(detalle.empleado.rancho.nombre, TA_LEFT),
            _celda_detalle(detalle.empleado.nombre, TA_LEFT),
            _celda_detalle(detalle.empleado.puesto.nombre, TA_LEFT),
            _celda_detalle(f"{detalle.dias_trabajados:,.1f}", TA_RIGHT),
            _celda_detalle(f"{detalle.salario_diario:,.2f}", TA_RIGHT),
            _celda_detalle(f"{detalle.descuento:,.2f}", TA_RIGHT),
            _celda_detalle(f"{detalle.total_bruto:,.2f}", TA_RIGHT),
            _celda_detalle(f"{detalle.total_neto:,.2f}", TA_RIGHT),
        ])

    if len(filas_detalle) == 1:
        elementos.append(Paragraph("No hay nóminas para los filtros seleccionados.", ESTILOS["Normal"]))
    else:
        anchos_detalle = [3 * cm, 2.6 * cm, 3.4 * cm, 2.6 * cm, 1.6 * cm, 2.4 * cm, 2.2 * cm, 2.4 * cm, 2.4 * cm]
        tabla_detalle = Table(filas_detalle, colWidths=anchos_detalle, repeatRows=1)
        tabla_detalle.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8.5),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        elementos.append(tabla_detalle)

    documento.build(elementos, onFirstPage=_pie_pagina_pdf, onLaterPages=_pie_pagina_pdf)
    buffer.seek(0)
    return buffer
