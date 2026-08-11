from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NOMBRE_EMPRESA = "IMPORTADORA Y EXPORTADORA DE MARISCOS DE CENTRO AMERICA Y EL CARIBE, S.A. DE C.V."
RFC_EMPRESA = "IEM110811AR1"

COLOR_ENCABEZADO = "0E6B6F"
COLOR_ENCABEZADO_RL = colors.HexColor(f"#{COLOR_ENCABEZADO}")
COLOR_FALTANTE_RL = colors.HexColor("#C0392B")
COLOR_SOBRANTE_RL = colors.HexColor("#1E8449")
COLOR_EXACTO_RL = colors.HexColor("#555555")


def calcular_resumen_arqueo(arqueo):
    """
    Suma los totales de cada divisa del arqueo sin conversión de moneda
    (cada divisa aporta su valor "de a como es"), replicando cómo se hace
    hoy en el arqueo físico: no hay tipo de cambio de por medio.
    """
    divisas = list(
        arqueo.divisas.select_related('divisa').prefetch_related('conteos__denominacion').order_by('divisa__codigo')
    )
    total_contado = sum((linea.total_contado for linea in divisas), Decimal('0'))
    resultado_esperado = sum((linea.resultado_esperado for linea in divisas), Decimal('0'))
    diferencia = total_contado - resultado_esperado

    if diferencia > 0:
        estado = 'SOBRANTE'
    elif diferencia < 0:
        estado = 'FALTANTE'
    else:
        estado = 'EXACTO'

    return {
        'divisas': divisas,
        'total_contado': total_contado,
        'resultado_esperado': resultado_esperado,
        'diferencia': diferencia,
        'estado': estado,
    }


def _conteos_ordenados(linea_divisa):
    return sorted(
        linea_divisa.conteos.all(),
        key=lambda conteo: (conteo.denominacion.tipo != 'B', -conteo.denominacion.valor),
    )


def nombre_archivo_arqueo(arqueo):
    estado_corte = 'final' if arqueo.corte.cerrado else 'preliminar'
    fecha = timezone.localtime(arqueo.corte.fecha).strftime('%Y-%m-%d')
    return f"arqueo_{arqueo.id}_{fecha}_{estado_corte}"


# --- Excel ---

def _estilizar_fila(ws, fila, num_columnas):
    for col in range(1, num_columnas + 1):
        celda = ws.cell(row=fila, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(vertical="center")
    ws.row_dimensions[fila].height = 20


def construir_libro_excel_arqueo(arqueo, resumen):
    wb = Workbook()
    ws = wb.active
    ws.title = "Arqueo"

    ws.append([NOMBRE_EMPRESA])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13)
    ws.append([f"RFC: {RFC_EMPRESA}"])
    ws.append([])

    estado_corte = "FINAL (corte cerrado)" if arqueo.corte.cerrado else "PRELIMINAR (corte todavía abierto)"
    ws.append([f"ARQUEO DE CAJA — {estado_corte}"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append([f"Arqueo N°: {arqueo.id}", "Corte:", f"#{arqueo.corte_id}"])
    ws.append([f"Fecha: {timezone.localtime(arqueo.corte.fecha).strftime('%d/%m/%Y')}"])
    ws.append([
        f"Hora inicio: {timezone.localtime(arqueo.hora_inicio).strftime('%H:%M')}",
        "Hora término:", timezone.localtime(arqueo.hora_termino).strftime('%H:%M'),
    ])
    ws.append([])

    for linea in resumen['divisas']:
        ws.append([f"{linea.divisa.codigo} — {linea.divisa.nombre}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11)

        encabezados = ["Denominación", "Tipo", "Piezas", "Total"]
        ws.append(encabezados)
        _estilizar_fila(ws, ws.max_row, len(encabezados))

        for conteo in _conteos_ordenados(linea):
            ws.append([
                float(conteo.denominacion.valor),
                "Billete" if conteo.denominacion.tipo == "B" else "Moneda",
                conteo.piezas,
                float(conteo.total),
            ])
            fila = ws.max_row
            ws.cell(row=fila, column=1).number_format = "#,##0.00"
            ws.cell(row=fila, column=4).number_format = "#,##0.00"

        for etiqueta, valor in (
            ("Saldo inicial", linea.saldo_inicial),
            ("Saldo disponible", linea.resultado_esperado),
            ("Total contado", linea.total_contado),
            ("Diferencia", linea.diferencia),
        ):
            ws.append(["", "", etiqueta, float(valor)])
            ws.cell(row=ws.max_row, column=4).number_format = "#,##0.00"
        ws.append(["", "", "Estado", linea.estado])
        ws.append([])

    ws.append(["RESUMEN GENERAL"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11)
    for etiqueta, valor in (
        ("Resultado esperado (suma de divisas, sin conversión)", resumen['resultado_esperado']),
        ("Total contado (suma de divisas, sin conversión)", resumen['total_contado']),
        ("Diferencia", resumen['diferencia']),
    ):
        ws.append([etiqueta, float(valor)])
        ws.cell(row=ws.max_row, column=2).number_format = "#,##0.00"
    ws.append(["Estado", resumen['estado']])

    if arqueo.observaciones:
        ws.append([])
        ws.append(["Observaciones:", arqueo.observaciones])

    if arqueo.leyenda_totales:
        ws.append([])
        ws.append([arqueo.leyenda_totales])

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 16

    return wb


# --- PDF ---

ESTILOS = getSampleStyleSheet()


def _pie_pagina_pdf(canvas, documento):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#888888"))
    ancho_pagina, _ = documento.pagesize
    canvas.drawRightString(ancho_pagina - 1.2 * cm, 0.7 * cm, f"Página {documento.page}")
    canvas.restoreState()


def _color_estado(estado):
    if estado == 'FALTANTE':
        return COLOR_FALTANTE_RL
    if estado == 'SOBRANTE':
        return COLOR_SOBRANTE_RL
    return COLOR_EXACTO_RL


def _tabla_denominaciones(linea):
    encabezados = ["Denominación", "Tipo", "Piezas", "Total"]
    filas = [encabezados]
    for conteo in _conteos_ordenados(linea):
        filas.append([
            f"{linea.divisa.simbolo}{conteo.denominacion.valor:,.2f}",
            "Billete" if conteo.denominacion.tipo == "B" else "Moneda",
            str(conteo.piezas),
            f"{linea.divisa.simbolo}{conteo.total:,.2f}",
        ])

    tabla = Table(filas, colWidths=[4 * cm, 3 * cm, 2.5 * cm, 3.5 * cm], repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return tabla


def _tabla_subtotales_divisa(linea):
    filas = [
        ["Saldo inicial", f"{linea.divisa.simbolo}{linea.saldo_inicial:,.2f}"],
        ["Saldo disponible", f"{linea.divisa.simbolo}{linea.resultado_esperado:,.2f}"],
        ["Total contado", f"{linea.divisa.simbolo}{linea.total_contado:,.2f}"],
        ["Diferencia", f"{linea.divisa.simbolo}{linea.diferencia:,.2f}"],
        ["Estado", linea.estado],
    ]
    tabla = Table(filas, colWidths=[4 * cm, 3.5 * cm])
    tabla.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("FONTNAME", (0, -2), (-1, -2), "Helvetica-Bold"),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TEXTCOLOR", (1, -1), (1, -1), _color_estado(linea.estado)),
        ("FONTNAME", (1, -1), (1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return tabla


def construir_pdf_arqueo(arqueo, resumen):
    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=f"Arqueo de caja #{arqueo.id}",
    )

    estilo_empresa = ParagraphStyle(
        "empresa", parent=ESTILOS["Heading2"], fontSize=12, textColor=COLOR_ENCABEZADO_RL, spaceAfter=1,
    )
    estilo_rfc = ParagraphStyle(
        "rfc", parent=ESTILOS["Normal"], fontSize=8.5, textColor=colors.HexColor("#555555"), spaceAfter=6,
    )
    estilo_titulo = ParagraphStyle(
        "titulo", parent=ESTILOS["Heading1"], fontSize=14, alignment=TA_CENTER, spaceAfter=4,
    )
    es_final = arqueo.corte.cerrado
    estilo_estado = ParagraphStyle(
        "estado_corte", parent=ESTILOS["Normal"], fontSize=10, alignment=TA_CENTER,
        textColor=COLOR_SOBRANTE_RL if es_final else colors.HexColor("#B7791F"), spaceAfter=8,
    )
    estilo_dato = ParagraphStyle("dato", parent=ESTILOS["Normal"], fontSize=9, spaceAfter=2)
    estilo_h2 = ParagraphStyle(
        "h2", parent=ESTILOS["Heading2"], fontSize=11.5, textColor=COLOR_ENCABEZADO_RL, spaceBefore=10, spaceAfter=4,
    )
    estilo_observaciones = ParagraphStyle("observaciones", parent=ESTILOS["Normal"], fontSize=9, spaceBefore=4)
    estilo_leyenda = ParagraphStyle(
        "leyenda", parent=ESTILOS["Normal"], fontSize=9, spaceBefore=10, alignment=TA_LEFT,
    )

    elementos = [
        Paragraph(NOMBRE_EMPRESA, estilo_empresa),
        Paragraph(f"RFC: {RFC_EMPRESA}", estilo_rfc),
        Paragraph("ARQUEO DE CAJA", estilo_titulo),
        Paragraph(
            "ARQUEO FINAL — corte cerrado" if es_final else "ARQUEO PRELIMINAR — el corte sigue abierto",
            estilo_estado,
        ),
        Paragraph(f"<b>Arqueo N°:</b> {arqueo.id} &nbsp;&nbsp; <b>Corte:</b> #{arqueo.corte_id}", estilo_dato),
        Paragraph(f"<b>Fecha:</b> {timezone.localtime(arqueo.corte.fecha).strftime('%d/%m/%Y')}", estilo_dato),
        Paragraph(
            f"<b>Hora inicio:</b> {timezone.localtime(arqueo.hora_inicio).strftime('%H:%M')} "
            f"&nbsp;&nbsp; <b>Hora término:</b> {timezone.localtime(arqueo.hora_termino).strftime('%H:%M')}",
            estilo_dato,
        ),
        Paragraph(f"<b>Responsable:</b> {arqueo.usuario.get_full_name() or arqueo.usuario.username}", estilo_dato),
    ]

    for linea in resumen['divisas']:
        elementos.append(KeepTogether([
            Paragraph(f"{linea.divisa.codigo} — {linea.divisa.nombre}", estilo_h2),
            Table(
                [[_tabla_denominaciones(linea), _tabla_subtotales_divisa(linea)]],
                colWidths=[13.2 * cm, 7.6 * cm],
                style=TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]),
            ),
        ]))

    elementos.append(Paragraph("Resumen general", estilo_h2))
    filas_resumen = [
        ["Resultado esperado (suma de divisas, sin conversión)", f"{resumen['resultado_esperado']:,.2f}"],
        ["Total contado (suma de divisas, sin conversión)", f"{resumen['total_contado']:,.2f}"],
        ["Diferencia", f"{resumen['diferencia']:,.2f}"],
        ["Estado", resumen['estado']],
    ]
    tabla_resumen = Table(filas_resumen, colWidths=[11 * cm, 4 * cm])
    tabla_resumen.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TEXTCOLOR", (1, -1), (1, -1), _color_estado(resumen['estado'])),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabla_resumen)

    if arqueo.observaciones:
        elementos.append(Paragraph("Observaciones", estilo_h2))
        elementos.append(Paragraph(arqueo.observaciones, estilo_observaciones))

    if arqueo.leyenda_totales:
        elementos.append(Spacer(1, 0.4 * cm))
        elementos.append(Paragraph(arqueo.leyenda_totales, estilo_leyenda))

    documento.build(elementos, onFirstPage=_pie_pagina_pdf, onLaterPages=_pie_pagina_pdf)
    buffer.seek(0)
    return buffer
