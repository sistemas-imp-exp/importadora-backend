from datetime import timedelta
from decimal import Decimal

from django.utils import timezone
from rest_framework.test import APITestCase

from security.permissions import AREA_TESORERIA
from security.testing import crear_usuario_con_area

from .models import Deuda, Deudor, Divisa, MovimientoTesoreria
from .saldos import saldos_al

HOY = timezone.localdate()


class DeudasApiTests(APITestCase):
    def setUp(self):
        self.usuario = crear_usuario_con_area("cajero", AREA_TESORERIA)
        self.client.force_authenticate(user=self.usuario)
        self.mxn = Divisa.objects.create(codigo="MXN", nombre="Peso", simbolo="$")
        self.usd = Divisa.objects.create(codigo="USD", nombre="Dólar", simbolo="US$")
        self.deudor = Deudor.objects.create(nombre="Juan Pérez", tipo=Deudor.TRABAJADOR, usuario=self.usuario)
        self.deuda = Deuda.objects.create(
            deudor=self.deudor, concepto="Venta de camarón", divisa=self.mxn, monto_inicial=Decimal("1000"),
            fecha=HOY - timedelta(days=5), usuario=self.usuario,
        )
        self.folio = 0

    def _abonar(self, cantidad, deuda=None, divisas=None, tipo="I", fecha=None):
        self.folio += 1
        return self.client.post("/api/treasury/movimientos/", {
            "fecha": (fecha or HOY).isoformat(), "folio": f"A-{self.folio}", "tipo": tipo,
            "autorizo": "Jefe", "beneficiario": "Juan Pérez", "concepto": "Abono",
            "aplicaciones": [{"deuda_id": (deuda or self.deuda).id}],
            "divisas": divisas or [{"divisa_id": self.mxn.id, "cantidad": cantidad}],
        }, format="json")

    def _deuda(self):
        return self.client.get(f"/api/treasury/deudas/{self.deuda.id}/").data

    def test_crear_deudor_y_deuda_registra_al_usuario_de_la_sesion(self):
        r = self.client.post("/api/treasury/deudores/", {"nombre": "  Pesca   Sur ", "tipo": "E"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["nombre"], "Pesca Sur")
        r = self.client.post("/api/treasury/deudas/", {
            "deudor_id": r.data["id"], "concepto": "Venta de pulpo", "divisa_id": self.usd.id,
            "monto_inicial": "500", "fecha": HOY.isoformat(),
        }, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["usuario"]["username"], "cajero")
        self.assertEqual((r.data["estado"], r.data["saldo"]), ("abierta", "500.00"))

    def test_la_deuda_no_mueve_la_caja(self):
        self.assertEqual(saldos_al(HOY), {})

    def test_abono_es_un_ingreso_que_baja_el_saldo_y_queda_en_el_historial(self):
        r = self._abonar("300")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["deudas"][0]["deudor"]["nombre"], "Juan Pérez")
        deuda = self._deuda()
        self.assertEqual((deuda["abonado"], deuda["saldo"], deuda["estado"]), ("300.00", "700.00", "abierta"))
        self.assertEqual([a["folio"] for a in deuda["movimientos"]], ["A-1"])
        self.assertEqual(saldos_al(HOY)[self.mxn.id], Decimal("300"))

    def test_abonar_el_saldo_completo_la_marca_pagada(self):
        self.assertEqual(self._abonar("600").status_code, 201)
        self.assertEqual(self._abonar("400").status_code, 201)
        self.assertEqual(self._deuda()["estado"], "pagada")
        r = self.client.get("/api/treasury/deudas/", {"estado": "abierta"})
        self.assertEqual(r.data, [])

    def test_no_se_abona_mas_que_el_saldo_pendiente(self):
        self._abonar("900")
        r = self._abonar("100.01")
        self.assertEqual(r.status_code, 400)
        self.assertIn("saldo pendiente", str(r.data["aplicaciones"]))

    def test_validaciones_del_abono(self):
        r = self._abonar("10", divisas=[{"divisa_id": self.usd.id, "cantidad": "10"}])
        self.assertIn("MXN", str(r.data["divisas"]))
        r = self._abonar("10", divisas=[
            {"divisa_id": self.mxn.id, "cantidad": "10"}, {"divisa_id": self.usd.id, "cantidad": "10"},
        ])
        self.assertEqual(r.status_code, 400)
        r = self._abonar("10", fecha=HOY - timedelta(days=6))
        self.assertIn("anterior", str(r.data["fecha"]))
        self.assertFalse(MovimientoTesoreria.objects.exists())

    def test_cancelar_el_ingreso_revierte_el_abono_y_queda_en_historial(self):
        movimiento_id = self._abonar("1000").data["id"]
        self.client.post(f"/api/treasury/movimientos/{movimiento_id}/cancelar/", {"motivo": "Error"}, format="json")
        deuda = self._deuda()
        self.assertEqual((deuda["saldo"], deuda["estado"]), ("1000.00", "abierta"))
        self.assertTrue(deuda["movimientos"][0]["cancelado"])

    def test_editar_el_ingreso_valida_el_saldo_sin_contarse_a_si_mismo(self):
        r = self._abonar("1000")
        datos = {
            "fecha": HOY.isoformat(), "folio": "A-1", "tipo": "I", "autorizo": "Jefe", "beneficiario": "Juan",
            "concepto": "Abono", "divisas": [{"divisa_id": self.mxn.id, "cantidad": "800"}],
        }
        self.assertEqual(self.client.put(f"/api/treasury/movimientos/{r.data['id']}/", datos, format="json").status_code, 200)
        datos["divisas"][0]["cantidad"] = "1000.01"
        self.assertEqual(self.client.put(f"/api/treasury/movimientos/{r.data['id']}/", datos, format="json").status_code, 400)

    def test_deuda_con_abonos_no_se_edita_ni_se_cancela(self):
        datos = {"deudor_id": self.deudor.id, "concepto": "Otro", "divisa_id": self.mxn.id,
                 "monto_inicial": "2000", "fecha": HOY.isoformat()}
        r = self.client.put(f"/api/treasury/deudas/{self.deuda.id}/", datos, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["editado_por"]["username"], "cajero")
        self._abonar("10")
        # Con abonos ya no cambian deudor, divisa, monto inicial ni fecha.
        datos["monto_inicial"] = "2500"
        self.assertEqual(self.client.put(f"/api/treasury/deudas/{self.deuda.id}/", datos, format="json").status_code, 400)
        r = self.client.post(f"/api/treasury/deudas/{self.deuda.id}/cancelar/", {"motivo": "x"}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_cancelar_deuda_pide_motivo_y_bloquea_abonos(self):
        url = f"/api/treasury/deudas/{self.deuda.id}/cancelar/"
        self.assertEqual(self.client.post(url, {"motivo": " "}, format="json").status_code, 400)
        r = self.client.post(url, {"motivo": "Se condonó"}, format="json")
        self.assertEqual((r.status_code, r.data["estado"]), (200, "cancelada"))
        self.assertIn("cancelada", str(self._abonar("10").data))

    def test_deudores_con_deuda_abierta_para_el_buscador(self):
        otro = Deudor.objects.create(nombre="Sin deuda", tipo=Deudor.EXTERNO, usuario=self.usuario)
        r = self.client.get("/api/treasury/deudores/", {"con_deuda_abierta": "1"})
        self.assertEqual([d["id"] for d in r.data], [self.deudor.id])
        self.assertEqual(r.data[0]["deudas_abiertas"], 1)
        r = self.client.get("/api/treasury/deudores/", {"q": "sin"})
        self.assertEqual([d["id"] for d in r.data], [otro.id])

    def test_estado_de_cuenta_pdf_y_excel(self):
        self._abonar("250")
        r = self.client.get(f"/api/treasury/deudores/{self.deudor.id}/estado-cuenta/pdf/")
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        r = self.client.get(f"/api/treasury/deudores/{self.deudor.id}/estado-cuenta/excel/")
        self.assertEqual(r.status_code, 200)

    def test_solo_lectura_consulta_pero_no_abona(self):
        lector = crear_usuario_con_area("lector", AREA_TESORERIA, solo_lectura=True)
        self.client.force_authenticate(user=lector)
        self.assertEqual(self.client.get("/api/treasury/deudas/").status_code, 200)
        self.assertEqual(self.client.get(f"/api/treasury/deudores/{self.deudor.id}/estado-cuenta/pdf/").status_code, 200)
        self.assertEqual(self._abonar("10").status_code, 403)
        self.assertEqual(self.client.post("/api/treasury/deudores/", {"nombre": "X", "tipo": "E"}).status_code, 403)


class HistorialMovimientoTests(APITestCase):
    def setUp(self):
        self.usuario = crear_usuario_con_area("cajero", AREA_TESORERIA)
        self.client.force_authenticate(user=self.usuario)
        self.mxn = Divisa.objects.create(codigo="MXN", nombre="Peso", simbolo="$")
        self.datos = {
            "fecha": HOY.isoformat(), "folio": "10", "tipo": "E", "autorizo": "Jefe", "beneficiario": "CFE",
            "concepto": "Pago luz", "divisas": [{"divisa_id": self.mxn.id, "cantidad": "500"}],
        }
        self.id = self.client.post("/api/treasury/movimientos/", self.datos, format="json").data["id"]

    def _historial(self):
        return self.client.get(f"/api/treasury/movimientos/{self.id}/historial/").data

    def test_cada_edicion_guarda_que_cambio(self):
        self.client.put(f"/api/treasury/movimientos/{self.id}/",
                        {**self.datos, "concepto": "Pago agua", "divisas": [{"divisa_id": self.mxn.id, "cantidad": "650"}]},
                        format="json")
        self.client.put(f"/api/treasury/movimientos/{self.id}/",
                        {**self.datos, "concepto": "Pago agua", "beneficiario": "CONAGUA",
                         "divisas": [{"divisa_id": self.mxn.id, "cantidad": "650"}]}, format="json")
        eventos = self._historial()[::-1]  # la API da lo más reciente primero
        self.assertEqual([e["tipo"] for e in eventos], ["registro", "edicion", "edicion"])
        self.assertEqual(eventos[1]["cambios"], [
            {"campo": "Concepto", "antes": "Pago luz", "despues": "Pago agua"},
            {"campo": "Cantidad MXN", "antes": "$500.00", "despues": "$650.00"},
        ])
        self.assertEqual(eventos[2]["cambios"], [{"campo": "Beneficiario", "antes": "CFE", "despues": "CONAGUA"}])
        self.assertEqual(eventos[2]["usuario"]["username"], "cajero")

    def test_guardar_sin_cambios_no_registra_edicion(self):
        self.client.put(f"/api/treasury/movimientos/{self.id}/", self.datos, format="json")
        self.assertEqual([e["tipo"] for e in self._historial()], ["registro"])

    def test_cancelacion_en_la_linea_de_tiempo(self):
        self.client.post(f"/api/treasury/movimientos/{self.id}/cancelar/", {"motivo": "Duplicado"}, format="json")
        eventos = self._historial()
        self.assertEqual((eventos[0]["tipo"], eventos[0]["motivo"]), ("cancelacion", "Duplicado"))


class ListadoMovimientosTests(APITestCase):
    def setUp(self):
        self.usuario = crear_usuario_con_area("cajero", AREA_TESORERIA)
        self.client.force_authenticate(user=self.usuario)
        self.mxn = Divisa.objects.create(codigo="MXN", nombre="Peso", simbolo="$")

    def _crear(self, folio, fecha):
        return self.client.post("/api/treasury/movimientos/", {
            "fecha": fecha.isoformat(), "folio": folio, "tipo": "I", "autorizo": "J", "beneficiario": "B",
            "concepto": "C", "divisas": [{"divisa_id": self.mxn.id, "cantidad": "10"}],
        }, format="json")

    def test_rango_son_los_extremos_de_los_registros(self):
        self.assertEqual(self.client.get("/api/treasury/movimientos/rango/").data,
                         {"desde": HOY.isoformat(), "hasta": HOY.isoformat()})
        self._crear("1", HOY - timedelta(days=40))
        self._crear("2", HOY - timedelta(days=3))
        self.assertEqual(self.client.get("/api/treasury/movimientos/rango/").data,
                         {"desde": (HOY - timedelta(days=40)).isoformat(), "hasta": (HOY - timedelta(days=3)).isoformat()})

    def test_listado_por_rango_sin_una_consulta_por_movimiento(self):
        for i in range(5):
            self._crear(str(i), HOY - timedelta(days=i))
        with self.assertNumQueries(5):  # permiso de área + lista + divisas + deudas (prefetch), sin importar cuántos
            r = self.client.get("/api/treasury/movimientos/", {"desde": (HOY - timedelta(days=3)).isoformat(), "hasta": HOY.isoformat()})
        self.assertEqual(len(r.data), 4)
        self.assertEqual(r.data[0]["archivos_count"], 0)


class CargosDeudaTests(APITestCase):
    """Un egreso ligado a una deuda es un cargo: sale de caja y aumenta lo que se debe."""

    def setUp(self):
        self.usuario = crear_usuario_con_area("cajero", AREA_TESORERIA)
        self.client.force_authenticate(user=self.usuario)
        self.mxn = Divisa.objects.create(codigo="MXN", nombre="Peso", simbolo="$")
        self.usd = Divisa.objects.create(codigo="USD", nombre="Dólar", simbolo="US$")
        self.deudor = Deudor.objects.create(nombre="Juan Pérez", tipo=Deudor.TRABAJADOR, usuario=self.usuario)
        self.folio = 0

    def _mov(self, tipo, cantidad, divisa=None, deuda_id=None, **extra):
        self.folio += 1
        if deuda_id:
            extra["aplicaciones"] = [{"deuda_id": deuda_id}]
        return self.client.post("/api/treasury/movimientos/", {
            "fecha": HOY.isoformat(), "folio": f"M-{self.folio}", "tipo": tipo, "autorizo": "Jefe",
            "beneficiario": "Juan Pérez", "concepto": "Préstamo",
            "divisas": [{"divisa_id": (divisa or self.mxn).id, "cantidad": cantidad}], **extra,
        }, format="json")

    def _deuda(self, deuda_id):
        return self.client.get(f"/api/treasury/deudas/{deuda_id}/").data

    def _nueva(self, cantidad="1000"):
        r = self._mov("E", cantidad, nueva_deuda={"deudor_id": self.deudor.id, "concepto": "Préstamo personal"})
        self.assertEqual(r.status_code, 201, r.data)
        return r.data["deudas"][0]["id"]

    def test_egreso_con_nueva_deuda_sale_de_caja_y_crea_la_deuda(self):
        deuda = self._deuda(self._nueva("1000"))
        self.assertEqual((deuda["monto_inicial"], deuda["cargado"], deuda["total"], deuda["saldo"], deuda["estado"]),
                         ("0.00", "1000.00", "1000.00", "1000.00", "abierta"))
        self.assertEqual(deuda["divisa"]["codigo"], "MXN")
        self.assertEqual(deuda["movimientos"][0]["tipo"], "E")
        self.assertEqual(saldos_al(HOY)[self.mxn.id], Decimal("-1000"))

    def test_otro_egreso_aumenta_la_deuda_y_el_abono_tiene_tope_en_el_total(self):
        deuda_id = self._nueva("1000")
        self.assertEqual(self._mov("E", "500", deuda_id=deuda_id).status_code, 201)
        self.assertEqual(self._deuda(deuda_id)["total"], "1500.00")
        self.assertEqual(self._mov("I", "1500.01", deuda_id=deuda_id).status_code, 400)
        self.assertEqual(self._mov("I", "1500", deuda_id=deuda_id).status_code, 201)
        self.assertEqual(self._deuda(deuda_id)["estado"], "pagada")
        # Un cargo nuevo la vuelve a abrir.
        self._mov("E", "200", deuda_id=deuda_id)
        self.assertEqual((self._deuda(deuda_id)["saldo"], self._deuda(deuda_id)["estado"]), ("200.00", "abierta"))

    def test_deuda_capturada_a_mano_mas_cargos(self):
        deuda = Deuda.objects.create(deudor=self.deudor, concepto="Venta a crédito", divisa=self.mxn,
                                     monto_inicial=Decimal("300"), fecha=HOY, usuario=self.usuario)
        self._mov("E", "700", deuda_id=deuda.id)
        datos = self._deuda(deuda.id)
        self.assertEqual((datos["monto_inicial"], datos["cargado"], datos["total"]), ("300.00", "700.00", "1000.00"))

    def test_el_cargo_va_en_la_divisa_de_la_deuda_y_solo_un_egreso_crea_deuda(self):
        deuda_id = self._nueva()
        self.assertIn("MXN", str(self._mov("E", "10", divisa=self.usd, deuda_id=deuda_id).data))
        r = self._mov("I", "10", nueva_deuda={"deudor_id": self.deudor.id, "concepto": "X"})
        self.assertIn("egreso", str(r.data))
        self.assertEqual(Deuda.objects.count(), 1)

    def test_no_se_cancela_un_cargo_si_los_abonos_lo_superan(self):
        deuda_id = self._nueva("1000")
        cargo = Deuda.objects.get(pk=deuda_id).movimientos.get()
        self._mov("I", "400", deuda_id=deuda_id)
        r = self.client.post(f"/api/treasury/movimientos/{cargo.id}/cancelar/", {"motivo": "x"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("abonado", r.data["detail"])

    def test_reducir_un_cargo_por_debajo_de_lo_abonado_no_se_permite(self):
        deuda_id = self._nueva("1000")
        cargo = Deuda.objects.get(pk=deuda_id).movimientos.get()
        self._mov("I", "600", deuda_id=deuda_id)
        datos = {"fecha": HOY.isoformat(), "folio": cargo.folio, "tipo": "E", "autorizo": "Jefe", "beneficiario": "J",
                 "concepto": "Préstamo", "aplicaciones": [{"deuda_id": deuda_id}], "divisas": [{"divisa_id": self.mxn.id, "cantidad": "500"}]}
        self.assertEqual(self.client.put(f"/api/treasury/movimientos/{cargo.id}/", datos, format="json").status_code, 400)
        datos["divisas"][0]["cantidad"] = "600"
        self.assertEqual(self.client.put(f"/api/treasury/movimientos/{cargo.id}/", datos, format="json").status_code, 200)

    def test_desde_deudas_se_exige_monto_inicial_y_con_movimientos_solo_concepto_y_notas(self):
        r = self.client.post("/api/treasury/deudas/", {"deudor_id": self.deudor.id, "concepto": "X", "divisa_id": self.mxn.id,
                                                        "monto_inicial": "0", "fecha": HOY.isoformat()}, format="json")
        self.assertEqual(r.status_code, 400)
        deuda_id = self._nueva("1000")
        base = {"deudor_id": self.deudor.id, "divisa_id": self.mxn.id, "monto_inicial": "0", "fecha": HOY.isoformat()}
        r = self.client.put(f"/api/treasury/deudas/{deuda_id}/", {**base, "concepto": "Préstamo para moto"}, format="json")
        self.assertEqual((r.status_code, r.data["concepto"]), (200, "Préstamo para moto"))
        r = self.client.put(f"/api/treasury/deudas/{deuda_id}/", {**base, "concepto": "Y", "monto_inicial": "50"}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_un_movimiento_ligado_no_se_elimina(self):
        deuda_id = self._nueva()
        cargo = Deuda.objects.get(pk=deuda_id).movimientos.get()
        self.assertEqual(self.client.delete(f"/api/treasury/movimientos/{cargo.id}/").status_code, 400)


class RepartoAbonoTests(APITestCase):
    """Un ingreso puede abonar a varias deudas, repartiendo exactamente su total."""

    def setUp(self):
        self.usuario = crear_usuario_con_area("cajero", AREA_TESORERIA)
        self.client.force_authenticate(user=self.usuario)
        self.mxn = Divisa.objects.create(codigo="MXN", nombre="Peso", simbolo="$")
        self.deudor = Deudor.objects.create(nombre="Esdras", tipo=Deudor.TRABAJADOR, usuario=self.usuario)
        self.prestamo = self._deuda_con("PRESTAMO PERSONAL", "3000")
        self.producto = self._deuda_con("CUENTA PRODUCTO", "1500")

    def _deuda_con(self, concepto, monto):
        return Deuda.objects.create(deudor=self.deudor, concepto=concepto, divisa=self.mxn, monto_inicial=Decimal(monto),
                                    fecha=HOY - timedelta(days=10), usuario=self.usuario)

    def _datos(self, cantidad="2000", aplicaciones=None):
        datos = {
            "fecha": HOY.isoformat(), "folio": "8050", "tipo": "I", "autorizo": "CP", "beneficiario": "Esdras",
            "concepto": "Abono", "divisas": [{"divisa_id": self.mxn.id, "cantidad": cantidad}],
        }
        if aplicaciones is not None:
            datos["aplicaciones"] = [{"deuda_id": d.id, "monto": m} for d, m in aplicaciones]
        return datos

    def _saldo(self, deuda):
        return self.client.get(f"/api/treasury/deudas/{deuda.id}/").data["saldo"]

    def test_un_ingreso_abona_a_dos_deudas(self):
        r = self.client.post("/api/treasury/movimientos/",
                             self._datos(aplicaciones=[(self.prestamo, "1200"), (self.producto, "800")]), format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual([(d["concepto"], d["monto"]) for d in r.data["deudas"]],
                         [("PRESTAMO PERSONAL", "1200.00"), ("CUENTA PRODUCTO", "800.00")])
        self.assertEqual((self._saldo(self.prestamo), self._saldo(self.producto)), ("1800.00", "700.00"))
        abono = self.client.get(f"/api/treasury/deudas/{self.producto.id}/").data["movimientos"][0]
        self.assertEqual((abono["cantidad"], abono["monto_movimiento"]), ("800.00", "2000.00"))
        self.assertEqual(saldos_al(HOY)[self.mxn.id], Decimal("2000"))

    def test_el_reparto_debe_sumar_el_total_y_caber_en_cada_saldo(self):
        url = "/api/treasury/movimientos/"
        r = self.client.post(url, self._datos(aplicaciones=[(self.prestamo, "1000"), (self.producto, "900")]), format="json")
        self.assertIn("sumar el total", str(r.data["aplicaciones"]))
        r = self.client.post(url, self._datos(aplicaciones=[(self.prestamo, "400"), (self.producto, "1600")]), format="json")
        self.assertIn("CUENTA PRODUCTO", str(r.data["aplicaciones"]))
        r = self.client.post(url, self._datos(aplicaciones=[(self.prestamo, "1000"), (self.prestamo, "1000")]), format="json")
        self.assertIn("dos veces", str(r.data["aplicaciones"]))
        r = self.client.post(url, self._datos(aplicaciones=[(self.prestamo, "1000"), (self.producto, None)]), format="json")
        self.assertEqual(r.status_code, 400)
        self.assertFalse(MovimientoTesoreria.objects.exists())

    def test_un_egreso_solo_carga_a_una_deuda(self):
        datos = self._datos(aplicaciones=[(self.prestamo, "1000"), (self.producto, "1000")])
        datos["tipo"] = "E"
        r = self.client.post("/api/treasury/movimientos/", datos, format="json")
        self.assertIn("una deuda", str(r.data["aplicaciones"]))

    def test_corregir_un_ingreso_ya_registrado_repartiendolo_queda_en_ambos_historiales(self):
        # Caso del folio 8050: capturado sin ligar y luego repartido al editarlo.
        r = self.client.post("/api/treasury/movimientos/", self._datos(), format="json")
        self.assertEqual((r.status_code, r.data["deudas"]), (201, []))
        url = f"/api/treasury/movimientos/{r.data['id']}/"
        r = self.client.put(url, self._datos(aplicaciones=[(self.prestamo, "1000"), (self.producto, "1000")]), format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(r.data["editado"])
        self.assertEqual((self._saldo(self.prestamo), self._saldo(self.producto)), ("2000.00", "500.00"))

        historial = self.client.get(f"{url}historial/").data
        cambio = [e for e in historial if e["tipo"] == "edicion"][0]["cambios"][0]
        self.assertEqual(cambio["campo"], "Deudas")
        self.assertEqual(cambio["antes"], "")
        self.assertIn("PRESTAMO PERSONAL $1,000.00", cambio["despues"])
        self.assertIn("CUENTA PRODUCTO $1,000.00", cambio["despues"])
        movs = self.client.get(f"/api/treasury/deudas/{self.prestamo.id}/").data["movimientos"]
        self.assertEqual([(m["folio"], m["cantidad"], m["editado"]) for m in movs], [("8050", "1000.00", True)])

        # Cambiar la división también queda registrado; sin la clave se conserva.
        r = self.client.put(url, self._datos(aplicaciones=[(self.prestamo, "1500"), (self.producto, "500")]), format="json")
        self.assertEqual((self._saldo(self.prestamo), self._saldo(self.producto)), ("1500.00", "1000.00"))
        datos = self._datos()
        datos["concepto"] = "Abono préstamo y producto"
        self.assertEqual(self.client.put(url, datos, format="json").status_code, 200)
        self.assertEqual(self._saldo(self.prestamo), "1500.00")
        # Con varias deudas, cambiar el total sin rehacer el reparto no cuadra.
        self.assertEqual(self.client.put(url, self._datos(cantidad="2500"), format="json").status_code, 400)
        # Quitar el reparto devuelve los saldos.
        self.client.put(url, self._datos(aplicaciones=[]), format="json")
        self.assertEqual((self._saldo(self.prestamo), self._saldo(self.producto)), ("3000.00", "1500.00"))

    def test_cancelar_el_ingreso_revierte_todas_sus_partes(self):
        r = self.client.post("/api/treasury/movimientos/",
                             self._datos(aplicaciones=[(self.prestamo, "1000"), (self.producto, "1000")]), format="json")
        self.client.post(f"/api/treasury/movimientos/{r.data['id']}/cancelar/", {"motivo": "Error"}, format="json")
        self.assertEqual((self._saldo(self.prestamo), self._saldo(self.producto)), ("3000.00", "1500.00"))

    def test_estado_de_cuenta_marca_el_pago_compartido(self):
        self.client.post("/api/treasury/movimientos/",
                         self._datos(aplicaciones=[(self.prestamo, "1000"), (self.producto, "1000")]), format="json")
        from io import BytesIO
        from openpyxl import load_workbook
        r = self.client.get(f"/api/treasury/deudores/{self.deudor.id}/estado-cuenta/excel/")
        hoja = load_workbook(BytesIO(r.content))["Cargos y abonos"]
        filas = [fila for fila in hoja.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(f[7] for f in filas), [1000, 1000])
        self.assertTrue(all("pago compartido" in f[5] for f in filas))
        r = self.client.get(f"/api/treasury/deudores/{self.deudor.id}/estado-cuenta/pdf/")
        self.assertEqual(r.status_code, 200)
