from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from django.core.exceptions import ValidationError
from django.db.models import Count, DecimalField, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import CorteCaja, Divisa, MovimientoDivisa, MovimientoTesoreria

PREVIEW_LIMIT = 20

COLOR_ENCABEZADO = "0E6B6F"
COLOR_FILA_CANCELADA = "F0F0F0"
COLOR_ENCABEZADO_RL = colors.HexColor(f"#{COLOR_ENCABEZADO}")
COLOR_FILA_CANCELADA_RL = colors.HexColor(f"#{COLOR_FILA_CANCELADA}")


def _parsear_fecha(valor, nombre_campo):
    fecha = parse_date(valor) if valor else None
    if not fecha:
        raise ValidationError(f"La {nombre_campo} no es válida. Usa el formato AAAA-MM-DD.")
    return fecha


def obtener_lineas_filtradas(params):
    """
    params: QueryDict (request.query_params). Devuelve un queryset de
    MovimientoDivisa —una fila por línea de divisa, no por movimiento—
    con los filtros del reporte ya aplicados. Lanza ValidationError con un
    mensaje entendible si el rango de fechas falta o es inválido.
    """
    fecha_inicio_raw = params.get("fecha_inicio")
    fecha_fin_raw = params.get("fecha_fin")

    if not fecha_inicio_raw or not fecha_fin_raw:
        raise ValidationError("Debes indicar la fecha de inicio y la fecha de fin.")

    fecha_inicio = _parsear_fecha(fecha_inicio_raw, "fecha de inicio")
    fecha_fin = _parsear_fecha(fecha_fin_raw, "fecha de fin")

    if fecha_inicio > fecha_fin:
        raise ValidationError("La fecha de inicio no puede ser posterior a la fecha de fin.")

    lineas = (
        MovimientoDivisa.objects
        .select_related("movimiento", "movimiento__corte", "divisa")
        .filter(movimiento__fecha__date__gte=fecha_inicio, movimiento__fecha__date__lte=fecha_fin)
    )

    corte_id = params.get("corte")
    if corte_id:
        lineas = lineas.filter(movimiento__corte_id=corte_id)

    tipo = params.get("tipo")
    if tipo in (MovimientoTesoreria.INGRESO, MovimientoTesoreria.EGRESO):
        lineas = lineas.filter(movimiento__tipo=tipo)

    estado = params.get("estado")
    if estado == "activos":
        lineas = lineas.filter(movimiento__cancelado=False)
    elif estado == "cancelados":
        lineas = lineas.filter(movimiento__cancelado=True)

    beneficiario = params.get("beneficiario")
    if beneficiario:
        lineas = lineas.filter(movimiento__beneficiario__icontains=beneficiario)

    divisa_ids = params.getlist("divisa") if hasattr(params, "getlist") else None
    if divisa_ids:
        lineas = lineas.filter(divisa_id__in=divisa_ids)

    return lineas.order_by("movimiento__fecha", "movimiento__creado", "id")


def calcular_resumen_por_divisa(lineas):
    """
    Ingresos/egresos solo cuentan movimientos activos (los cancelados no
    representan dinero real que entró o salió), pero se reportan aparte
    cuántos movimientos cancelados hay por divisa para trazabilidad.
    """
    campo_decimal = DecimalField(max_digits=18, decimal_places=2)
    cero = Value(Decimal("0"), output_field=campo_decimal)

    filtro_ingreso_activo = Q(movimiento__tipo=MovimientoTesoreria.INGRESO, movimiento__cancelado=False)
    filtro_egreso_activo = Q(movimiento__tipo=MovimientoTesoreria.EGRESO, movimiento__cancelado=False)

    return list(
        lineas.values("divisa_id", "divisa__codigo", "divisa__nombre", "divisa__simbolo")
        .annotate(
            ingresos=Coalesce(Sum("cantidad", filter=filtro_ingreso_activo), cero),
            egresos=Coalesce(Sum("cantidad", filter=filtro_egreso_activo), cero),
            movimientos_activos=Count("id", filter=Q(movimiento__cancelado=False)),
            movimientos_cancelados=Count("id", filter=Q(movimiento__cancelado=True)),
        )
        .order_by("divisa__codigo")
    )


def _estilizar_encabezado(ws, num_columnas):
    for col in range(1, num_columnas + 1):
        celda = ws.cell(row=1, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 20


def construir_libro_excel(lineas, resumen):
    wb = Workbook()

    ws = wb.active
    ws.title = "Detalle"
    encabezados = ["Folio", "Fecha", "Corte", "Tipo", "Autorizó", "Beneficiario", "Concepto", "Divisa", "Cantidad", "Estado"]
    ws.append(encabezados)
    _estilizar_encabezado(ws, len(encabezados))

    fill_cancelado = PatternFill("solid", fgColor=COLOR_FILA_CANCELADA)
    font_cancelado = Font(strike=True, color="808080")

    for linea in lineas:
        m = linea.movimiento
        ws.append([
            m.folio,
            timezone.localtime(m.fecha).strftime("%d/%m/%Y"),
            m.corte_id or "",
            m.get_tipo_display(),
            m.autorizo,
            m.beneficiario,
            m.concepto,
            linea.divisa.codigo,
            float(linea.cantidad),
            "Cancelado" if m.cancelado else "Activo",
        ])
        fila = ws.max_row
        ws.cell(row=fila, column=9).number_format = "#,##0.00"
        if m.cancelado:
            for col in range(1, len(encabezados) + 1):
                celda = ws.cell(row=fila, column=col)
                celda.font = font_cancelado
                celda.fill = fill_cancelado

    anchos = {"A": 14, "B": 12, "C": 8, "D": 10, "E": 20, "F": 28, "G": 32, "H": 8, "I": 14, "J": 10}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    ws_resumen = wb.create_sheet("Resumen por divisa")
    encabezados_resumen = ["Divisa", "Ingresos", "Egresos", "Neto", "Movimientos activos", "Movimientos cancelados"]
    ws_resumen.append(encabezados_resumen)
    _estilizar_encabezado(ws_resumen, len(encabezados_resumen))

    for fila_datos in resumen:
        neto = fila_datos["ingresos"] - fila_datos["egresos"]
        ws_resumen.append([
            fila_datos["divisa__codigo"],
            float(fila_datos["ingresos"]),
            float(fila_datos["egresos"]),
            float(neto),
            fila_datos["movimientos_activos"],
            fila_datos["movimientos_cancelados"],
        ])
        fila = ws_resumen.max_row
        for col in (2, 3, 4):
            ws_resumen.cell(row=fila, column=col).number_format = "#,##0.00"

    anchos_resumen = {"A": 10, "B": 16, "C": 16, "D": 16, "E": 18, "F": 20}
    for col, ancho in anchos_resumen.items():
        ws_resumen.column_dimensions[col].width = ancho

    return wb


def describir_filtros(params):
    """
    Arma una lista de frases legibles ("Tipo: Ingreso", "Divisas: MXN, USD")
    con los filtros distintos al rango de fechas, para mostrarlos como
    subtítulo en el PDF. El rango de fechas se muestra aparte porque siempre
    está presente.
    """
    partes = []

    corte_id = params.get("corte")
    if corte_id:
        corte = CorteCaja.objects.filter(pk=corte_id).first()
        if corte:
            partes.append(f"Corte del {timezone.localtime(corte.fecha).strftime('%d/%m/%Y')} (#{corte.id})")
        else:
            partes.append(f"Corte #{corte_id}")

    tipo_nombre = dict(MovimientoTesoreria.TIPO_CHOICES).get(params.get("tipo"))
    if tipo_nombre:
        partes.append(f"Tipo: {tipo_nombre}")

    estado = params.get("estado")
    if estado == "activos":
        partes.append("Estado: Activos")
    elif estado == "cancelados":
        partes.append("Estado: Cancelados")

    beneficiario = params.get("beneficiario")
    if beneficiario:
        partes.append(f'Beneficiario contiene: "{beneficiario}"')

    divisa_ids = params.getlist("divisa") if hasattr(params, "getlist") else None
    if divisa_ids:
        codigos = list(Divisa.objects.filter(id__in=divisa_ids).order_by("codigo").values_list("codigo", flat=True))
        if codigos:
            partes.append("Divisas: " + ", ".join(codigos))

    return partes


ESTILOS = getSampleStyleSheet()

_ESTILOS_CELDA_DETALLE = {
    (alineacion, cancelado): ParagraphStyle(
        f"detalle-{alineacion}-{'cancelado' if cancelado else 'activo'}",
        parent=ESTILOS["Normal"],
        fontSize=7.5,
        leading=9,
        alignment=alineacion,
        textColor=colors.HexColor("#808080") if cancelado else colors.black,
    )
    for alineacion in (TA_LEFT, TA_CENTER, TA_RIGHT)
    for cancelado in (False, True)
}


def _celda_detalle(texto, alineacion=TA_LEFT, cancelado=False):
    texto_escapado = escape(str(texto))
    if cancelado:
        texto_escapado = f"<strike>{texto_escapado}</strike>"
    return Paragraph(texto_escapado, _ESTILOS_CELDA_DETALLE[(alineacion, cancelado)])


def _pie_pagina_pdf(canvas, documento):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#888888"))
    ancho_pagina, _ = documento.pagesize
    canvas.drawRightString(ancho_pagina - 1.2 * cm, 0.7 * cm, f"Página {documento.page}")
    canvas.restoreState()


def construir_pdf_movimientos(lineas, resumen, fecha_inicio, fecha_fin, filtros_descripcion):
    lineas = list(lineas)

    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title="Reporte de movimientos de Tesorería",
    )

    estilo_empresa = ParagraphStyle(
        "empresa", parent=ESTILOS["Heading1"], fontSize=15, textColor=COLOR_ENCABEZADO_RL, spaceAfter=1,
    )
    estilo_titulo = ParagraphStyle(
        "titulo", parent=ESTILOS["Normal"], fontSize=12, textColor=colors.black, spaceAfter=6,
    )
    estilo_subtitulo = ParagraphStyle(
        "subtitulo", parent=ESTILOS["Normal"], fontSize=9.5, textColor=colors.HexColor("#444444"), spaceAfter=2,
    )
    estilo_generado = ParagraphStyle(
        "generado", parent=ESTILOS["Normal"], fontSize=8, textColor=colors.HexColor("#999999"), spaceBefore=4,
    )
    estilo_h2 = ParagraphStyle(
        "h2", parent=ESTILOS["Heading2"], fontSize=12, textColor=COLOR_ENCABEZADO_RL, spaceAfter=4,
    )

    elementos = [
        Paragraph("Importadora y Exportadora de Mariscos", estilo_empresa),
        Paragraph("Reporte de movimientos de Tesorería", estilo_titulo),
        Paragraph(f"Del {fecha_inicio.strftime('%d/%m/%Y')} al {fecha_fin.strftime('%d/%m/%Y')}", estilo_subtitulo),
    ]
    if filtros_descripcion:
        elementos.append(Paragraph("Filtros: " + " · ".join(filtros_descripcion), estilo_subtitulo))
    elementos.append(Paragraph(
        f"Generado el {timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M')}",
        estilo_generado,
    ))
    elementos.append(Spacer(1, 0.6 * cm))

    # --- Resumen por divisa ---
    elementos.append(Paragraph("Resumen por divisa", estilo_h2))
    encabezados_resumen = ["Divisa", "Ingresos", "Egresos", "Neto", "Mov. activos", "Mov. cancelados"]
    filas_resumen = [encabezados_resumen]
    for fila in resumen:
        neto = fila["ingresos"] - fila["egresos"]
        filas_resumen.append([
            fila["divisa__codigo"],
            f'{fila["ingresos"]:,.2f}',
            f'{fila["egresos"]:,.2f}',
            f'{neto:,.2f}',
            str(fila["movimientos_activos"]),
            str(fila["movimientos_cancelados"]),
        ])

    if len(filas_resumen) == 1:
        elementos.append(Paragraph("No hay movimientos para los filtros seleccionados.", ESTILOS["Normal"]))
    else:
        tabla_resumen = Table(filas_resumen, colWidths=[2.5 * cm, 3.4 * cm, 3.4 * cm, 3.4 * cm, 3 * cm, 3.3 * cm])
        tabla_resumen.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 0), (0, -1), "LEFT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elementos.append(tabla_resumen)

    elementos.append(Spacer(1, 0.7 * cm))

    # --- Detalle ---
    elementos.append(Paragraph("Detalle de movimientos", estilo_h2))
    encabezados_detalle = ["Folio", "Fecha", "Corte", "Tipo", "Autorizó", "Beneficiario", "Concepto", "Divisa", "Cantidad", "Estado"]
    filas_detalle = [encabezados_detalle]

    for linea in lineas:
        m = linea.movimiento
        cancelado = m.cancelado
        filas_detalle.append([
            _celda_detalle(m.folio, TA_LEFT, cancelado),
            _celda_detalle(timezone.localtime(m.fecha).strftime("%d/%m/%Y"), TA_CENTER, cancelado),
            _celda_detalle(f"#{m.corte_id}" if m.corte_id else "-", TA_CENTER, cancelado),
            _celda_detalle(m.get_tipo_display(), TA_CENTER, cancelado),
            _celda_detalle(m.autorizo, TA_LEFT, cancelado),
            _celda_detalle(m.beneficiario, TA_LEFT, cancelado),
            _celda_detalle(m.concepto, TA_LEFT, cancelado),
            _celda_detalle(linea.divisa.codigo, TA_CENTER, cancelado),
            _celda_detalle(f"{linea.cantidad:,.2f}", TA_RIGHT, cancelado),
            _celda_detalle("Cancelado" if cancelado else "Activo", TA_CENTER, cancelado),
        ])

    if len(filas_detalle) == 1:
        elementos.append(Paragraph("No hay movimientos para los filtros seleccionados.", ESTILOS["Normal"]))
    else:
        anchos_detalle = [2 * cm, 2.3 * cm, 1.5 * cm, 1.7 * cm, 3 * cm, 3.8 * cm, 4.6 * cm, 1.4 * cm, 2.3 * cm, 1.9 * cm]
        tabla_detalle = Table(filas_detalle, colWidths=anchos_detalle, repeatRows=1)

        estilo_detalle = [
            ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8.5),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        for indice, linea in enumerate(lineas, start=1):
            if linea.movimiento.cancelado:
                estilo_detalle.append(("BACKGROUND", (0, indice), (-1, indice), COLOR_FILA_CANCELADA_RL))
        tabla_detalle.setStyle(TableStyle(estilo_detalle))
        elementos.append(tabla_detalle)

    documento.build(elementos, onFirstPage=_pie_pagina_pdf, onLaterPages=_pie_pagina_pdf)
    buffer.seek(0)
    return buffer
