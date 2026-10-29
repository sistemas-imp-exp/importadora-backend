"""
Estado de cuenta de un deudor: cada deuda con su monto inicial, cargos,
total, abonado y saldo, el historial de cargos (egresos) y abonos (ingresos)
de caja ligados a ella, y lo pendiente por divisa. Recibe las deudas
anotadas con Deuda.objects.con_saldo().
"""
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Deuda, MovimientoTesoreria
from .reportes import (
    COLOR_ENCABEZADO_RL,
    COLOR_FILA_CANCELADA,
    COLOR_FILA_CANCELADA_RL,
    ESTILOS,
    _celda_detalle,
    _estilizar_encabezado,
    _pie_pagina_pdf,
)

ESTADOS = {Deuda.ABIERTA: "Abierta", Deuda.PAGADA: "Pagada", Deuda.CANCELADA: "Cancelada"}


def _estado(deuda):
    return Deuda.estado_de(deuda.cancelada, deuda.saldo)


def _cantidad(movimiento):
    """Lo aplicado a la deuda (un abono repartido cuenta solo su parte)."""
    return movimiento.monto_aplicado


def _movimientos(deuda):
    return deuda.movimientos_aplicados()


def _concepto(movimiento, simbolo):
    """Concepto del folio; si el pago se repartió entre deudas, lo dice con el total del folio."""
    if movimiento.monto_aplicado == movimiento.monto_movimiento:
        return movimiento.concepto
    return f"{movimiento.concepto} (pago compartido; total del folio {_dinero(simbolo, movimiento.monto_movimiento)})"


def _es_cargo(movimiento):
    return movimiento.tipo == MovimientoTesoreria.EGRESO


def pendiente_por_divisa(deudas):
    """{codigo: (simbolo, saldo)} de las deudas abiertas."""
    totales = {}
    for deuda in deudas:
        if _estado(deuda) != Deuda.ABIERTA:
            continue
        simbolo, total = totales.get(deuda.divisa.codigo, (deuda.divisa.simbolo, Decimal("0")))
        totales[deuda.divisa.codigo] = (simbolo, total + deuda.saldo)
    return dict(sorted(totales.items()))


def _dinero(simbolo, valor):
    return f"{simbolo}{valor:,.2f}"


def construir_pdf_estado_cuenta(deudor, deudas):
    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=f"Estado de cuenta - {deudor.nombre}",
    )
    estilo_empresa = ParagraphStyle(
        "empresa", parent=ESTILOS["Heading1"], fontSize=15, textColor=COLOR_ENCABEZADO_RL, spaceAfter=1,
    )
    estilo_titulo = ParagraphStyle("titulo", parent=ESTILOS["Normal"], fontSize=12, spaceAfter=6)
    estilo_subtitulo = ParagraphStyle(
        "subtitulo", parent=ESTILOS["Normal"], fontSize=9.5, textColor=colors.HexColor("#444444"), spaceAfter=2,
    )
    estilo_generado = ParagraphStyle(
        "generado", parent=ESTILOS["Normal"], fontSize=8, textColor=colors.HexColor("#999999"), spaceBefore=4,
    )
    estilo_h2 = ParagraphStyle(
        "h2", parent=ESTILOS["Heading2"], fontSize=12, textColor=COLOR_ENCABEZADO_RL, spaceAfter=4,
    )
    estilo_deuda = ParagraphStyle("deuda", parent=ESTILOS["Normal"], fontSize=9.5, spaceAfter=3)
    estilo_nota = ParagraphStyle(
        "nota", parent=ESTILOS["Normal"], fontSize=8, textColor=colors.HexColor("#777777"), spaceAfter=2,
    )
    estilo_tabla = [
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_ENCABEZADO_RL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]

    elementos = [
        Paragraph("Importadora y Exportadora de Mariscos de Centro América y el Caribe", estilo_empresa),
        Paragraph("Estado de cuenta", estilo_titulo),
        Paragraph(f"<b>{escape(deudor.nombre)}</b> · {deudor.get_tipo_display()}"
                  + (f" · Tel. {escape(deudor.telefono)}" if deudor.telefono else ""), estilo_subtitulo),
        Paragraph(f"Generado el {timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M')}", estilo_generado),
        Spacer(1, 0.5 * cm),
        Paragraph("Saldo pendiente", estilo_h2),
    ]

    pendiente = pendiente_por_divisa(deudas)
    if pendiente:
        tabla = Table(
            [["Divisa", "Pendiente"]] + [[codigo, _dinero(simbolo, total)] for codigo, (simbolo, total) in pendiente.items()],
            colWidths=[2.5 * cm, 4 * cm],
        )
        tabla.setStyle(TableStyle(estilo_tabla + [("ALIGN", (1, 0), (1, -1), "RIGHT")]))
        elementos.append(tabla)
    else:
        elementos.append(Paragraph("Sin deudas abiertas.", ESTILOS["Normal"]))

    elementos += [Spacer(1, 0.5 * cm), Paragraph("Deudas", estilo_h2)]
    if not deudas:
        elementos.append(Paragraph("Sin deudas registradas.", ESTILOS["Normal"]))

    for deuda in deudas:
        simbolo = deuda.divisa.simbolo
        estado = _estado(deuda)
        bloque = [
            Paragraph(
                f"<b>{escape(deuda.concepto)}</b> · {deuda.fecha.strftime('%d/%m/%Y')} · {deuda.divisa.codigo} · "
                f"{ESTADOS[estado]}",
                estilo_deuda,
            ),
            Paragraph(
                (f"Monto inicial {_dinero(simbolo, deuda.monto_inicial)} · " if deuda.monto_inicial else "")
                + (f"Cargos {_dinero(simbolo, deuda.cargado)} · " if deuda.cargado else "")
                + f"Total {_dinero(simbolo, deuda.total)} · Abonado {_dinero(simbolo, deuda.abonado)} · "
                f"<b>Saldo {_dinero(simbolo, deuda.saldo)}</b>",
                estilo_deuda,
            ),
        ]
        if deuda.cancelada:
            bloque.append(Paragraph(f"Cancelada: {escape(deuda.motivo_cancelacion)}", estilo_nota))
        movimientos = _movimientos(deuda)
        if movimientos:
            filas = [["Folio", "Fecha", "Tipo", "Concepto", "Cargo", "Abono", "Estado"]]
            for mov in movimientos:
                cancelado = mov.cancelado
                importe = _dinero(simbolo, _cantidad(mov))
                filas.append([
                    _celda_detalle(mov.folio, TA_LEFT, cancelado),
                    _celda_detalle(mov.fecha.strftime("%d/%m/%Y"), TA_CENTER, cancelado),
                    _celda_detalle("Cargo" if _es_cargo(mov) else "Abono", TA_CENTER, cancelado),
                    _celda_detalle(_concepto(mov, simbolo), TA_LEFT, cancelado),
                    _celda_detalle(importe if _es_cargo(mov) else "", TA_RIGHT, cancelado),
                    _celda_detalle("" if _es_cargo(mov) else importe, TA_RIGHT, cancelado),
                    _celda_detalle("Cancelado" if cancelado else "Activo", TA_CENTER, cancelado),
                ])
            tabla = Table(filas, colWidths=[1.9 * cm, 2.1 * cm, 1.5 * cm, 5.6 * cm, 2.4 * cm, 2.4 * cm, 1.9 * cm], repeatRows=1)
            estilo = list(estilo_tabla)
            for indice, mov in enumerate(movimientos, start=1):
                if mov.cancelado:
                    estilo.append(("BACKGROUND", (0, indice), (-1, indice), COLOR_FILA_CANCELADA_RL))
            tabla.setStyle(TableStyle(estilo))
            bloque.append(tabla)
        else:
            bloque.append(Paragraph("Sin cargos ni abonos de caja.", estilo_nota))
        bloque.append(Spacer(1, 0.45 * cm))
        elementos.append(KeepTogether(bloque))

    documento.build(elementos, onFirstPage=_pie_pagina_pdf, onLaterPages=_pie_pagina_pdf)
    buffer.seek(0)
    return buffer


def construir_excel_estado_cuenta(deudor, deudas):
    wb = Workbook()
    fill_cancelado = PatternFill("solid", fgColor=COLOR_FILA_CANCELADA)
    font_cancelado = Font(strike=True, color="808080")

    ws = wb.active
    ws.title = "Deudas"
    encabezados = ["Fecha", "Concepto", "Divisa", "Monto inicial", "Cargos", "Total", "Abonado", "Saldo", "Estado", "Notas"]
    ws.append(encabezados)
    _estilizar_encabezado(ws, len(encabezados))
    for deuda in deudas:
        ws.append([
            deuda.fecha, deuda.concepto, deuda.divisa.codigo,
            float(deuda.monto_inicial), float(deuda.cargado), float(deuda.total), float(deuda.abonado), float(deuda.saldo),
            ESTADOS[_estado(deuda)], deuda.motivo_cancelacion if deuda.cancelada else deuda.notas,
        ])
        fila = ws.max_row
        ws.cell(row=fila, column=1).number_format = "DD/MM/YYYY"
        for col in (4, 5, 6, 7, 8):
            ws.cell(row=fila, column=col).number_format = "#,##0.00"
        if deuda.cancelada:
            for col in range(1, len(encabezados) + 1):
                ws.cell(row=fila, column=col).font = font_cancelado
                ws.cell(row=fila, column=col).fill = fill_cancelado
    for col, ancho in zip("ABCDEFGHIJ", (12, 36, 8, 14, 14, 14, 14, 14, 11, 36)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    ws_movs = wb.create_sheet("Cargos y abonos")
    encabezados = ["Deuda", "Divisa", "Folio", "Fecha", "Tipo", "Concepto", "Cargo", "Abono", "Estado", "Motivo cancelación"]
    ws_movs.append(encabezados)
    _estilizar_encabezado(ws_movs, len(encabezados))
    for deuda in deudas:
        for mov in _movimientos(deuda):
            cargo = _es_cargo(mov)
            ws_movs.append([
                deuda.concepto, deuda.divisa.codigo, mov.folio, mov.fecha, "Cargo" if cargo else "Abono",
                _concepto(mov, deuda.divisa.simbolo),
                float(_cantidad(mov)) if cargo else None, None if cargo else float(_cantidad(mov)),
                "Cancelado" if mov.cancelado else "Activo", mov.motivo_cancelacion,
            ])
            fila = ws_movs.max_row
            ws_movs.cell(row=fila, column=4).number_format = "DD/MM/YYYY"
            for col in (7, 8):
                ws_movs.cell(row=fila, column=col).number_format = "#,##0.00"
            if mov.cancelado:
                for col in range(1, len(encabezados) + 1):
                    ws_movs.cell(row=fila, column=col).font = font_cancelado
                    ws_movs.cell(row=fila, column=col).fill = fill_cancelado
    for col, ancho in zip("ABCDEFGHIJ", (30, 8, 12, 12, 8, 32, 14, 14, 11, 30)):
        ws_movs.column_dimensions[col].width = ancho
    ws_movs.freeze_panes = "A2"

    ws_resumen = wb.create_sheet("Pendiente por divisa")
    ws_resumen.append(["Deudor", deudor.nombre])
    ws_resumen.append(["Tipo", deudor.get_tipo_display()])
    ws_resumen.append([])
    ws_resumen.append(["Divisa", "Pendiente"])
    for col in (1, 2):
        celda = ws_resumen.cell(row=4, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0E6B6F")
    for codigo, (_, total) in pendiente_por_divisa(deudas).items():
        ws_resumen.append([codigo, float(total)])
        ws_resumen.cell(row=ws_resumen.max_row, column=2).number_format = "#,##0.00"
    ws_resumen.column_dimensions["A"].width = 12
    ws_resumen.column_dimensions["B"].width = 30
    return wb
